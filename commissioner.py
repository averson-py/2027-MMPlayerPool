import tkinter as tk
from tkinter import filedialog, messagebox, ttk
from pathlib import Path
from datetime import datetime, timedelta
import difflib
import json
import os
import re
import unicodedata
import urllib.parse
import urllib.request

from openpyxl import load_workbook

DRAFT_ORDER_SHEET = "DraftOrder"
TEAM_DATA_SHEET = "TeamData"
DRAFT_SHEET = "Draft"
TEAMS_OUT_SHEET = "TeamsOut"
GAME_RESULTS_SHEET = "GameResults"
CONFIG_SHEET = "Config"
PICK_TIMER_HEADER_CELL = "K1"
PICK_TIMER_VALUE_CELL = "K2"
PICK_TIMER_HEADER = "Pick Timer (min)"
DEFAULT_PICK_TIMER_MINUTES = 2

DRAFTER_START_ROW = 2
PLAYER_START_ROW = 2
DRAFTER_COLUMN = 1
TOTAL_ROUNDS = 5

SEED_GAMES_FIRST_ROW = 2
SEED_GAMES_LAST_ROW = 17
SEED_GAMES_SEED_COLUMN = "H"
SEED_GAMES_AVERAGE_COLUMN = "I"

ROUND_OPTIONS = [
    "First Four",
    "First Round",
    "Second Round",
    "Sweet Sixteen",
    "Elite Eight",
    "Final Four",
    "National Title",
]

STAT_COLUMN_MAP = {
    round_name: {
        "PTS": 4 + index,
        "REB": 10 + index,
        "AST": 16 + index,
    }
    for index, round_name in enumerate(ROUND_OPTIONS)
}

TEAM_HEADERS = [
    "Region",
    "Seed",
    "Team",
    "Player",
    "GP",
    "MPG",
    "PPG",
    "RPG",
    "APG",
    "SPG",
    "BPG",
    "TO",
    "FG%",
    "FT%",
    "3PT%"
]

DISPLAY_HEADERS = (
    TEAM_HEADERS[:4]
    + ["Tot. FPts", "Proj. FPts"]
    + TEAM_HEADERS[4:]
)

PICK_HEADERS = [
    "Draft Position",
    "Draftee",
    "Player Name",
    "Region",
    "Seed",
    "Team",
    "Tot. FPts",
    "PPG",
    "RPG",
    "APG"
]

PREVIOUS_SELECTION_HEADERS = PICK_HEADERS.copy()

ESPN_BASE_URL = (
    "https://site.api.espn.com/apis/site/v2/sports/"
    "basketball/mens-college-basketball"
)
ESPN_LOOKBACK_DAYS = 365
ESPN_TEAM_MATCH_THRESHOLD = 0.85
ESPN_PLAYER_MATCH_THRESHOLD = 0.8
NAME_SUFFIXES = {"jr", "sr", "ii", "iii", "iv", "v"}


UNDO_SHEET = "UndoLog"
TEAM_LIST_SHEET = "teamlist"
TEAM_STATS_PLAYERS_PER_TEAM = 5
PLAYERS_PER_TEAM_CELL = "L2"
ERROR_COLUMN = 16  # TeamData column P
ERROR_HEADER = "Error"

# Teams missing from ESPN's Division I team index, keyed by normalized name.
ESPN_TEAM_ID_OVERRIDES = {
    "queens": "2511",
    "queens university": "2511",
    "queens university of charlotte": "2511",
    "queens charlotte": "2511",
    "queens nc": "2511",
}

STATUS_COLUMN = 17  # TeamData column Q
STATUS_HEADER = "Status"
STATUS_OPTIONS = ["Out", "Questionable"]


def normalize_status(value):
    text = str(value or "").strip().casefold()

    for option in STATUS_OPTIONS:
        if text == option.casefold():
            return option

    return ""

# Standard seed order within one region of the bracket, used to work out the
# earliest round two teams from the same region could meet.
BRACKET_SEED_ORDER = [1, 16, 8, 9, 5, 12, 4, 13, 6, 11, 3, 14, 7, 10, 2, 15]
BRACKET_MEETING_ROUND = {1: "R1", 2: "R2", 3: "S16", 4: "E8"}
HINT_ROUND_ORDER = ["First Four", "R1", "R2", "S16", "E8"]
HINT_COLORS = {
    "out": "#D0D0D0",
    "questionable": "#E6DAF5",
    "conflict": "#FFF1CC",
    "stack": "#FFD8A8",
    "crowded_region": "#F8D0D0",
    "new_region": "#DFF3E4",
}


def parse_seed(value):
    match = re.search(r"\d+", str(value or ""))

    if not match:
        return None

    seed = int(match.group())
    return seed if 1 <= seed <= 16 else None


def bracket_meeting_round(seed_a, seed_b):
    seed_a = parse_seed(seed_a)
    seed_b = parse_seed(seed_b)

    if seed_a is None or seed_b is None:
        return None

    if seed_a == seed_b:
        return "First Four"

    distance = (
        BRACKET_SEED_ORDER.index(seed_a)
        ^ BRACKET_SEED_ORDER.index(seed_b)
    ).bit_length()

    return BRACKET_MEETING_ROUND.get(distance)

SETTINGS_PATH = (
    Path(os.environ.get("APPDATA", Path.home()))
    / "MMPlayerPoolCommissioner"
    / "settings.json"
)

ESPN_ROUND_MAP = [
    ("1st round", "First Round"),
    ("first round", "First Round"),
    ("2nd round", "Second Round"),
    ("second round", "Second Round"),
    ("sweet 16", "Sweet Sixteen"),
    ("sweet sixteen", "Sweet Sixteen"),
    ("elite 8", "Elite Eight"),
    ("elite eight", "Elite Eight"),
    ("final four", "Final Four"),
    ("national championship", "National Title"),
]


def load_settings():
    try:
        return json.loads(SETTINGS_PATH.read_text(encoding="utf-8"))
    except Exception:
        return {}


def save_settings(settings):
    try:
        SETTINGS_PATH.parent.mkdir(parents=True, exist_ok=True)
        SETTINGS_PATH.write_text(
            json.dumps(settings, indent=2),
            encoding="utf-8"
        )
    except Exception:
        pass


def espn_round_name(event):
    competitions = event.get("competitions") or [{}]
    headlines = [
        str(note.get("headline", ""))
        for note in competitions[0].get("notes") or []
    ]

    for headline in headlines:
        text = headline.casefold()

        if "men's basketball championship" not in text:
            continue

        for key, round_name in ESPN_ROUND_MAP:
            if key in text:
                return round_name

    return None


def espn_get_json(path, params=None, timeout=15):
    url = f"{ESPN_BASE_URL}/{path}"

    if params:
        url += "?" + urllib.parse.urlencode(params)

    request = urllib.request.Request(
        url,
        headers={"User-Agent": "Mozilla/5.0"}
    )

    with urllib.request.urlopen(request, timeout=timeout) as response:
        return json.loads(response.read().decode("utf-8"))


def normalize_name(text):
    text = unicodedata.normalize("NFKD", str(text or ""))
    text = "".join(c for c in text if not unicodedata.combining(c))
    text = re.sub(r"[^a-z0-9 ]", " ", text.casefold())
    words = []

    for word in text.split():
        if word in NAME_SUFFIXES:
            continue

        if word in ("state", "saint"):
            word = "st"

        words.append(word)

    return " ".join(words)


def name_similarity(first, second):
    first = normalize_name(first)
    second = normalize_name(second)

    if not first or not second:
        return 0.0

    if first == second:
        return 1.0

    return difflib.SequenceMatcher(None, first, second).ratio()


def espn_team_score(espn_team, team_name):
    candidates = [
        espn_team.get(key)
        for key in (
            "location",
            "shortDisplayName",
            "displayName",
            "name",
            "abbreviation",
            "nickname"
        )
        if espn_team.get(key)
    ]

    return max(
        (name_similarity(candidate, team_name) for candidate in candidates),
        default=0.0
    )


def find_espn_team_id(team_name):
    data = espn_get_json("teams", {"limit": 500, "groups": 50})
    best_id, best_score = None, 0.0

    for league in data.get("sports", [{}])[0].get("leagues", []):
        for entry in league.get("teams", []):
            team = entry.get("team", {})
            score = espn_team_score(team, team_name)

            if score > best_score:
                best_id, best_score = team.get("id"), score

    if best_score < ESPN_TEAM_MATCH_THRESHOLD:
        return None

    return best_id


def find_latest_espn_game(team_a, team_b):
    team_id = find_espn_team_id(team_a)

    if team_id is None:
        raise ValueError(f'ESPN has no team matching "{team_a}".')

    now = datetime.now()
    cutoff = (now - timedelta(days=ESPN_LOOKBACK_DAYS)).strftime("%Y-%m-%d")
    latest_allowed = datetime.utcnow().strftime("%Y-%m-%dT%H:%M")
    current_season = now.year + 1 if now.month >= 10 else now.year
    events = {}

    for season in (current_season, current_season - 1):
        for season_type in (2, 3):
            try:
                data = espn_get_json(
                    f"teams/{team_id}/schedule",
                    {"season": season, "seasontype": season_type}
                )
            except Exception:
                continue

            for event in data.get("events", []):
                event_date = event.get("date", "")

                if (
                    event.get("id")
                    and cutoff <= event_date[:10]
                    and event_date[:16] <= latest_allowed
                ):
                    events[event["id"]] = event

    return match_espn_event(list(events.values()), team_a, team_b)


def match_espn_event(events, team_a, team_b):
    best_event = None

    for event in events:
        competitions = event.get("competitions") or []

        if not competitions:
            continue

        competitors = competitions[0].get("competitors") or []

        if len(competitors) != 2:
            continue

        teams = [c.get("team", {}) for c in competitors]

        direct = min(
            espn_team_score(teams[0], team_a),
            espn_team_score(teams[1], team_b)
        )
        swapped = min(
            espn_team_score(teams[0], team_b),
            espn_team_score(teams[1], team_a)
        )

        if max(direct, swapped) < ESPN_TEAM_MATCH_THRESHOLD:
            continue

        if best_event is None or event.get("date", "") > best_event.get("date", ""):
            best_event = event

    return best_event


def read_espn_box_score(event_id):
    summary = espn_get_json("summary", {"event": event_id})
    players = {}

    for team_block in summary.get("boxscore", {}).get("players", []):
        team = team_block.get("team", {})

        for stat_group in team_block.get("statistics", []):
            labels = [
                str(label).upper()
                for label in stat_group.get("labels", [])
            ]
            keys = stat_group.get("keys", [])

            def stat_index(label, key):
                if label in labels:
                    return labels.index(label)

                if key in keys:
                    return keys.index(key)

                return None

            indexes = {
                "PTS": stat_index("PTS", "points"),
                "REB": stat_index("REB", "rebounds"),
                "AST": stat_index("AST", "assists"),
            }

            for athlete_block in stat_group.get("athletes", []):
                athlete = athlete_block.get("athlete", {})
                name = athlete.get("displayName") or athlete.get("shortName")

                if not name:
                    continue

                raw_stats = athlete_block.get("stats") or []
                values = {}

                for stat_name, index in indexes.items():
                    try:
                        values[stat_name] = int(float(raw_stats[index]))
                    except (TypeError, ValueError, IndexError):
                        values[stat_name] = 0

                players[name] = {
                    "team": team,
                    "stats": values
                }

    return players


THEME = {
    "bg": "#F4F6FA",
    "surface": "#FFFFFF",
    "primary": "#0B1F3A",
    "primary_active": "#1B3A66",
    "accent": "#F28C28",
    "accent_active": "#D9741A",
    "text": "#1F2937",
    "muted": "#6B7280",
    "border": "#D7DDE8",
    "select": "#FDE3C8",
    "font": "Segoe UI",
}


def fetch_espn_team_index():
    data = espn_get_json("teams", {"limit": 500, "groups": 50})
    teams = []

    for league in data.get("sports", [{}])[0].get("leagues", []):
        for entry in league.get("teams", []):
            teams.append(entry.get("team", {}))

    return teams


def match_espn_team_id(espn_teams, team_name):
    override = ESPN_TEAM_ID_OVERRIDES.get(normalize_name(team_name))

    if override:
        return override

    best_id, best_score = None, 0.0

    for team in espn_teams:
        score = espn_team_score(team, team_name)

        if score > best_score:
            best_id, best_score = team.get("id"), score

    if best_score < ESPN_TEAM_MATCH_THRESHOLD:
        return None

    return best_id


def fetch_espn_team_top_players(team_id, count=TEAM_STATS_PLAYERS_PER_TEAM):
    data = espn_get_json(f"teams/{team_id}/athletes/statistics")
    results = data.get("results") or []

    if not results:
        return []

    players = []

    for leader in results[0].get("leaders", []):
        stats = {}

        for category in leader.get("statistics", []):
            for stat in category.get("stats", []):
                stats.setdefault(stat.get("name"), stat.get("value"))

        def value(name):
            try:
                return float(stats.get(name) or 0)
            except (TypeError, ValueError):
                return 0.0

        games = value("gamesPlayed")
        minutes_per_game = (
            value("minutes") / games if games else value("avgMinutes")
        )

        players.append({
            "Player": leader.get("athlete", {}).get("displayName", ""),
            "GP": int(games),
            "MPG": round(minutes_per_game, 2),
            "PPG": round(value("avgPoints"), 2),
            "RPG": round(value("avgRebounds"), 2),
            "APG": round(value("avgAssists"), 2),
            "SPG": round(value("avgSteals"), 2),
            "BPG": round(value("avgBlocks"), 2),
            "TO": round(value("avgTurnovers"), 2),
            "FG%": round(value("fieldGoalPct"), 1),
            "FT%": round(value("freeThrowPct"), 1),
            "3PT%": round(value("threePointFieldGoalPct"), 1),
        })

    players.sort(key=lambda player: player["PPG"], reverse=True)
    return players[:count]


def apply_theme(root):
    font = THEME["font"]

    root.configure(bg=THEME["bg"])

    root.option_add("*Background", THEME["bg"])
    root.option_add("*Foreground", THEME["text"])
    root.option_add("*Font", (font, 10))

    root.option_add("*Button.Background", THEME["primary"])
    root.option_add("*Button.Foreground", "white")
    root.option_add("*Button.activeBackground", THEME["primary_active"])
    root.option_add("*Button.activeForeground", "white")
    root.option_add("*Button.disabledForeground", "#9AA5B5")
    root.option_add("*Button.Relief", "flat")
    root.option_add("*Button.borderWidth", 0)
    root.option_add("*Button.Cursor", "hand2")
    root.option_add("*Button.padY", 6)
    root.option_add("*Button.Font", (font, 10, "bold"))

    root.option_add("*Entry.Background", THEME["surface"])
    root.option_add("*Entry.Relief", "solid")
    root.option_add("*Entry.borderWidth", 1)
    root.option_add("*Entry.highlightThickness", 1)
    root.option_add("*Entry.highlightColor", THEME["accent"])
    root.option_add("*Entry.highlightBackground", THEME["border"])

    root.option_add("*Listbox.Background", THEME["surface"])
    root.option_add("*Listbox.Relief", "flat")
    root.option_add("*Listbox.highlightThickness", 1)
    root.option_add("*Listbox.highlightBackground", THEME["border"])
    root.option_add("*Listbox.selectBackground", THEME["select"])
    root.option_add("*Listbox.selectForeground", THEME["text"])

    root.option_add("*Labelframe.Foreground", THEME["primary"])
    root.option_add("*Labelframe.Font", (font, 10, "bold"))
    root.option_add("*Labelframe.Relief", "groove")
    root.option_add("*Labelframe.borderWidth", 1)
    root.option_add("*Labelframe.padX", 6)
    root.option_add("*Labelframe.padY", 6)

    root.option_add("*TCombobox*Listbox.Background", THEME["surface"])
    root.option_add("*TCombobox*Listbox.selectBackground", THEME["accent"])
    root.option_add("*TCombobox*Listbox.selectForeground", "white")

    style = ttk.Style(root)

    try:
        style.theme_use("clam")
    except tk.TclError:
        pass

    style.configure(
        ".",
        background=THEME["bg"],
        foreground=THEME["text"],
        font=(font, 10)
    )
    style.configure(
        "TCombobox",
        fieldbackground=THEME["surface"],
        bordercolor=THEME["border"],
        arrowcolor=THEME["primary"],
        padding=4
    )
    style.map(
        "TCombobox",
        bordercolor=[("focus", THEME["accent"])]
    )
    style.configure(
        "Vertical.TScrollbar",
        background=THEME["border"],
        troughcolor=THEME["bg"],
        bordercolor=THEME["bg"],
        arrowcolor=THEME["primary"]
    )
    style.configure(
        "Horizontal.TScrollbar",
        background=THEME["border"],
        troughcolor=THEME["bg"],
        bordercolor=THEME["bg"],
        arrowcolor=THEME["primary"]
    )
    style.configure(
        "Treeview",
        background=THEME["surface"],
        fieldbackground=THEME["surface"],
        rowheight=26,
        bordercolor=THEME["border"]
    )
    style.map(
        "Treeview",
        background=[("selected", THEME["select"])],
        foreground=[("selected", THEME["text"])]
    )
    style.configure(
        "Treeview.Heading",
        background=THEME["primary"],
        foreground="white",
        relief="flat",
        font=(font, 10, "bold"),
        padding=(4, 6)
    )
    style.map(
        "Treeview.Heading",
        background=[("active", THEME["primary_active"])]
    )

    return style

class PlayerSelectionDialog:
    NUMERIC_HEADERS = {
        "Seed",
        "Tot. FPts",
        "Proj. FPts",
        "GP",
        "MPG",
        "PPG",
        "RPG",
        "APG",
        "SPG",
        "BPG",
        "TO",
        "FG%",
        "FT%",
        "3PT%"
    }

    def __init__(
        self,
        parent,
        drafter,
        draft_position,
        player_records,
        recent_picks,
        on_deck_text="",
        pick_timer_seconds=DEFAULT_PICK_TIMER_MINUTES * 60
    ):
        self.selected_record = None
        self.undo_requested = False
        self.pick_timer_seconds = pick_timer_seconds
        self.timer_remaining = 0
        self.timer_job = None
        self.player_records = player_records
        self.filtered_records = player_records.copy()
        self.set_roster_context(drafter, recent_picks)
        self.sort_column = "Proj. FPts"
        self.sort_reverse = True

        self.window = tk.Toplevel(parent)
        self.pick_var = tk.IntVar(self.window, value=0)
        self.window.title("Select Player")
        self.window.geometry("1600x750")
        self.window.transient(parent)
        self.window.grab_set()

        self.table_style = "Blue.Treeview"
        table_style = ttk.Style(self.window)
        table_style.configure(
            self.table_style,
            background=THEME["surface"],
            fieldbackground=THEME["surface"],
            foreground=THEME["text"],
            rowheight=26
        )
        table_style.map(
            self.table_style,
            background=[("selected", THEME["select"])],
            foreground=[("selected", THEME["text"])]
        )
        table_style.configure(
            f"{self.table_style}.Heading",
            background=THEME["primary"],
            foreground="white",
            relief="flat",
            padding=(4, 6),
            font=(THEME["font"], 10, "bold")
        )
        table_style.map(
            f"{self.table_style}.Heading",
            background=[("active", THEME["primary_active"]), ("pressed", THEME["primary_active"])],
            foreground=[("active", "white"), ("pressed", "white")]
        )

        self.header_label = tk.Label(
            self.window,
            text=f"Pick {draft_position}  •  On the clock: {drafter}",
            font=(THEME["font"], 14, "bold"),
            bg=THEME["primary"],
            fg="white",
            pady=12
        )
        self.header_label.pack(fill=tk.X)

        deck_strip = tk.Frame(self.window, bg=THEME["primary_active"])
        deck_strip.pack(fill=tk.X)

        self.timer_label = tk.Label(
            deck_strip,
            text="",
            font=(THEME["font"], 12, "bold"),
            bg=THEME["primary_active"],
            fg="white",
            padx=14,
            pady=4
        )
        self.timer_label.pack(side=tk.RIGHT)

        self.on_deck_label = tk.Label(
            deck_strip,
            text=on_deck_text,
            font=(THEME["font"], 11),
            bg=THEME["primary_active"],
            fg="white",
            pady=4
        )
        self.on_deck_label.pack(side=tk.LEFT, fill=tk.X, expand=True)
        self.start_timer()

        legend_frame = tk.Frame(self.window)
        legend_frame.pack(fill=tk.X, padx=10, pady=(6, 0))

        tk.Label(
            legend_frame,
            text="Hints key:",
            font=(THEME["font"], 9, "bold")
        ).pack(side=tk.LEFT, padx=(0, 8))

        for tag, text in (
            ("new_region", "New region for this drafter"),
            ("conflict", "Bracket conflict (could meet before Final Four)"),
            ("stack", "Same-team stack"),
            ("crowded_region", "3rd+ pick in this region"),
            ("questionable", "Questionable"),
            ("out", "Out"),
        ):
            tk.Label(
                legend_frame,
                text=text,
                bg=HINT_COLORS[tag],
                fg="black",
                padx=8,
                pady=2,
                relief="solid",
                borderwidth=1,
                font=(THEME["font"], 9)
            ).pack(side=tk.LEFT, padx=(0, 8))

        self.show_hints_var = tk.BooleanVar(
            self.window,
            value=bool(load_settings().get("show_pick_hints", True))
        )
        tk.Checkbutton(
            legend_frame,
            text="Show pick hints",
            variable=self.show_hints_var,
            command=self.toggle_hints
        ).pack(side=tk.RIGHT)

        self.hide_injured_var = tk.BooleanVar(self.window, value=False)
        tk.Checkbutton(
            legend_frame,
            text="Hide Out/Questionable players",
            variable=self.hide_injured_var,
            command=self.update_suggestions
        ).pack(side=tk.RIGHT, padx=(0, 12))

        content_frame = tk.Frame(self.window)
        content_frame.pack(
            fill=tk.BOTH,
            expand=True,
            padx=10,
            pady=10
        )

        history_frame = tk.LabelFrame(
            content_frame,
            text="Draft History"
        )
        history_frame.pack(
            side=tk.BOTTOM,
            fill=tk.X,
            pady=(10, 0)
        )

        history_ids = [
            f"history_{index}"
            for index in range(len(PICK_HEADERS))
        ]

        self.history_tree = ttk.Treeview(
            history_frame,
            columns=history_ids,
            show="headings",
            height=6,
            style=self.table_style
        )

        history_widths = {
            "Draft Position": 95,
            "Draftee": 120,
            "Player Name": 135,
            "Region": 75,
            "Seed": 50,
            "Team": 90,
            "Tot. FPts": 80,
            "PPG": 60,
            "RPG": 60,
            "APG": 60
        }

        for column_id, header in zip(
            history_ids,
            PICK_HEADERS
        ):
            self.history_tree.heading(
                column_id,
                text=header
            )

            self.history_tree.column(
                column_id,
                width=history_widths[header],
                anchor=tk.CENTER
            )

        history_scrollbar = ttk.Scrollbar(
            history_frame,
            orient=tk.VERTICAL,
            command=self.history_tree.yview
        )

        self.history_tree.configure(
            yscrollcommand=history_scrollbar.set
        )

        self.history_tree.pack(
            side=tk.LEFT,
            fill=tk.BOTH,
            expand=True
        )

        history_scrollbar.pack(
            side=tk.RIGHT,
            fill=tk.Y
        )

        self.populate_history(recent_picks)

        previous_frame = tk.LabelFrame(
            content_frame,
            text="Previous Selections"
        )
        previous_frame.pack(
            side=tk.BOTTOM,
            fill=tk.X,
            pady=(10, 0)
        )

        self.previous_selection_status = tk.Label(
            previous_frame,
            anchor="w",
            fg="dim gray"
        )
        self.previous_selection_status.pack(
            fill=tk.X,
            padx=5,
            pady=(4, 2)
        )

        previous_ids = [
            f"previous_{index}"
            for index in range(len(PREVIOUS_SELECTION_HEADERS))
        ]
        self.previous_selection_tree = ttk.Treeview(
            previous_frame,
            columns=previous_ids,
            show="headings",
            height=4,
            style=self.table_style
        )

        previous_widths = {
            "Draft Position": 95,
            "Draftee": 120,
            "Player Name": 135,
            "Region": 75,
            "Seed": 50,
            "Team": 90,
            "Tot. FPts": 80,
            "PPG": 60,
            "RPG": 60,
            "APG": 60
        }

        for column_id, header in zip(
            previous_ids,
            PREVIOUS_SELECTION_HEADERS
        ):
            self.previous_selection_tree.heading(
                column_id,
                text=header
            )
            self.previous_selection_tree.column(
                column_id,
                width=previous_widths[header],
                anchor=tk.CENTER
            )

        self.previous_selection_tree.pack(
            fill=tk.X,
            padx=5,
            pady=(0, 5)
        )
        self.populate_previous_selections(
            recent_picks,
            drafter
        )

        selection_frame = tk.LabelFrame(
            content_frame,
            text="Available Players"
        )
        selection_frame.pack(
            fill=tk.BOTH,
            expand=True
        )

        tk.Label(
            selection_frame,
            text=(
                "Filter by player or statistic. Examples: "
                "PPG>=20, RPG<8, Tot. FPts>=50, Proj. FPts>=100, FG%>=45"
            )
        ).pack(pady=5)

        self.player_var = tk.StringVar()

        self.player_entry = tk.Entry(
            selection_frame,
            textvariable=self.player_var,
            width=70
        )
        self.player_entry.pack(pady=5)
        self.player_entry.focus_set()

        self.player_entry.bind(
            "<KeyRelease>",
            self.update_suggestions
        )

        self.player_entry.bind(
            "<Return>",
            lambda event: self.select_player()
        )

        table_frame = tk.Frame(selection_frame)
        table_frame.pack(
            fill=tk.BOTH,
            expand=True,
            padx=5,
            pady=5
        )

        self.column_ids = [
            f"column_{index}"
            for index in range(len(DISPLAY_HEADERS))
        ]

        self.hint_column_id = "hints"

        self.tree = ttk.Treeview(
            table_frame,
            columns=self.column_ids + [self.hint_column_id],
            displaycolumns=(
                self.column_ids[:4]
                + [self.hint_column_id]
                + self.column_ids[4:]
            ),
            show="headings",
            height=25,
            style=self.table_style
        )

        self.tree.heading(self.hint_column_id, text="Hints")
        self.tree.column(
            self.hint_column_id,
            width=230,
            minwidth=230,
            stretch=True,
            anchor=tk.W
        )

        for tag, color in HINT_COLORS.items():
            self.tree.tag_configure(tag, background=color)

        self.tree.tag_configure(
            "out",
            foreground="#B00020",
            font=(THEME["font"], 10, "bold")
        )

        vertical_scrollbar = ttk.Scrollbar(
            table_frame,
            orient=tk.VERTICAL,
            command=self.tree.yview
        )

        self.tree.configure(
            yscrollcommand=vertical_scrollbar.set
        )

        for column_id, header in zip(
            self.column_ids,
            DISPLAY_HEADERS
        ):
            self.tree.heading(
                column_id,
                text=header,
                command=lambda h=header: self.sort_by(h)
            )

            self.tree.column(
                column_id,
                width=85 if header != "Player" else 160,
                anchor=tk.CENTER
            )

        self.tree.pack(
            side=tk.LEFT,
            fill=tk.BOTH,
            expand=True
        )

        vertical_scrollbar.pack(
            side=tk.RIGHT,
            fill=tk.Y
        )

        self.tree.bind(
            "<<TreeviewSelect>>",
            self.fill_from_suggestion
        )

        self.tree.bind(
            "<Double-Button-1>",
            lambda event: self.select_player()
        )

        tk.Button(
            self.window,
            text="Select Player",
            command=self.select_player,
            width=20
        ).pack(pady=10)

        self.undo_button = tk.Button(
            self.window,
            text="\u21B6  Undo Last Pick",
            command=self.request_undo,
            bg="white",
            fg=THEME["primary"],
            activebackground=THEME["select"],
            relief="flat",
            padx=10,
            state=tk.NORMAL if recent_picks else tk.DISABLED
        )

        self.filtered_records = self.apply_sort(
            self.filtered_records
        )

        self.populate_table(self.filtered_records)
        self.autosize_window()
        self.maximize_to_work_area()
        self.window.after_idle(self.zoom_window)

        self.window.protocol(
            "WM_DELETE_WINDOW",
            self.cancel
        )

    def zoom_window(self):
        try:
            if self.window.winfo_exists():
                self.window.attributes("-fullscreen", True)
        except tk.TclError:
            pass

        try:
            if self.window.winfo_exists():
                close_button = tk.Button(
                    self.window,
                    text="\u2715  Close",
                    command=self.cancel,
                    bg="#C0392B",
                    fg="white",
                    activebackground="#A93226",
                    activeforeground="white",
                    relief="flat",
                    padx=10
                )
                close_button.place(relx=1.0, x=-10, y=8, anchor="ne")
                close_button.lift()

                self.window.update_idletasks()
                undo_offset = -(10 + close_button.winfo_reqwidth() + 8)
                self.undo_button.place(
                    relx=1.0,
                    x=undo_offset,
                    y=8,
                    anchor="ne"
                )
                self.undo_button.lift()
        except tk.TclError:
            pass

    def maximize_to_work_area(self):
        """Fit the dialog inside the Windows work area, above the taskbar."""
        try:
            windowing_system = self.window.tk.call(
                "tk",
                "windowingsystem"
            )

            if windowing_system == "win32":
                import ctypes
                from ctypes import wintypes

                work_area = wintypes.RECT()
                get_work_area = 0x0030  # SPI_GETWORKAREA
                result = ctypes.windll.user32.SystemParametersInfoW(
                    get_work_area,
                    0,
                    ctypes.byref(work_area),
                    0
                )

                if result:
                    self.window.update_idletasks()
                    hwnd = self.window.winfo_id()
                    user32 = ctypes.windll.user32

                    # Tk can report DPI-scaled dimensions. Use the actual
                    # monitor rectangle for width and the work rectangle for
                    # height so the taskbar remains uncovered.
                    monitor = user32.MonitorFromWindow(
                        hwnd,
                        2  # MONITOR_DEFAULTTONEAREST
                    )
                    monitor_info = type("MONITORINFO", (ctypes.Structure,), {
                        "_fields_": [
                            ("cbSize", wintypes.DWORD),
                            ("rcMonitor", wintypes.RECT),
                            ("rcWork", wintypes.RECT),
                            ("dwFlags", wintypes.DWORD)
                        ]
                    })()
                    monitor_info.cbSize = ctypes.sizeof(monitor_info)

                    if monitor and user32.GetMonitorInfoW(
                        monitor,
                        ctypes.byref(monitor_info)
                    ):
                        monitor_rect = monitor_info.rcMonitor
                        work_rect = monitor_info.rcWork
                        left = monitor_rect.left
                        top = work_rect.top
                        width = monitor_rect.right - monitor_rect.left
                        height = work_rect.bottom - work_rect.top
                    else:
                        left = 0
                        top = work_area.top
                        width = work_area.right - work_area.left
                        height = work_area.bottom - work_area.top

                    no_z_order = 0x0004
                    no_activate = 0x0010
                    show_window = 0x0040
                    positioned = user32.SetWindowPos(
                        hwnd,
                        0,
                        left,
                        top,
                        width,
                        height,
                        no_z_order | no_activate | show_window
                    )

                    if positioned:
                        return

                    # Geometry is a client-area size on some Tk/Windows
                    # combinations, so leave a safety margin for the frame.
                    self.window.geometry(
                        f"{max(400, width - 16)}x"
                        f"{max(300, height - 48)}+"
                        f"{left + 8}+{top + 8}"
                    )
                    return
        except (AttributeError, OSError, tk.TclError):
            pass

        try:
            self.window.state("zoomed")
        except tk.TclError:
            pass

    def populate_history(self, recent_picks):
        for item in self.history_tree.get_children():
            self.history_tree.delete(item)

        for pick in reversed(recent_picks):
            values = [
                self.format_value(
                    header,
                    pick.get(header)
                )
                for header in PICK_HEADERS
            ]

            self.history_tree.insert(
                "",
                tk.END,
                values=values
            )

    def populate_previous_selections(self, recent_picks, drafter):
        for item in self.previous_selection_tree.get_children():
            self.previous_selection_tree.delete(item)

        previous_picks = [
            pick
            for pick in recent_picks
            if str(pick.get("Draftee", "")).strip().casefold()
            == str(drafter).strip().casefold()
        ]

        if not previous_picks:
            self.previous_selection_status.config(
                text="No previous selections for this draftee."
            )
            return

        self.previous_selection_status.config(
            text=(
                f"{len(previous_picks)} previous selection(s)"
            )
        )

        for pick in reversed(previous_picks):
            values = [
                self.format_value(
                    header,
                    pick.get(header)
                )
                for header in PREVIOUS_SELECTION_HEADERS
            ]
            self.previous_selection_tree.insert(
                "",
                tk.END,
                values=values
            )

    def format_value(self, header, value):
        if value is None:
            return ""

        if isinstance(value, (int, float)):
            if header in {"FG%", "FT%", "3PT%"}:
                if 0 <= value <= 1:
                    return f"{value:.1%}"

                return f"{value:.1f}%"

            statistic_headers = {
                "Tot. FPts",
                "Proj. FPts",
                "GP",
                "MPG",
                "PPG",
                "RPG",
                "APG",
                "SPG",
                "BPG",
                "TO"
            }

            if header in statistic_headers:
                return f"{value:.1f}"

            return f"{value:g}"

        return str(value)

    def autosize_window(self):
        # Compact widths keep every player statistic visible without
        # requiring horizontal scrolling.
        compact_widths = {
            "Region": 72,
            "Seed": 65,
            "Team": 140,
            "Player": 180,
            "Tot. FPts": 88,
            "Proj. FPts": 92,
            "GP": 55,
            "MPG": 68,
            "PPG": 65,
            "RPG": 65,
            "APG": 65,
            "SPG": 65,
            "BPG": 65,
            "TO": 65,
            "FG%": 68,
            "FT%": 68,
            "3PT%": 68
        }

        screen_width = self.window.winfo_screenwidth()
        available_width = max(800, screen_width - 80)
        widths = [
            compact_widths.get(header, 65)
            for header in DISPLAY_HEADERS
        ]
        total_width = sum(widths)

        if total_width > available_width:
            scale = available_width / total_width
            widths = [
                max(44, int(width * scale))
                for width in widths
            ]

            while sum(widths) > available_width:
                widest_index = max(
                    range(len(widths)),
                    key=lambda index: widths[index]
                )

                if widths[widest_index] <= 44:
                    break

                widths[widest_index] -= 1

        for column_id, width in zip(self.column_ids, widths):
            self.tree.column(
                column_id,
                width=width,
                minwidth=width,
                stretch=True
            )

        total_width = sum(widths) + 40
        window_width = min(
            total_width,
            screen_width - 40
        )
        window_height = min(
            800,
            self.window.winfo_screenheight() - 80
        )

        self.window.geometry(
            f"{window_width}x{window_height}"
        )

    def populate_table(self, records):
        for item in self.tree.get_children():
            self.tree.delete(item)

        self.record_lookup = {}

        for record in records:
            row_id = str(record["_excel_row"])
            self.record_lookup[row_id] = record

            values = [
                self.format_value(
                    header,
                    record.get(header)
                )
                for header in DISPLAY_HEADERS
            ]

            hint_text, row_tag = self.compute_hints(record)
            values.append(hint_text)

            self.tree.insert(
                "",
                tk.END,
                iid=row_id,
                values=values,
                tags=(row_tag,) if row_tag else ()
            )

    def set_roster_context(self, drafter, recent_picks):
        drafter_key = str(drafter or "").strip().casefold()
        self.roster = [
            pick
            for pick in (recent_picks or [])
            if str(pick.get("Draftee") or "").strip().casefold()
            == drafter_key
        ]

    def toggle_hints(self):
        settings = load_settings()
        settings["show_pick_hints"] = bool(self.show_hints_var.get())
        save_settings(settings)
        self.populate_table(self.filtered_records)

    def compute_hints(self, record):
        """Return (hint text, row tag), with injury status first."""
        hint_text, row_tag = self.compute_roster_hints(record)
        status = normalize_status(record.get("Status"))

        if not status:
            return hint_text, row_tag

        status_text = status.upper() if status == "Out" else status
        combined = " | ".join(
            part for part in (status_text, hint_text) if part
        )
        return combined, status.casefold()

    def compute_roster_hints(self, record):
        """Return (hint text, row tag) for the drafter on the clock."""
        try:
            show_hints_var = getattr(self, "show_hints_var", None)

            if show_hints_var is not None and not show_hints_var.get():
                return "", None

            roster = getattr(self, "roster", [])

            if not roster:
                return "", None

            team = str(record.get("Team") or "").strip()
            region = str(record.get("Region") or "").strip()
            team_key = team.casefold()
            region_key = region.casefold()
            hints = []
            row_tag = None
            conflicts = []

            for pick in roster:
                pick_team = str(pick.get("Team") or "").strip()
                pick_region = str(pick.get("Region") or "").strip()

                if (
                    not pick_team
                    or pick_team.casefold() == team_key
                    or pick_region.casefold() != region_key
                ):
                    continue

                meeting = bracket_meeting_round(
                    record.get("Seed"),
                    pick.get("Seed")
                )

                if meeting and (meeting, pick_team) not in conflicts:
                    conflicts.append((meeting, pick_team))

            if conflicts:
                conflicts.sort(
                    key=lambda item: HINT_ROUND_ORDER.index(item[0])
                )
                hints.append(
                    "Conflict: " + ", ".join(
                        f"{meeting} vs {pick_team}"
                        for meeting, pick_team in conflicts
                    )
                )
                row_tag = "conflict"

            stack_count = sum(
                1
                for pick in roster
                if str(pick.get("Team") or "").strip().casefold()
                == team_key
            )

            if team and stack_count:
                hints.append(f"Stack #{stack_count + 1}")
                row_tag = "stack"

            region_count = sum(
                1
                for pick in roster
                if str(pick.get("Region") or "").strip().casefold()
                == region_key
            )

            if region and region_count >= 2:
                hints.append(f"Region pick #{region_count + 1}")
                row_tag = "crowded_region"

            roster_regions = {
                str(pick.get("Region") or "").strip().casefold()
                for pick in roster
            }

            if region and region_key not in roster_regions:
                hints.append("New region")
                row_tag = row_tag or "new_region"

            return " | ".join(hints), row_tag
        except Exception:
            return "", None

    def normalize_header(self, header):
        return re.sub(
            r"\s+",
            " ",
            header.strip().lower()
        )

    def numeric_value(self, value):
        try:
            return float(value)
        except (TypeError, ValueError):
            return None

    def matches_filter(self, record, filter_text):
        filter_text = filter_text.strip()

        if not filter_text:
            return True

        comparison_match = re.match(
            r"^\s*(.+?)\s*(>=|<=|=|>|<)\s*(-?\d+(?:\.\d+)?)\s*$",
            filter_text
        )

        if comparison_match:
            raw_header, operator, target_text = (
                comparison_match.groups()
            )

            header_lookup = {
                self.normalize_header(header): header
                for header in DISPLAY_HEADERS
            }

            header = header_lookup.get(
                self.normalize_header(raw_header)
            )

            if header is None:
                return False

            actual_value = self.numeric_value(
                record.get(header)
            )

            if actual_value is None:
                return False

            target_value = float(target_text)

            if header in {"FG%", "FT%", "3PT%"}:
                if target_value > 1:
                    target_value /= 100

            if operator == ">=":
                return actual_value >= target_value

            if operator == "<=":
                return actual_value <= target_value

            if operator == ">":
                return actual_value > target_value

            if operator == "<":
                return actual_value < target_value

            return actual_value == target_value

        searchable_text = " ".join(
            self.format_value(
                header,
                record.get(header)
            ).lower()
            for header in DISPLAY_HEADERS
        )

        return all(
            term in searchable_text
            for term in filter_text.lower().split()
        )

    def update_suggestions(self, event=None):
        filter_text = self.player_var.get()
        hide_injured_var = getattr(self, "hide_injured_var", None)
        hide_injured = bool(hide_injured_var and hide_injured_var.get())

        self.filtered_records = [
            record
            for record in self.player_records
            if self.matches_filter(record, filter_text)
            and not (
                hide_injured and normalize_status(record.get("Status"))
            )
        ]

        self.filtered_records = self.apply_sort(
            self.filtered_records
        )

        self.populate_table(self.filtered_records)

    def get_sort_key(self, record, header):
        value = record.get(header)

        if header in self.NUMERIC_HEADERS:
            numeric_value = self.numeric_value(value)

            if numeric_value is None:
                return (1, 0)

            return (0, numeric_value)

        return (
            0,
            str(value).lower()
            if value is not None
            else ""
        )

    def apply_sort(self, records):
        if self.sort_column is None:
            return records

        return sorted(
            records,
            key=lambda record: self.get_sort_key(
                record,
                self.sort_column
            ),
            reverse=self.sort_reverse
        )

    def sort_by(self, header):
        if self.sort_column == header:
            self.sort_reverse = not self.sort_reverse
        else:
            self.sort_column = header
            self.sort_reverse = False

        self.filtered_records = self.apply_sort(
            self.filtered_records
        )

        self.populate_table(self.filtered_records)

    def fill_from_suggestion(self, event=None):
        selection = self.tree.selection()

        if selection:
            record = self.record_lookup[selection[0]]

            self.player_var.set(
                record["Player"]
            )

            self.player_entry.icursor(tk.END)

    def confirm_status(self, record):
        status = normalize_status(record.get("Status"))

        if status != "Out":
            return True

        return messagebox.askyesno(
            "Player Marked Out",
            f"{record.get('Player')} ({record.get('Team')}) is marked OUT."
            "\n\nDraft anyway?",
            icon="warning",
            parent=self.window
        )

    def choose_record(self, record):
        if not self.confirm_status(record):
            return

        self.selected_record = record
        self.pick_var.set(self.pick_var.get() + 1)

    def select_player(self):
        selection = self.tree.selection()

        if selection:
            self.choose_record(self.record_lookup[selection[0]])
            return

        typed_name = self.player_var.get().strip().lower()

        matches = [
            record
            for record in self.player_records
            if str(
                record.get("Player", "")
            ).lower() == typed_name
        ]

        if len(matches) == 1:
            self.choose_record(matches[0])
            return

        if len(self.filtered_records) == 1:
            self.choose_record(self.filtered_records[0])
            return

        messagebox.showwarning(
            "Selection Required",
            "Select a player row or enter an exact player name.",
            parent=self.window
        )

    def wait_for_pick(self):
        self.selected_record = None
        self.undo_requested = False

        if not self.window.winfo_exists():
            return None

        self.window.wait_variable(self.pick_var)
        return self.selected_record

    def request_undo(self):
        self.selected_record = None
        self.undo_requested = True
        self.pick_var.set(self.pick_var.get() + 1)

    def start_timer(self):
        self.stop_timer()

        if not self.pick_timer_seconds or self.pick_timer_seconds <= 0:
            self.timer_label.config(text="")
            return

        self.timer_remaining = int(self.pick_timer_seconds)
        self.tick_timer()

    def stop_timer(self):
        if self.timer_job is not None:
            try:
                self.window.after_cancel(self.timer_job)
            except tk.TclError:
                pass

            self.timer_job = None

    def tick_timer(self):
        self.timer_job = None

        try:
            if not self.window.winfo_exists():
                return
        except tk.TclError:
            return

        remaining = max(0, self.timer_remaining)
        minutes, seconds = divmod(remaining, 60)

        if remaining == 0:
            text, color = "\u23F1  Time's up", "#FF8A80"
        else:
            text = f"\u23F1  {minutes}:{seconds:02d}"
            color = "#FFE082" if remaining <= 30 else "white"

        self.timer_label.config(text=text, fg=color)

        if remaining > 0:
            self.timer_remaining -= 1
            self.timer_job = self.window.after(1000, self.tick_timer)

    def next_pick(
        self,
        drafter,
        draft_position,
        recent_picks,
        on_deck_text=""
    ):
        self.set_roster_context(drafter, recent_picks)
        self.on_deck_label.config(text=on_deck_text)
        self.start_timer()
        self.undo_button.config(
            state=tk.NORMAL if recent_picks else tk.DISABLED
        )
        self.header_label.config(
            text=f"Pick {draft_position}  •  On the clock: {drafter}"
        )
        self.populate_history(recent_picks)
        self.populate_previous_selections(recent_picks, drafter)
        self.player_var.set("")
        self.update_suggestions()
        self.tree.yview_moveto(0)
        self.player_entry.focus_set()

    def close(self):
        self.stop_timer()

        if self.window.winfo_exists():
            self.window.destroy()

    def cancel(self):
        self.stop_timer()
        self.selected_record = None
        self.pick_var.set(self.pick_var.get() + 1)
        self.window.destroy()

class DraftApp:
    def __init__(self, root):
        self.root = root
        self.root.title("March Madness Player Pool — Commissioner")
        self.root.geometry("620x620")
        self.root.minsize(520, 520)
        apply_theme(self.root)

        self.file_path = None
        self.workbook = None
        self.drafters = []
        self.players = []
        self.teams = []
        self.available_teams = []
        self.game_results_window = None
        self.last_game_result = {
            "Round": "",
            "Winner": "",
            "Loser": ""
        }

        header = tk.Frame(root, bg=THEME["primary"])
        header.pack(fill=tk.X)

        tk.Frame(header, bg=THEME["accent"], height=4).pack(fill=tk.X, side=tk.BOTTOM)

        tk.Label(
            header,
            text="🏀  March Madness Player Pool",
            font=(THEME["font"], 18, "bold"),
            bg=THEME["primary"],
            fg="white"
        ).pack(anchor="w", padx=20, pady=(16, 0))

        tk.Label(
            header,
            text="Commissioner Dashboard",
            font=(THEME["font"], 10),
            bg=THEME["primary"],
            fg="#B8C4D6"
        ).pack(anchor="w", padx=22, pady=(0, 14))

        body = tk.Frame(root)
        body.pack(fill=tk.BOTH, expand=True, padx=20, pady=16)

        file_card = tk.LabelFrame(body, text="Workbook")
        file_card.pack(fill=tk.X)

        tk.Button(
            file_card,
            text="📂  Select Excel File",
            command=self.select_file,
            padx=12
        ).pack(side=tk.LEFT, padx=6, pady=6)

        self.file_label = tk.Label(
            file_card,
            text="No file selected",
            fg=THEME["muted"],
            anchor="w"
        )
        self.file_label.pack(side=tk.LEFT, fill=tk.X, expand=True, padx=8)

        actions = tk.LabelFrame(body, text="Actions")
        actions.pack(fill=tk.X, pady=(12, 0))
        actions.grid_columnconfigure((0, 1), weight=1, uniform="actions")

        self.team_stats_button = tk.Button(
            actions,
            text="\U0001F4CA  Pull Team Stats from ESPN",
            command=self.pull_team_stats,
            font=(THEME["font"], 11, "bold"),
            pady=10,
            state=tk.DISABLED
        )
        self.team_stats_button.grid(
            row=0, column=0, columnspan=2, sticky="ew", padx=6, pady=6
        )

        self.begin_button = tk.Button(
            actions,
            text="📝  Begin / Resume Draft",
            command=self.begin_draft,
            bg=THEME["accent"],
            activebackground=THEME["accent_active"],
            font=(THEME["font"], 11, "bold"),
            pady=10,
            state=tk.DISABLED
        )
        self.begin_button.grid(row=1, column=0, sticky="ew", padx=6, pady=6)

        self.game_results_button = tk.Button(
            actions,
            text="🏆  Input Game Results",
            command=self.open_game_results,
            font=(THEME["font"], 11, "bold"),
            pady=10,
            state=tk.DISABLED
        )
        self.game_results_button.grid(row=1, column=1, sticky="ew", padx=6, pady=6)

        self.status_button = tk.Button(
            actions,
            text="\U0001FA79  Update Player Statuses",
            command=self.open_player_statuses,
            font=(THEME["font"], 11, "bold"),
            pady=10,
            state=tk.DISABLED
        )
        self.status_button.grid(
            row=2, column=0, columnspan=2, sticky="ew", padx=6, pady=6
        )

        drafters_card = tk.LabelFrame(body, text="Draft Order")
        drafters_card.pack(fill=tk.BOTH, expand=True, pady=(12, 0))

        self.drafter_list = tk.Listbox(
            drafters_card,
            height=12,
            activestyle="none",
            font=(THEME["font"], 11)
        )
        self.drafter_list.pack(fill=tk.BOTH, expand=True, padx=4, pady=4)

        status_bar = tk.Frame(root, bg=THEME["border"])
        status_bar.pack(fill=tk.X, side=tk.BOTTOM)

        self.status_label = tk.Label(
            status_bar,
            text="Select an Excel file to begin.",
            bg=THEME["border"],
            fg=THEME["text"],
            anchor="w",
            font=(THEME["font"], 9)
        )
        self.status_label.pack(fill=tk.X, padx=12, pady=4)

        self.root.after(100, self.reopen_last_workbook)

    def select_file(self):
        file_path = filedialog.askopenfilename(
            title="Select Excel File",
            filetypes=[
                ("Excel files", "*.xlsx *.xlsm"),
                ("All files", "*.*")
            ]
        )

        if not file_path:
            return

        self.load_workbook_file(file_path)

    def reopen_last_workbook(self):
        last_path = load_settings().get("last_workbook")

        if not last_path or not Path(last_path).exists():
            return

        if self.load_workbook_file(last_path, show_errors=False):
            self.status_label.config(
                text=f"Reopened last workbook. {self.status_label.cget('text')}"
            )

    def load_workbook_file(self, file_path, show_errors=True):
        try:
            keep_vba = Path(file_path).suffix.lower() == ".xlsm"

            self.workbook = load_workbook(
                file_path,
                keep_vba=keep_vba
            )

            self.file_path = file_path
            self.team_stats_button.config(
                state=(
                    tk.NORMAL
                    if self.find_team_list_sheet() is not None
                    else tk.DISABLED
                )
            )

            if DRAFT_ORDER_SHEET not in self.workbook.sheetnames:
                raise ValueError(
                    f'Missing sheet: "{DRAFT_ORDER_SHEET}"'
                )

            if TEAM_DATA_SHEET not in self.workbook.sheetnames:
                raise ValueError(
                    f'Missing sheet: "{TEAM_DATA_SHEET}"'
                )

            if CONFIG_SHEET not in self.workbook.sheetnames:
                raise ValueError(
                    f'Missing sheet: "{CONFIG_SHEET}" with FPts values in A2:C2'
                )

            if DRAFT_SHEET not in self.workbook.sheetnames:
                self.workbook.create_sheet(DRAFT_SHEET)

            self.file_path = file_path
            self.last_game_result = self.read_last_game_result()

            draft_order_sheet = self.workbook[DRAFT_ORDER_SHEET]
            team_data_sheet = self.workbook[TEAM_DATA_SHEET]
            config_sheet = self.workbook[CONFIG_SHEET]
            scoring_weights = self.read_fantasy_point_weights(
                config_sheet
            )
            seed_average_games = self.read_seed_average_games(
                config_sheet
            )

            self.drafters = self.read_column_values(
                draft_order_sheet,
                DRAFTER_START_ROW,
                DRAFTER_COLUMN
            )

            self.players = self.read_team_players(
                team_data_sheet,
                scoring_weights,
                seed_average_games
            )

            self.teams = self.read_unique_team_names(
                team_data_sheet
            )

            if not self.drafters:
                raise ValueError(
                    "No drafters found in DraftOrder."
                )

            if not self.players:
                raise ValueError(
                    "No players found in TeamData."
                )

            if not self.teams:
                raise ValueError(
                    "No teams found in TeamData."
                )

            self.available_teams = self.get_available_teams()
            self.status_button.config(state=tk.NORMAL)

            self.game_results_button.config(
                state=(
                    tk.NORMAL
                    if len(self.available_teams) >= 2
                    else tk.DISABLED
                )
            )

            if self.is_draft_complete():
                self.begin_button.config(
                    state=tk.DISABLED
                )

                draft_status = "Draft already completed."
            else:
                self.begin_button.config(
                    state=tk.NORMAL
                )

                draft_status = (
                    f"{len(self.drafters)} drafters loaded, "
                    f"{len(self.players)} players loaded."
                )

            self.file_label.config(
                text=Path(file_path).name,
                fg=THEME["text"]
            )

            self.status_label.config(
                text=draft_status
            )

            self.drafter_list.delete(0, tk.END)

            for position, drafter in enumerate(
                self.drafters,
                start=1
            ):
                self.drafter_list.insert(
                    tk.END,
                    f"{position}. {drafter}"
                )

            settings = load_settings()
            settings["last_workbook"] = str(Path(file_path).resolve())
            save_settings(settings)
            return True

        except Exception as error:
            self.begin_button.config(
                state=tk.DISABLED
            )

            self.game_results_button.config(
                state=tk.DISABLED
            )

            if show_errors:
                messagebox.showerror(
                    "Error Reading Workbook",
                    str(error)
                )
            else:
                self.status_label.config(
                    text=f"Could not reopen last workbook: {error}"
                )

            return False

    @staticmethod
    def read_column_values(worksheet, start_row, column):
        values = []

        for row in range(
            start_row,
            worksheet.max_row + 1
        ):
            value = worksheet.cell(
                row=row,
                column=column
            ).value

            if value is not None and str(value).strip():
                values.append(str(value).strip())

        return values

    @staticmethod
    def numeric_value(value):
        try:
            return float(value)
        except (TypeError, ValueError):
            return 0.0

    @staticmethod
    def seed_number(seed):
        if seed is None:
            return None

        numeric_seed = DraftApp.numeric_value(seed)
        if numeric_seed:
            return int(numeric_seed)

        match = re.search(r"\d+", str(seed))
        return int(match.group()) if match else None

    @staticmethod
    def projected_fantasy_points(total_fpts, seed, seed_average_games):
        total_fpts = DraftApp.numeric_value(total_fpts)
        average_games = seed_average_games.get(
            DraftApp.seed_number(seed)
        )

        if total_fpts is None or average_games is None:
            return None

        return total_fpts * average_games

    @staticmethod
    def read_pick_timer_seconds(worksheet):
        """Read the pick timer (minutes) from Config!K2, adding a default."""
        value = worksheet[PICK_TIMER_VALUE_CELL].value

        if value in (None, ""):
            if worksheet[PICK_TIMER_HEADER_CELL].value in (None, ""):
                worksheet[PICK_TIMER_HEADER_CELL] = PICK_TIMER_HEADER

            worksheet[PICK_TIMER_VALUE_CELL] = DEFAULT_PICK_TIMER_MINUTES
            value = DEFAULT_PICK_TIMER_MINUTES

        try:
            minutes = float(value)
        except (TypeError, ValueError):
            minutes = DEFAULT_PICK_TIMER_MINUTES

        return max(0, int(round(minutes * 60)))

    @staticmethod
    def read_fantasy_point_weights(worksheet):
        config_cells = ("A2", "B2", "C2")
        stat_names = ("PPG", "RPG", "APG")
        weights = []

        for cell_ref, stat_name in zip(config_cells, stat_names):
            value = worksheet[cell_ref].value

            try:
                weight = float(value)
            except (TypeError, ValueError):
                raise ValueError(
                    f'Config!{cell_ref} must contain a numeric FPts value '
                    f'for {stat_name}.'
                )

            weights.append(weight)

        return tuple(weights)

    @staticmethod
    def read_seed_average_games(worksheet):
        seed_average_games = {}

        for row in range(SEED_GAMES_FIRST_ROW, SEED_GAMES_LAST_ROW + 1):
            seed_ref = f"{SEED_GAMES_SEED_COLUMN}{row}"
            games_ref = f"{SEED_GAMES_AVERAGE_COLUMN}{row}"
            seed = DraftApp.seed_number(worksheet[seed_ref].value)

            if seed is None:
                raise ValueError(
                    f"Config!{seed_ref} must contain a seed number."
                )

            try:
                average_games = float(worksheet[games_ref].value)
            except (TypeError, ValueError):
                raise ValueError(
                    f"Config!{games_ref} must contain a numeric average "
                    f"games value for seed {seed}."
                )

            seed_average_games[seed] = average_games

        return seed_average_games

    @staticmethod
    def read_team_players(worksheet, scoring_weights, seed_average_games):
        header_map = {}

        for column in range(
            1,
            worksheet.max_column + 1
        ):
            value = worksheet.cell(
                row=1,
                column=column
            ).value

            if value is not None:
                header_map[str(value).strip()] = column

        missing_headers = [
            header
            for header in TEAM_HEADERS
            if header not in header_map
        ]

        if missing_headers:
            raise ValueError(
                "Missing TeamData headers: "
                + ", ".join(missing_headers)
            )

        player_records = []

        for row in range(
            PLAYER_START_ROW,
            worksheet.max_row + 1
        ):
            player_value = worksheet.cell(
                row=row,
                column=header_map["Player"]
            ).value

            if player_value is None:
                continue

            record = {
                header: worksheet.cell(
                    row=row,
                    column=header_map[header]
                ).value
                for header in TEAM_HEADERS
            }

            ppg_weight, rpg_weight, apg_weight = scoring_weights
            record["Tot. FPts"] = (
                DraftApp.numeric_value(record.get("PPG"))
                * ppg_weight
                + DraftApp.numeric_value(record.get("RPG"))
                * rpg_weight
                + DraftApp.numeric_value(record.get("APG"))
                * apg_weight
            )
            record["Proj. FPts"] = DraftApp.projected_fantasy_points(
                record["Tot. FPts"],
                record.get("Seed"),
                seed_average_games
            )

            record["Status"] = normalize_status(
                worksheet.cell(row=row, column=STATUS_COLUMN).value
            )
            record["_excel_row"] = row
            player_records.append(record)

        return player_records

    @staticmethod
    def read_unique_team_names(worksheet):
        header_map = {}

        for column in range(
            1,
            worksheet.max_column + 1
        ):
            value = worksheet.cell(
                row=1,
                column=column
            ).value

            if value is not None:
                header_map[str(value).strip()] = column

        if "Team" not in header_map:
            raise ValueError(
                'Missing "Team" header in TeamData.'
            )

        team_column = header_map["Team"]
        teams = []
        seen = set()

        for row in range(
            PLAYER_START_ROW,
            worksheet.max_row + 1
        ):
            value = worksheet.cell(
                row=row,
                column=team_column
            ).value

            if value is None:
                continue

            team_name = str(value).strip()
            team_key = team_name.casefold()

            if team_name and team_key not in seen:
                teams.append(team_name)
                seen.add(team_key)

        return teams

    def read_teams_out(self):
        excluded_teams = set()

        if TEAMS_OUT_SHEET not in self.workbook.sheetnames:
            return excluded_teams

        teams_out_sheet = self.workbook[TEAMS_OUT_SHEET]

        for row in range(
            2,
            teams_out_sheet.max_row + 1
        ):
            value = teams_out_sheet.cell(
                row=row,
                column=1
            ).value

            if value is not None and str(value).strip():
                excluded_teams.add(
                    str(value).strip().casefold()
                )

        return excluded_teams

    def read_last_game_result(self):
        result = {
            "Round": "",
            "Winner": "",
            "Loser": ""
        }

        if GAME_RESULTS_SHEET not in self.workbook.sheetnames:
            return result

        worksheet = self.workbook[GAME_RESULTS_SHEET]
        headers = {}

        for column in range(1, worksheet.max_column + 1):
            value = worksheet.cell(row=1, column=column).value

            if value is not None:
                headers[str(value).strip()] = column

        required_headers = ["Round", "Winner", "Loser"]
        if any(header not in headers for header in required_headers):
            return result

        for row in range(worksheet.max_row, 1, -1):
            values = {
                header: worksheet.cell(
                    row=row,
                    column=headers[header]
                ).value
                for header in required_headers
            }

            if any(value not in (None, "") for value in values.values()):
                return {
                    header: "" if values[header] is None else str(
                        values[header]
                    ).strip()
                    for header in required_headers
                }

        return result

    def save_last_game_result(self, selected_round, winner, loser):
        if GAME_RESULTS_SHEET not in self.workbook.sheetnames:
            worksheet = self.workbook.create_sheet(GAME_RESULTS_SHEET)
        else:
            worksheet = self.workbook[GAME_RESULTS_SHEET]

        worksheet["A1"] = "Round"
        worksheet["B1"] = "Winner"
        worksheet["C1"] = "Loser"
        worksheet["D1"] = "Statistics Written"
        worksheet["A2"] = selected_round
        worksheet["B2"] = winner
        worksheet["C2"] = loser
        worksheet["D2"] = datetime.now()
        worksheet["D2"].number_format = "m/d/yyyy h:mm:ss AM/PM"

        if worksheet.max_row > 2:
            worksheet.delete_rows(3, worksheet.max_row - 2)

        self.last_game_result = {
            "Round": selected_round,
            "Winner": winner,
            "Loser": loser
        }

    def last_game_result_text(self):
        winner = self.last_game_result.get("Winner", "")
        loser = self.last_game_result.get("Loser", "")
        selected_round = self.last_game_result.get("Round", "")

        if not winner or not loser:
            return "Most Recent Saved Result: None recorded"

        return (
            "Most Recent Saved Result: "
            f"{selected_round} | Winner: {winner} | Loser: {loser}"
        )

    def get_available_teams(self):
        excluded_teams = self.read_teams_out()

        return [
            team
            for team in self.teams
            if team.casefold() not in excluded_teams
        ]

    def get_autocomplete_matches(self, text, options):
        text = text.strip().casefold()
        options = list(options)

        if not text:
            return options

        starts_with = [
            option
            for option in options
            if option.casefold().startswith(text)
        ]

        contains = [
            option
            for option in options
            if text in option.casefold()
            and option not in starts_with
        ]

        return starts_with + contains

    def resolve_autocomplete_value(self, text, options):
        text = text.strip()

        if not text:
            return None

        for option in options:
            if option.casefold() == text.casefold():
                return option

        matches = self.get_autocomplete_matches(
            text,
            options
        )

        if len(matches) == 1:
            return matches[0]

        return None

    def attach_autocomplete(
        self,
        combo,
        variable,
        options_getter
    ):
        def filter_options(event=None):
            if event and event.keysym in {
                "Up",
                "Down",
                "Left",
                "Right",
                "Return",
                "Escape",
                "Tab"
            }:
                return

            combo["values"] = self.get_autocomplete_matches(
                variable.get(),
                options_getter()
            )

        def normalize_value(event=None):
            resolved_value = self.resolve_autocomplete_value(
                variable.get(),
                options_getter()
            )

            if resolved_value is not None:
                variable.set(resolved_value)

            combo["values"] = list(options_getter())

        combo.bind(
            "<KeyRelease>",
            filter_options
        )

        combo.bind(
            "<FocusOut>",
            normalize_value
        )

        combo["values"] = list(options_getter())

    def open_game_results(self):
        if len(self.available_teams) < 2:
            messagebox.showwarning(
                "No Available Teams",
                "At least two available teams are required."
            )
            return

        if (
            self.game_results_window is not None
            and self.game_results_window.winfo_exists()
        ):
            self.game_results_window.lift()
            return

        window = tk.Toplevel(self.root)
        self.game_results_window = window

        window.title("Input Game Results")
        window.geometry("900x780")
        window.minsize(760, 600)
        window.configure(bg=THEME["bg"])
        window.transient(self.root)
        window.grab_set()

        round_var = tk.StringVar()
        winner_var = tk.StringVar()
        loser_var = tk.StringVar()

        header = tk.Frame(window, bg=THEME["primary"])
        header.pack(fill=tk.X)
        tk.Frame(header, bg=THEME["accent"], height=4).pack(fill=tk.X, side=tk.BOTTOM)

        tk.Label(
            header,
            text="🏆  Input Game Results",
            font=(THEME["font"], 16, "bold"),
            bg=THEME["primary"],
            fg="white"
        ).pack(anchor="w", padx=20, pady=(12, 0))

        recent_result_label = tk.Label(
            header,
            text=self.last_game_result_text(),
            bg=THEME["primary"],
            fg="#B8C4D6"
        )
        recent_result_label.pack(anchor="w", padx=22, pady=(0, 10))

        matchup_card = tk.LabelFrame(window, text="Matchup")
        matchup_card.pack(fill=tk.X, padx=16, pady=(14, 0))
        matchup_card.grid_columnconfigure((0, 1, 2), weight=1, uniform="matchup")

        for column, label_text in enumerate(
            ("Game Round", "Winner", "Loser")
        ):
            tk.Label(
                matchup_card,
                text=label_text,
                fg=THEME["muted"],
                font=(THEME["font"], 9, "bold")
            ).grid(row=0, column=column, sticky="w", padx=8, pady=(2, 2))

        round_box = ttk.Combobox(
            matchup_card,
            textvariable=round_var,
            values=ROUND_OPTIONS,
            state="normal",
            width=24
        )
        round_box.grid(row=1, column=0, sticky="ew", padx=8, pady=(0, 6))

        winner_box = ttk.Combobox(
            matchup_card,
            textvariable=winner_var,
            values=self.available_teams,
            state="normal",
            width=24
        )
        winner_box.grid(row=1, column=1, sticky="ew", padx=8, pady=(0, 6))

        loser_box = ttk.Combobox(
            matchup_card,
            textvariable=loser_var,
            values=self.available_teams,
            state="normal",
            width=24
        )
        loser_box.grid(row=1, column=2, sticky="ew", padx=8, pady=(0, 6))

        espn_row = tk.Frame(window)
        espn_row.pack(fill=tk.X, padx=16, pady=(10, 0))

        espn_button = tk.Button(
            espn_row,
            text="⚡  Fetch Stats from ESPN",
            bg=THEME["accent"],
            activebackground=THEME["accent_active"],
            padx=14
        )
        espn_button.pack(side=tk.LEFT)

        espn_status_label = tk.Label(
            espn_row,
            text="",
            fg=THEME["muted"],
            wraplength=620,
            justify=tk.LEFT,
            anchor="w"
        )
        espn_status_label.pack(side=tk.LEFT, fill=tk.X, expand=True, padx=12)

        players_frame = tk.LabelFrame(
            window,
            text="Drafted Players and Statistics"
        )
        players_frame.pack(
            fill=tk.BOTH,
            expand=True,
            padx=16,
            pady=12
        )

        stats_canvas = tk.Canvas(
            players_frame,
            highlightthickness=0
        )

        stats_scrollbar = ttk.Scrollbar(
            players_frame,
            orient=tk.VERTICAL,
            command=stats_canvas.yview
        )

        stats_inner = tk.Frame(stats_canvas)

        stats_canvas_window = stats_canvas.create_window(
            (0, 0),
            window=stats_inner,
            anchor="nw"
        )

        stats_canvas.configure(
            yscrollcommand=stats_scrollbar.set
        )

        stats_inner.bind(
            "<Configure>",
            lambda event: stats_canvas.configure(
                scrollregion=stats_canvas.bbox("all")
            )
        )

        stats_canvas.bind(
            "<Configure>",
            lambda event: stats_canvas.itemconfigure(
                stats_canvas_window,
                width=event.width
            )
        )

        stats_canvas.pack(
            side=tk.LEFT,
            fill=tk.BOTH,
            expand=True
        )

        stats_scrollbar.pack(
            side=tk.RIGHT,
            fill=tk.Y
        )

        for column, header in enumerate(
            ["Player Name", "Team Owner", "Team", "PTS", "REB", "AST"]
        ):
            table_columns = [
                ("Player Name", 180),
                ("Team Owner", 140),
                ("Team", 120),
                ("PTS", 95),
                ("REB", 95),
                ("AST", 95),
            ]

            for column, (header, width) in enumerate(table_columns):
                stats_inner.grid_columnconfigure(
                    column,
                    minsize=width,
                    weight=1,
                    uniform="stats_table"
                )

                tk.Label(
                    stats_inner,
                    text=header,
                    font=(THEME["font"], 10, "bold"),
                    fg=THEME["primary"],
                    anchor=tk.CENTER,
                    justify=tk.CENTER
                ).grid(
                    row=0,
                    column=column,
                    padx=8,
                    pady=5,
                    sticky="nsew"
                )



        stats_entries = {}

        def refresh_player_list(event=None):
            self.update_team_options(
                winner_box,
                loser_box,
                winner_var,
                loser_var
            )

            self.update_game_result_players(
                winner_var,
                loser_var,
                stats_inner,
                stats_entries
            )

        def validate_and_refresh(event=None):
            self.validate_team_selections(
                winner_box,
                loser_box,
                winner_var,
                loser_var
            )

            refresh_player_list()

        espn_button.config(
            command=lambda: self.fetch_espn_stats(
                window,
                winner_var,
                loser_var,
                stats_entries,
                validate_and_refresh,
                espn_status_label,
                round_var
            )
        )

        self.attach_autocomplete(
            round_box,
            round_var,
            lambda: ROUND_OPTIONS
        )

        self.attach_autocomplete(
            winner_box,
            winner_var,
            lambda: self.available_teams
        )

        self.attach_autocomplete(
            loser_box,
            loser_var,
            lambda: self.available_teams
        )

        winner_box.bind(
            "<KeyRelease>",
            refresh_player_list,
            add="+"
        )

        loser_box.bind(
            "<KeyRelease>",
            refresh_player_list,
            add="+"
        )

        winner_box.bind(
            "<<ComboboxSelected>>",
            refresh_player_list
        )

        loser_box.bind(
            "<<ComboboxSelected>>",
            refresh_player_list
        )

        winner_box.bind(
            "<FocusOut>",
            validate_and_refresh,
            add="+"
        )

        loser_box.bind(
            "<FocusOut>",
            validate_and_refresh,
            add="+"
        )

        button_frame = tk.Frame(window)
        button_frame.pack(pady=10)

        tk.Button(
            button_frame,
            text="Save Result",
            command=lambda: self.save_game_result(
                window,
                round_var,
                winner_var,
                loser_var,
                stats_entries,
                recent_result_label=recent_result_label,
                close_after_save=True
            ),
            bg=THEME["accent"],
            activebackground=THEME["accent_active"],
            width=18
        ).pack(side=tk.LEFT, padx=5)

        tk.Button(
            button_frame,
            text="Save & Enter Another",
            command=lambda: self.save_game_result(
                window,
                round_var,
                winner_var,
                loser_var,
                stats_entries,
                winner_box,
                loser_box,
                stats_inner,
                recent_result_label=recent_result_label,
                close_after_save=False
            ),
            width=22
        ).pack(side=tk.LEFT, padx=5)

        tk.Button(
            button_frame,
            text="↶  Undo Last Result",
            command=lambda: self.undo_last_game_result(
                window,
                winner_box,
                loser_box,
                recent_result_label,
                refresh_player_list
            ),
            width=20
        ).pack(side=tk.LEFT, padx=5)

        tk.Button(
            button_frame,
            text="Close",
            command=window.destroy,
            width=15
        ).pack(side=tk.LEFT, padx=5)

        window.protocol(
            "WM_DELETE_WINDOW",
            window.destroy
        )

    def update_team_options(
        self,
        winner_box,
        loser_box,
        winner_var,
        loser_var
    ):
        winner_text = winner_var.get().strip().casefold()
        loser_text = loser_var.get().strip().casefold()

        winner_options = [
            team
            for team in self.available_teams
            if winner_text in team.casefold()
            and team.casefold() != loser_text
        ]

        loser_options = [
            team
            for team in self.available_teams
            if loser_text in team.casefold()
            and team.casefold() != winner_text
        ]

        winner_box["values"] = winner_options
        loser_box["values"] = loser_options

    def validate_team_selections(
        self,
        winner_box,
        loser_box,
        winner_var,
        loser_var
    ):
        winner = self.resolve_autocomplete_value(
            winner_var.get(),
            self.available_teams
        )

        loser = self.resolve_autocomplete_value(
            loser_var.get(),
            self.available_teams
        )

        if winner is not None:
            winner_var.set(winner)

        if loser is not None:
            loser_var.set(loser)

        if (
            winner
            and loser
            and winner.casefold() == loser.casefold()
        ):
            messagebox.showwarning(
                "Duplicate Team",
                "Winner and Loser must be different teams."
            )

            loser_var.set("")

        self.update_team_options(
            winner_box,
            loser_box,
            winner_var,
            loser_var
        )

    def update_game_result_players(
        self,
        winner_var,
        loser_var,
        stats_inner,
        stats_entries
    ):
        for widget in stats_inner.winfo_children():
            grid_info = widget.grid_info()

            if int(grid_info.get("row", 0)) > 0:
                widget.destroy()

        stats_entries.clear()

        winner = self.resolve_autocomplete_value(
            winner_var.get(),
            self.teams
        )

        loser = self.resolve_autocomplete_value(
            loser_var.get(),
            self.teams
        )

        if winner is None or loser is None:
            return

        if winner.casefold() == loser.casefold():
            return

        matching_players = self.get_drafted_players_for_teams({
            winner.casefold(),
            loser.casefold()
        })

        matching_players = sorted(
            matching_players,
            key=lambda player: (
                str(player["Team"]).strip().casefold()
                != winner.casefold(),
                str(player["Team"]).strip().casefold(),
                str(player["Player Name"]).strip().casefold()
            )
        )

        for row_number, player in enumerate(
            matching_players,
            start=1
        ):
            player_name = str(
                player["Player Name"]
            ).strip()

            player_key = player_name.casefold()

            display_values = [
                player_name,
                player["Draftee"],
                player["Team"]
            ]

            for column, value in enumerate(display_values):
                tk.Label(
                    stats_inner,
                    text=value,
                    anchor=tk.CENTER,
                    justify=tk.CENTER
                ).grid(
                    row=row_number,
                    column=column,
                    padx=8,
                    pady=4,
                    sticky="nsew"
                )



            player_stat_entries = {}

            for column, stat_name in enumerate(
                ["PTS", "REB", "AST"],
                start=3
            ):
                entry = tk.Entry(
                    stats_inner,
                    width=10,
                    justify=tk.CENTER
                )

                entry.grid(
                    row=row_number,
                    column=column,
                    padx=8,
                    pady=4,
                    sticky=""
                )

                player_stat_entries[stat_name] = entry


            stats_entries[player_key] = {
                "player_name": player_name,
                "entries": player_stat_entries
            }


    def fetch_espn_stats(
        self,
        window,
        winner_var,
        loser_var,
        stats_entries,
        refresh_callback,
        status_label,
        round_var=None
    ):
        refresh_callback()

        winner = self.resolve_autocomplete_value(
            winner_var.get(),
            self.teams
        )
        loser = self.resolve_autocomplete_value(
            loser_var.get(),
            self.teams
        )

        if winner is None or loser is None:
            messagebox.showwarning(
                "Teams Required",
                "Select both the Winner and Loser before fetching stats.",
                parent=window
            )
            return

        status_label.config(text="Contacting ESPN...")
        window.config(cursor="watch")
        window.update_idletasks()

        try:
            event = find_latest_espn_game(winner, loser)

            if event is None:
                status_label.config(text="")
                messagebox.showwarning(
                    "Game Not Found",
                    f"No ESPN game between {winner} and {loser} was found "
                    f"in the last {ESPN_LOOKBACK_DAYS} days.\n\n"
                    "Check the team names or enter stats manually.",
                    parent=window
                )
                return

            competition = event["competitions"][0]
            status = competition.get("status", event.get("status", {}))
            completed = status.get("type", {}).get("completed", False)

            if not completed and not messagebox.askyesno(
                "Game Not Final",
                f"{event.get('name', 'This game')} is not final yet "
                f"({status.get('type', {}).get('detail', 'in progress')})."
                "\n\nLoad the current stats anyway?",
                parent=window
            ):
                status_label.config(text="")
                return

            espn_winner = next(
                (
                    c.get("team", {})
                    for c in competition.get("competitors", [])
                    if c.get("winner")
                ),
                None
            )

            if (
                completed
                and espn_winner
                and espn_team_score(espn_winner, loser)
                > espn_team_score(espn_winner, winner)
            ):
                if messagebox.askyesno(
                    "Winner Mismatch",
                    f"ESPN reports {espn_winner.get('displayName', loser)} "
                    f"won this game.\n\nSwap Winner and Loser?",
                    parent=window
                ):
                    winner_var.set(loser)
                    loser_var.set(winner)
                    refresh_callback()

            box_score = read_espn_box_score(event["id"])

            detected_round = espn_round_name(event)

            if detected_round and round_var is not None:
                round_var.set(detected_round)

        except Exception as error:
            status_label.config(text="")
            messagebox.showerror(
                "ESPN Request Failed",
                f"Could not retrieve stats from ESPN:\n\n{error}",
                parent=window
            )
            return

        finally:
            window.config(cursor="")

        if not box_score:
            status_label.config(text="")
            messagebox.showwarning(
                "No Box Score",
                "ESPN has no box score for this game yet.",
                parent=window
            )
            return

        filled = []
        unmatched = []

        for player_data in stats_entries.values():
            player_name = player_data["player_name"]

            best_name, best_score = None, 0.0

            for espn_name in box_score:
                score = name_similarity(player_name, espn_name)

                if score > best_score:
                    best_name, best_score = espn_name, score

            if best_name is None or best_score < ESPN_PLAYER_MATCH_THRESHOLD:
                unmatched.append(player_name)

                for entry in player_data["entries"].values():
                    entry.delete(0, tk.END)
                    entry.insert(0, "0")

                continue

            for stat_name, entry in player_data["entries"].items():
                entry.delete(0, tk.END)
                entry.insert(
                    0,
                    str(box_score[best_name]["stats"].get(stat_name, 0))
                )

            filled.append(
                player_name
                if best_score == 1.0
                else f"{player_name} (ESPN: {best_name})"
            )

        summary = (
            f"Loaded {event.get('name', 'game')} "
            f"({event.get('date', '')[:10]}). "
            f"Filled {len(filled)} of {len(stats_entries)} drafted player(s)."
        )

        if detected_round:
            summary += f" Round set to {detected_round}."

        if unmatched:
            summary += (
                " Not in box score (recorded as 0): "
                + ", ".join(unmatched)
            )

        status_label.config(text=summary)

    def get_drafted_players_for_teams(self, selected_teams):
        if DRAFT_SHEET not in self.workbook.sheetnames:
            return []

        team_by_player = {}

        for record in self.players:
            player_name = record.get("Player")
            team_name = record.get("Team")

            if player_name and team_name:
                player_key = str(
                    player_name
                ).strip().casefold()

                team_by_player[player_key] = str(
                    team_name
                ).strip()

        draft_sheet = self.workbook[DRAFT_SHEET]
        matching_players = []

        for row in range(
            2,
            draft_sheet.max_row + 1
        ):
            draft_position = draft_sheet.cell(
                row=row,
                column=1
            ).value

            draftee = draft_sheet.cell(
                row=row,
                column=2
            ).value

            player_name = draft_sheet.cell(
                row=row,
                column=3
            ).value

            if not player_name:
                continue

            player_key = str(
                player_name
            ).strip().casefold()

            team_name = team_by_player.get(player_key)

            if (
                team_name
                and team_name.casefold() in selected_teams
            ):
                matching_players.append({
                    "Draft Position": draft_position,
                    "Draftee": draftee,
                    "Player Name": player_name,
                    "Team": team_name
                })

        return matching_players

    def save_game_result(
        self,
        window,
        round_var,
        winner_var,
        loser_var,
        stats_entries,
        winner_box=None,
        loser_box=None,
        stats_inner=None,
        recent_result_label=None,
        close_after_save=True
    ):
        selected_round = self.resolve_autocomplete_value(
            round_var.get(),
            ROUND_OPTIONS
        )

        winner_text = self.resolve_autocomplete_value(
            winner_var.get(),
            self.available_teams
        )

        loser_text = self.resolve_autocomplete_value(
            loser_var.get(),
            self.available_teams
        )

        if selected_round is None:
            messagebox.showwarning(
                "Round Required",
                "Select a valid round."
            )
            return

        if winner_text is None or loser_text is None:
            winner_any = self.resolve_autocomplete_value(
                winner_var.get(),
                self.teams
            )
            loser_any = self.resolve_autocomplete_value(
                loser_var.get(),
                self.teams
            )
            eliminated = [
                team
                for team, available in (
                    (winner_any, winner_text),
                    (loser_any, loser_text)
                )
                if team is not None and available is None
            ]

            if winner_any and loser_any and eliminated:
                _, last_entry = self.peek_undo_entry()

                is_reversed = (
                    last_entry is not None
                    and last_entry.get("winner", "").casefold()
                    == loser_any.casefold()
                    and last_entry.get("loser", "").casefold()
                    == winner_any.casefold()
                )

                if is_reversed:
                    if not messagebox.askyesno(
                        "Replace Previous Result",
                        f"The last saved result was "
                        f"{last_entry['round']}: {last_entry['winner']} "
                        f"def. {last_entry['loser']}.\n\n"
                        f"Undo it and save {winner_any} def. {loser_any} "
                        f"instead?",
                        parent=window
                    ):
                        return

                    if not self.undo_last_game_result(
                        window,
                        winner_box,
                        loser_box,
                        recent_result_label,
                        confirm=False
                    ):
                        return

                    winner_text, loser_text = winner_any, loser_any
                else:
                    messagebox.showwarning(
                        "Team Already Eliminated",
                        f"{', '.join(eliminated)} is already marked as "
                        f"eliminated.\n\nIf that was a mistake, use "
                        f"\"Undo Last Result\" to reverse the earlier game "
                        f"first.",
                        parent=window
                    )
                    return

        if winner_text is None or loser_text is None:
            messagebox.showwarning(
                "Incomplete Result",
                "Select valid Winner and Loser teams."
            )
            return

        if winner_text.casefold() == loser_text.casefold():
            messagebox.showwarning(
                "Invalid Result",
                "Winner and Loser must be different teams."
            )
            return

        round_var.set(selected_round)
        winner_var.set(winner_text)
        loser_var.set(loser_text)

        if not messagebox.askyesno(
            "Confirm Game Result",
            f"{selected_round}: {winner_text} defeats {loser_text}.\n\n"
            f"{loser_text} will be marked as eliminated and its "
            f"players' remaining stats set to OUT.\n\nSave this result?",
            parent=window
        ):
            return

        draft_sheet = self.workbook[DRAFT_SHEET]
        parsed_stats = []

        for player_data in stats_entries.values():
            player_name = player_data["player_name"]
            player_values = {}

            for stat_name, entry in player_data["entries"].items():
                value_text = entry.get().strip()

                if not value_text:
                    player_values[stat_name] = 0
                    continue

                try:
                    value = float(value_text)

                    if value < 0:
                        raise ValueError

                    if value.is_integer():
                        value = int(value)

                except ValueError:
                    messagebox.showwarning(
                        "Invalid Statistic",
                        f"{stat_name} for {player_name} "
                        "must be a non-negative number."
                    )
                    return

                player_values[stat_name] = value

            parsed_stats.append(
                (player_name, player_values)
            )

        draft_before = self.snapshot_draft_stats(draft_sheet)
        previous_result = dict(self.last_game_result)

        for player_name, player_values in parsed_stats:
            draft_row = self.find_draft_row(
                draft_sheet,
                player_name
            )

            if draft_row is None:
                messagebox.showerror(
                    "Player Not Found",
                    f'Could not find "{player_name}" '
                    "on the Draft sheet."
                )
                return

            for stat_name, value in player_values.items():
                column = STAT_COLUMN_MAP[
                    selected_round
                ][stat_name]

                draft_sheet.cell(
                    row=draft_row,
                    column=column
                ).value = value

        self.mark_unfilled_stats_as_out(
            loser_text
        )

        if TEAMS_OUT_SHEET not in self.workbook.sheetnames:
            teams_out_sheet = self.workbook.create_sheet(
                TEAMS_OUT_SHEET
            )

            teams_out_sheet["A1"] = "Team"
            teams_out_sheet["B1"] = "Round"
        else:
            teams_out_sheet = self.workbook[
                TEAMS_OUT_SHEET
            ]

            teams_out_sheet["A1"] = "Team"
            teams_out_sheet["B1"] = "Round"

        next_row = 2

        while teams_out_sheet.cell(
            row=next_row,
            column=1
        ).value not in (None, ""):
            next_row += 1

        teams_out_sheet.cell(
            row=next_row,
            column=1
        ).value = loser_text

        teams_out_sheet.cell(
            row=next_row,
            column=2
        ).value = selected_round

        self.save_last_game_result(
            selected_round,
            winner_text,
            loser_text
        )

        draft_after = self.snapshot_draft_stats(draft_sheet)

        self.push_undo_entry({
            "round": selected_round,
            "winner": winner_text,
            "loser": loser_text,
            "teams_out_row": next_row,
            "previous_result": previous_result,
            "cells": [
                [row, column, value]
                for (row, column), value in draft_before.items()
                if draft_after.get((row, column)) != value
            ]
        })

        if not self.safe_save():
            return

        if recent_result_label is not None:
            recent_result_label.config(
                text=self.last_game_result_text()
            )

        self.available_teams = self.get_available_teams()

        if close_after_save:
            window.destroy()
            return

        round_var.set("")
        winner_var.set("")
        loser_var.set("")

        if winner_box is not None:
            winner_box["values"] = self.available_teams

        if loser_box is not None:
            loser_box["values"] = self.available_teams

        if stats_inner is not None:
            for widget in stats_inner.winfo_children():
                grid_info = widget.grid_info()

                if int(grid_info.get("row", 0)) > 0:
                    widget.destroy()

        stats_entries.clear()

        if len(self.available_teams) < 2:
            self.game_results_button.config(
                state=tk.DISABLED
            )

    @staticmethod
    def snapshot_draft_stats(draft_sheet):
        snapshot = {}

        for row in range(2, draft_sheet.max_row + 1):
            for column in range(4, 22):
                value = draft_sheet.cell(row=row, column=column).value

                if not isinstance(value, (int, float, str, type(None))):
                    value = str(value)

                snapshot[(row, column)] = value

        return snapshot

    def get_undo_sheet(self):
        if UNDO_SHEET not in self.workbook.sheetnames:
            sheet = self.workbook.create_sheet(UNDO_SHEET)
            sheet["A1"] = "Undo Data (do not edit)"
            sheet["B1"] = "Description"
            sheet.sheet_state = "hidden"

        return self.workbook[UNDO_SHEET]

    def push_undo_entry(self, entry):
        sheet = self.get_undo_sheet()
        row = 2

        while sheet.cell(row=row, column=1).value not in (None, ""):
            row += 1

        sheet.cell(row=row, column=1).value = json.dumps(entry)
        sheet.cell(row=row, column=2).value = (
            f"{entry['round']}: {entry['winner']} def. {entry['loser']}"
        )

    def peek_undo_entry(self):
        if UNDO_SHEET not in self.workbook.sheetnames:
            return None, None

        sheet = self.workbook[UNDO_SHEET]

        for row in range(sheet.max_row, 1, -1):
            raw = sheet.cell(row=row, column=1).value

            if raw not in (None, ""):
                try:
                    return row, json.loads(raw)
                except ValueError:
                    return row, None

        return None, None

    def undo_last_game_result(
        self,
        window,
        winner_box=None,
        loser_box=None,
        recent_result_label=None,
        refresh_callback=None,
        confirm=True
    ):
        undo_row, entry = self.peek_undo_entry()

        if undo_row is None or entry is None:
            messagebox.showinfo(
                "Nothing to Undo",
                "There is no saved game result that can be undone.\n\n"
                "(Only results saved with this version can be undone.)",
                parent=window
            )
            return

        if confirm and not messagebox.askyesno(
            "Undo Last Result",
            f"Undo this result?\n\n{entry['round']}: "
            f"{entry['winner']} def. {entry['loser']}\n\n"
            f"Statistics entered for this game will be cleared, OUT marks "
            f"removed, and {entry['loser']} restored as an active team.",
            parent=window
        ):
            return

        draft_sheet = self.workbook[DRAFT_SHEET]

        for row, column, value in entry.get("cells", []):
            draft_sheet.cell(row=row, column=column).value = value

        if TEAMS_OUT_SHEET in self.workbook.sheetnames:
            teams_out_sheet = self.workbook[TEAMS_OUT_SHEET]
            loser_key = entry["loser"].strip().casefold()
            target_row = entry.get("teams_out_row")
            target_value = teams_out_sheet.cell(
                row=target_row, column=1
            ).value if target_row else None

            if str(target_value or "").strip().casefold() != loser_key:
                target_row = None

                for row in range(teams_out_sheet.max_row, 1, -1):
                    value = teams_out_sheet.cell(row=row, column=1).value

                    if str(value or "").strip().casefold() == loser_key:
                        target_row = row
                        break

            if target_row:
                teams_out_sheet.delete_rows(target_row)

        previous = entry.get("previous_result") or {}

        if previous.get("Winner") and previous.get("Loser"):
            self.save_last_game_result(
                previous.get("Round", ""),
                previous["Winner"],
                previous["Loser"]
            )
        else:
            if GAME_RESULTS_SHEET in self.workbook.sheetnames:
                results_sheet = self.workbook[GAME_RESULTS_SHEET]

                if results_sheet.max_row > 1:
                    results_sheet.delete_rows(2, results_sheet.max_row - 1)

            self.last_game_result = {"Round": "", "Winner": "", "Loser": ""}

        self.workbook[UNDO_SHEET].delete_rows(undo_row)

        if not self.safe_save():
            return

        self.available_teams = self.get_available_teams()

        if winner_box is not None:
            winner_box["values"] = self.available_teams

        if loser_box is not None:
            loser_box["values"] = self.available_teams

        if recent_result_label is not None:
            recent_result_label.config(text=self.last_game_result_text())

        if refresh_callback is not None:
            refresh_callback()

        self.game_results_button.config(
            state=tk.NORMAL if len(self.available_teams) >= 2 else tk.DISABLED
        )

        self.status_label.config(
            text=f"Undid result: {entry['winner']} def. {entry['loser']}."
        )

        return True

    @staticmethod
    def find_draft_row(draft_sheet, player_name):
        for row in range(
            2,
            draft_sheet.max_row + 1
        ):
            draft_player = draft_sheet.cell(
                row=row,
                column=3
            ).value

            if (
                draft_player is not None
                and str(draft_player).strip().casefold()
                == player_name.strip().casefold()
            ):
                return row

        return None

    def mark_unfilled_stats_as_out(self, team_name):
        draft_sheet = self.workbook[DRAFT_SHEET]

        team_by_player = {
            str(record["Player"]).strip().casefold():
            str(record["Team"]).strip().casefold()
            for record in self.players
            if record.get("Player") and record.get("Team")
        }

        for row in range(
            2,
            draft_sheet.max_row + 1
        ):
            player_name = draft_sheet.cell(
                row=row,
                column=3
            ).value

            if not player_name:
                continue

            player_key = str(
                player_name
            ).strip().casefold()

            player_team = team_by_player.get(player_key)

            if player_team != team_name.strip().casefold():
                continue

            for column in range(4, 22):
                cell = draft_sheet.cell(
                    row=row,
                    column=column
                )

                if cell.value in (None, ""):
                    cell.value = "OUT"

    def find_team_list_sheet(self):
        if self.workbook is None:
            return None

        for name in self.workbook.sheetnames:
            if name.strip().casefold() == TEAM_LIST_SHEET.casefold():
                return self.workbook[name]

        return None

    def read_team_list(self):
        worksheet = self.find_team_list_sheet()

        if worksheet is None:
            raise ValueError(f'Missing sheet: "{TEAM_LIST_SHEET}"')

        columns = {"Region": 1, "Seed": 2, "Team": 3}
        header_row = [
            str(cell.value or "").strip().casefold()
            for cell in worksheet[1]
        ]

        for key in columns:
            if key.casefold() in header_row:
                columns[key] = header_row.index(key.casefold()) + 1

        teams = []

        for row in range(2, worksheet.max_row + 1):
            team = str(
                worksheet.cell(row=row, column=columns["Team"]).value or ""
            ).strip()

            if not team:
                continue

            teams.append({
                "Region": worksheet.cell(
                    row=row, column=columns["Region"]
                ).value,
                "Seed": worksheet.cell(
                    row=row, column=columns["Seed"]
                ).value,
                "Team": team,
            })

        return teams

    def read_players_per_team(self):
        if CONFIG_SHEET not in self.workbook.sheetnames:
            raise ValueError(f'Missing sheet: "{CONFIG_SHEET}"')

        value = self.workbook[CONFIG_SHEET][PLAYERS_PER_TEAM_CELL].value

        try:
            count = int(float(str(value).strip()))
        except (TypeError, ValueError):
            count = 0

        if count < 1:
            raise ValueError(
                f"{CONFIG_SHEET}!{PLAYERS_PER_TEAM_CELL} must be a whole "
                f"number of players per team (found: {value!r})."
            )

        return count

    def pull_team_stats(self):
        try:
            teams = self.read_team_list()
            players_per_team = self.read_players_per_team()
        except Exception as error:
            messagebox.showerror("Pull Team Stats", str(error))
            return

        if not teams:
            messagebox.showerror(
                "Pull Team Stats",
                f'No teams found on sheet "{TEAM_LIST_SHEET}".'
            )
            return

        if not messagebox.askyesno(
            "Pull Team Stats",
            f"Pull the top {players_per_team} players for "
            f"{len(teams)} teams from ESPN?\n\n"
            f'This replaces all player rows on "{TEAM_DATA_SHEET}".'
        ):
            return

        self.team_stats_button.config(state=tk.DISABLED)
        self.root.config(cursor="watch")

        try:
            espn_teams = fetch_espn_team_index()
        except Exception as error:
            self.root.config(cursor="")
            self.team_stats_button.config(state=tk.NORMAL)
            messagebox.showerror(
                "Pull Team Stats",
                f"Unable to reach ESPN: {error}"
            )
            return

        rows = []

        progress_window = tk.Toplevel(self.root)
        progress_window.title("Pulling Team Stats")
        progress_window.resizable(False, False)
        progress_window.transient(self.root)
        progress_window.protocol("WM_DELETE_WINDOW", lambda: None)

        progress_label = tk.Label(
            progress_window, text="Starting...", anchor="w", width=50
        )
        progress_label.pack(fill=tk.X, padx=16, pady=(14, 6))

        progress_bar = ttk.Progressbar(
            progress_window,
            orient=tk.HORIZONTAL,
            length=380,
            mode="determinate",
            maximum=len(teams)
        )
        progress_bar.pack(padx=16, pady=6)

        eta_label = tk.Label(
            progress_window, text="", anchor="w", fg=THEME["muted"]
        )
        eta_label.pack(fill=tk.X, padx=16, pady=(0, 14))

        progress_window.update_idletasks()
        progress_window.geometry(
            f"+{self.root.winfo_rootx() + 60}+{self.root.winfo_rooty() + 120}"
        )
        progress_window.grab_set()
        started = datetime.now()

        for index, team in enumerate(teams, start=1):
            self.status_label.config(
                text=f"Pulling ESPN stats ({index}/{len(teams)}): {team['Team']}"
            )
            progress_label.config(
                text=f"Team {index} of {len(teams)}: {team['Team']}"
            )
            completed = index - 1

            if completed:
                elapsed = (datetime.now() - started).total_seconds()
                remaining = int(
                    elapsed / completed * (len(teams) - completed)
                )
                eta_label.config(
                    text=(
                        f"{completed * 100 // len(teams)}% complete — "
                        f"about {remaining // 60}m {remaining % 60:02d}s "
                        "remaining"
                    )
                )
            else:
                eta_label.config(text="Estimating time remaining...")

            self.root.update()

            base = {
                "Region": team["Region"],
                "Seed": team["Seed"],
                "Team": team["Team"],
            }

            try:
                team_id = match_espn_team_id(espn_teams, team["Team"])

                if team_id is None:
                    raise ValueError("Team not found on ESPN")

                players = fetch_espn_team_top_players(
                    team_id, players_per_team
                )

                if not players:
                    raise ValueError("No player stats found")

                rows.extend({**base, **player} for player in players)
            except Exception as error:
                rows.append({**base, ERROR_HEADER: str(error)})

            progress_bar["value"] = index

        progress_window.grab_release()
        progress_window.destroy()
        self.root.config(cursor="")
        self.team_stats_button.config(state=tk.NORMAL)
        self.write_team_data_rows(rows)

    def write_team_data_rows(self, rows):
        if TEAM_DATA_SHEET in self.workbook.sheetnames:
            worksheet = self.workbook[TEAM_DATA_SHEET]
        else:
            worksheet = self.workbook.create_sheet(TEAM_DATA_SHEET)

        existing_statuses = {}

        for row in range(PLAYER_START_ROW, worksheet.max_row + 1):
            status = normalize_status(
                worksheet.cell(row=row, column=STATUS_COLUMN).value
            )

            if status:
                key = (
                    normalize_name(worksheet.cell(row=row, column=3).value),
                    normalize_name(worksheet.cell(row=row, column=4).value),
                )
                existing_statuses[key] = status

        if worksheet.max_row >= PLAYER_START_ROW:
            worksheet.delete_rows(
                PLAYER_START_ROW,
                worksheet.max_row - PLAYER_START_ROW + 1
            )

        headers = TEAM_HEADERS + [ERROR_HEADER, STATUS_HEADER]

        for column, header in enumerate(headers, start=1):
            worksheet.cell(row=1, column=column, value=header)

        errors = 0

        for offset, record in enumerate(rows):
            row = PLAYER_START_ROW + offset

            for column, header in enumerate(TEAM_HEADERS, start=1):
                worksheet.cell(row=row, column=column, value=record.get(header))

            if record.get(ERROR_HEADER):
                errors += 1
                worksheet.cell(
                    row=row, column=ERROR_COLUMN, value=record[ERROR_HEADER]
                )

            status = existing_statuses.get((
                normalize_name(record.get("Team")),
                normalize_name(record.get("Player")),
            ))

            if status:
                worksheet.cell(row=row, column=STATUS_COLUMN, value=status)

        if not self.safe_save():
            self.status_label.config(text="Team stats were not saved.")
            return

        self.load_workbook_file(self.file_path)

        message = f"Wrote {len(rows) - errors} player rows to {TEAM_DATA_SHEET}."

        if errors:
            message += (
                f"\n\n{errors} team(s) could not be pulled. "
                f"See the {ERROR_HEADER} column (P)."
            )

        messagebox.showinfo("Pull Team Stats", message)

    def safe_save(self):
        while True:
            try:
                self.workbook.save(self.file_path)
                return True
            except PermissionError:
                retry = messagebox.askretrycancel(
                    "Workbook Locked",
                    f"Unable to save '{Path(self.file_path).name}'.\n\n"
                    "Close the file in Excel (or any other program) "
                    "and click Retry."
                )

                if not retry:
                    return False
            except Exception as error:
                messagebox.showerror(
                    "Save Failed",
                    str(error)
                )
                return False

    def read_existing_picks(self):
        draft_sheet = self.workbook[DRAFT_SHEET]
        picks = {}

        for row in range(
            2,
            draft_sheet.max_row + 1
        ):
            position, drafter, player = (
                draft_sheet.cell(row=row, column=column).value
                for column in range(1, 4)
            )

            if position is None or player in (None, ""):
                continue

            picks[str(position).strip()] = (
                str(drafter).strip() if drafter else "",
                str(player).strip()
            )

        return picks

    @staticmethod
    def get_next_draft_row(worksheet):
        row = 2

        while any(
            worksheet.cell(
                row=row,
                column=column
            ).value is not None
            for column in range(1, 4)
        ):
            row += 1

        return row

    def build_draft_slots(self):
        slots = []

        for round_number in range(1, TOTAL_ROUNDS + 1):
            draft_order = (
                self.drafters
                if round_number % 2 == 1
                else list(reversed(self.drafters))
            )

            for position_in_round, drafter in enumerate(
                draft_order,
                start=1
            ):
                slots.append(
                    (f"{round_number}.{position_in_round}", drafter)
                )

        return slots

    @staticmethod
    def on_deck_text(slots, slot_index):
        drafter = slots[slot_index][1]
        parts = []

        if slot_index + 1 < len(slots):
            parts.append(f"On deck: {slots[slot_index + 1][1]}")

            if slot_index + 2 < len(slots):
                parts.append(f"In the hole: {slots[slot_index + 2][1]}")
        else:
            parts.append("Final pick of the draft")

        for offset, (_, future_drafter) in enumerate(
            slots[slot_index + 1:],
            start=1
        ):
            if future_drafter == drafter:
                parts.append(
                    f"{drafter} picks again in {offset} pick"
                    f"{'' if offset == 1 else 's'}"
                )
                break
        else:
            parts.append(f"This is {drafter}'s last pick")

        return "   \u2022   ".join(parts)

    @staticmethod
    def make_history_entry(draft_position, drafter, player_name, record):
        record = record or {}

        return {
            "Draft Position": draft_position,
            "Draftee": drafter,
            "Player Name": player_name,
            "Region": record.get("Region"),
            "Seed": record.get("Seed"),
            "Team": record.get("Team"),
            "Tot. FPts": record.get("Tot. FPts"),
            "PPG": record.get("PPG"),
            "RPG": record.get("RPG"),
            "APG": record.get("APG")
        }

    def undo_last_pick(
        self,
        slots,
        draft_history,
        available_players,
        record_by_player,
        existing_picks,
        dialog
    ):
        """Remove the most recent pick. Returns the slot index to redo."""
        if not draft_history:
            return None

        last_pick = draft_history[-1]
        draft_position = last_pick["Draft Position"]
        player_name = last_pick["Player Name"]

        confirmed = messagebox.askyesno(
            "Undo Last Pick",
            f"Undo pick {draft_position}: "
            f"{last_pick['Draftee']} \u2192 {player_name}?\n\n"
            "The player will be returned to the available pool.",
            parent=dialog.window
        )

        if not confirmed:
            return None

        draft_sheet = self.workbook[DRAFT_SHEET]
        sheet_row = None

        for row in range(draft_sheet.max_row, 1, -1):
            if (
                str(draft_sheet.cell(row=row, column=1).value or "").strip()
                == draft_position
            ):
                sheet_row = row
                break

        if sheet_row is not None:
            for column in range(1, 4):
                draft_sheet.cell(row=sheet_row, column=column).value = None

        if not self.safe_save():
            if sheet_row is not None:
                draft_sheet.cell(row=sheet_row, column=1).value = draft_position
                draft_sheet.cell(
                    row=sheet_row, column=2
                ).value = last_pick["Draftee"]
                draft_sheet.cell(row=sheet_row, column=3).value = player_name

            return None

        draft_history.pop()
        existing_picks.pop(draft_position, None)
        record = record_by_player.get(str(player_name).strip().casefold())

        if record is not None and record not in available_players:
            available_players.append(record)

        self.status_label.config(
            text=f"Undid pick {draft_position}: {player_name}"
        )

        slot_positions = [position for position, _ in slots]
        return slot_positions.index(draft_position)

    def open_player_statuses(self):
        if self.workbook is None:
            messagebox.showwarning(
                "No Workbook",
                "Please select an Excel file first."
            )
            return

        window = tk.Toplevel(self.root)
        window.title("Update Player Statuses")
        window.geometry("760x620")
        window.transient(self.root)
        window.grab_set()

        tk.Label(
            window,
            text="Player Statuses",
            font=(THEME["font"], 14, "bold"),
            bg=THEME["primary"],
            fg="white",
            pady=10
        ).pack(fill=tk.X)

        filter_row = tk.Frame(window)
        filter_row.pack(fill=tk.X, padx=10, pady=(10, 4))

        tk.Label(filter_row, text="Search player or team:").pack(side=tk.LEFT)
        filter_var = tk.StringVar(window)
        filter_entry = tk.Entry(filter_row, textvariable=filter_var, width=40)
        filter_entry.pack(side=tk.LEFT, padx=6)

        flagged_only_var = tk.BooleanVar(window, value=False)
        tk.Checkbutton(
            filter_row,
            text="Only show Out/Questionable",
            variable=flagged_only_var,
            command=lambda: refresh()
        ).pack(side=tk.LEFT, padx=10)

        table_frame = tk.Frame(window)
        table_frame.pack(fill=tk.BOTH, expand=True, padx=10, pady=4)

        columns = ("team", "player", "status")
        tree = ttk.Treeview(
            table_frame,
            columns=columns,
            show="headings",
            selectmode="extended"
        )

        for column_id, heading, width in (
            ("team", "Team", 200),
            ("player", "Player", 280),
            ("status", "Status", 140),
        ):
            tree.heading(column_id, text=heading, anchor=tk.CENTER)
            tree.column(column_id, width=width, anchor=tk.CENTER)

        tree.tag_configure("out", background=HINT_COLORS["out"])
        tree.tag_configure(
            "questionable",
            background=HINT_COLORS["questionable"]
        )

        scrollbar = ttk.Scrollbar(
            table_frame,
            orient=tk.VERTICAL,
            command=tree.yview
        )
        tree.configure(yscrollcommand=scrollbar.set)
        tree.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        scrollbar.pack(side=tk.RIGHT, fill=tk.Y)

        record_by_row = {
            str(record["_excel_row"]): record
            for record in self.players
        }

        def refresh():
            selected = set(tree.selection())
            tree.delete(*tree.get_children())
            text = filter_var.get().strip().casefold()

            for record in sorted(
                self.players,
                key=lambda item: (
                    str(item.get("Team") or "").casefold(),
                    str(item.get("Player") or "").casefold()
                )
            ):
                status = normalize_status(record.get("Status"))

                if flagged_only_var.get() and not status:
                    continue

                if text and text not in (
                    f"{record.get('Player')} {record.get('Team')}".casefold()
                ):
                    continue

                row_id = str(record["_excel_row"])
                tree.insert(
                    "",
                    tk.END,
                    iid=row_id,
                    values=(
                        record.get("Team"),
                        record.get("Player"),
                        status
                    ),
                    tags=(status.casefold(),) if status else ()
                )

            for row_id in selected:
                if tree.exists(row_id):
                    tree.selection_add(row_id)

        def apply_status(status):
            selection = tree.selection()

            if not selection:
                messagebox.showwarning(
                    "Select Players",
                    "Select one or more players first.",
                    parent=window
                )
                return

            team_sheet = self.workbook[TEAM_DATA_SHEET]

            if team_sheet.cell(row=1, column=STATUS_COLUMN).value in (
                None,
                ""
            ):
                team_sheet.cell(
                    row=1,
                    column=STATUS_COLUMN
                ).value = STATUS_HEADER

            previous = {}

            for row_id in selection:
                record = record_by_row[row_id]
                cell = team_sheet.cell(
                    row=record["_excel_row"],
                    column=STATUS_COLUMN
                )
                previous[row_id] = (cell.value, record.get("Status"))
                cell.value = status or None
                record["Status"] = status

            if not self.safe_save():
                for row_id, (cell_value, record_status) in previous.items():
                    record = record_by_row[row_id]
                    team_sheet.cell(
                        row=record["_excel_row"],
                        column=STATUS_COLUMN
                    ).value = cell_value
                    record["Status"] = record_status
                return

            label = status or "cleared"
            self.status_label.config(
                text=f"Status {label} for {len(selection)} player(s)."
            )
            refresh()

        filter_entry.bind("<KeyRelease>", lambda event: refresh())

        menu = tk.Menu(window, tearoff=0)
        menu.add_command(label="Mark Out", command=lambda: apply_status("Out"))
        menu.add_command(
            label="Mark Questionable",
            command=lambda: apply_status("Questionable")
        )
        menu.add_separator()
        menu.add_command(label="Clear status", command=lambda: apply_status(""))

        def show_menu(event):
            row_id = tree.identify_row(event.y)

            if row_id and row_id not in tree.selection():
                tree.selection_set(row_id)

            menu.tk_popup(event.x_root, event.y_root)

        tree.bind("<Button-3>", show_menu)

        button_row = tk.Frame(window)
        button_row.pack(fill=tk.X, padx=10, pady=10)

        tk.Button(
            button_row,
            text="Mark Out",
            bg="#F8C8C8",
            fg="black",
            activeforeground="black",
            activebackground="#F0B0B0",
            width=16,
            command=lambda: apply_status("Out")
        ).pack(side=tk.LEFT, padx=4)
        tk.Button(
            button_row,
            text="Mark Questionable",
            bg="#FFD8A8",
            fg="black",
            activeforeground="black",
            activebackground="#FFC88A",
            width=18,
            command=lambda: apply_status("Questionable")
        ).pack(side=tk.LEFT, padx=4)
        tk.Button(
            button_row,
            text="Clear Status",
            width=14,
            command=lambda: apply_status("")
        ).pack(side=tk.LEFT, padx=4)
        tk.Button(
            button_row,
            text="Done",
            width=12,
            command=window.destroy
        ).pack(side=tk.RIGHT, padx=4)

        tk.Label(
            window,
            text=(
                "Tip: Ctrl/Shift-click to select several players, "
                "or right-click a player. Changes save to TeamData column Q."
            ),
            fg=THEME["muted"]
        ).pack(pady=(0, 8))

        refresh()
        filter_entry.focus_set()

    def begin_draft(self):
        if self.workbook is None:
            messagebox.showwarning(
                "No Workbook",
                "Please select an Excel file first."
            )
            return

        total_picks = len(self.drafters) * TOTAL_ROUNDS

        if len(self.players) < total_picks:
            messagebox.showerror(
                "Not Enough Players",
                f"You need {total_picks} players, but only "
                f"{len(self.players)} are available."
            )
            return

        draft_sheet = self.workbook[DRAFT_SHEET]

        draft_sheet["A1"] = "Draft Position"
        draft_sheet["B1"] = "Drafter"
        draft_sheet["C1"] = "Player Name"

        pick_timer_seconds = self.read_pick_timer_seconds(
            self.workbook[CONFIG_SHEET]
        )

        available_players = self.players.copy()
        draft_history = []
        existing_picks = self.read_existing_picks()
        record_by_player = {
            str(record["Player"]).strip().casefold(): record
            for record in self.players
            if record.get("Player")
        }

        if existing_picks:
            resume = messagebox.askyesno(
                "Resume Draft",
                f"{len(existing_picks)} pick(s) already exist in the "
                f"Draft sheet.\n\nResume the draft from the next "
                f"open pick?"
            )

            if not resume:
                self.status_label.config(
                    text="Draft not started."
                )
                return

        dialog = None
        slots = self.build_draft_slots()
        slot_index = 0

        while slot_index < len(slots):
            draft_position, drafter = slots[slot_index]

            if draft_position in existing_picks:
                _, existing_player = existing_picks[draft_position]
                existing_record = record_by_player.get(
                    existing_player.casefold()
                )

                if existing_record in available_players:
                    available_players.remove(existing_record)

                draft_history.append(
                    self.make_history_entry(
                        draft_position,
                        drafter,
                        existing_player,
                        existing_record
                    )
                )
                slot_index += 1
                continue

            on_deck_text = self.on_deck_text(slots, slot_index)

            if dialog is None:
                dialog = PlayerSelectionDialog(
                    self.root,
                    drafter,
                    draft_position,
                    available_players,
                    draft_history,
                    on_deck_text,
                    pick_timer_seconds
                )
            else:
                dialog.next_pick(
                    drafter,
                    draft_position,
                    draft_history,
                    on_deck_text
                )

            selected_record = dialog.wait_for_pick()

            if dialog.undo_requested:
                undone_index = self.undo_last_pick(
                    slots,
                    draft_history,
                    available_players,
                    record_by_player,
                    existing_picks,
                    dialog
                )

                if undone_index is not None:
                    slot_index = undone_index

                continue

            if selected_record is None:
                dialog.close()
                self.status_label.config(
                    text="Draft paused. Click Begin / Resume Draft to continue."
                )
                return

            selected_player = selected_record["Player"]
            available_players.remove(selected_record)

            draft_history.append(
                self.make_history_entry(
                    draft_position,
                    drafter,
                    selected_player,
                    selected_record
                )
            )

            draft_row = self.get_next_draft_row(
                draft_sheet
            )

            draft_sheet.cell(
                row=draft_row,
                column=1
            ).value = draft_position

            draft_sheet.cell(
                row=draft_row,
                column=2
            ).value = drafter

            draft_sheet.cell(
                row=draft_row,
                column=3
            ).value = selected_player

            if not self.safe_save():
                dialog.close()
                self.status_label.config(
                    text="Draft paused: workbook could not be saved."
                )
                return

            self.status_label.config(
                text=f"Pick {draft_position}: {drafter} \u2192 {selected_player}"
            )
            slot_index += 1

        if dialog is not None:
            dialog.close()

        self.status_label.config(
            text="Draft complete and saved to the Draft sheet."
        )

        self.begin_button.config(
            state=tk.DISABLED
        )

        messagebox.showinfo(
            "Draft Complete",
            "The draft has been completed and saved."
        )

    def is_draft_complete(self):
        if DRAFT_SHEET not in self.workbook.sheetnames:
            return False

        draft_sheet = self.workbook[DRAFT_SHEET]

        required_picks = len(self.drafters) * TOTAL_ROUNDS
        completed_picks = 0

        for row in range(
            2,
            draft_sheet.max_row + 1
        ):
            row_values = [
                draft_sheet.cell(
                    row=row,
                    column=column
                ).value
                for column in range(1, 4)
            ]

            if all(
                value is not None and str(value).strip()
                for value in row_values
            ):
                completed_picks += 1

        return completed_picks >= required_picks

def hide_own_console_window():
    """Hide the console only if it was created solely for this process."""
    try:
        import ctypes

        kernel32 = ctypes.windll.kernel32
        console_window = kernel32.GetConsoleWindow()

        if not console_window:
            return

        process_ids = (ctypes.c_uint * 2)()

        if kernel32.GetConsoleProcessList(process_ids, 2) == 1:
            ctypes.windll.user32.ShowWindow(console_window, 0)
    except Exception:
        pass


if __name__ == "__main__":
    hide_own_console_window()
    root = tk.Tk()
    app = DraftApp(root)
    root.mainloop()
