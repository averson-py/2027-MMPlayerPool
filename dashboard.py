
import colorsys
import os
import re
import secrets
from datetime import datetime
from pathlib import Path


import openpyxl
import pandas as pd
import requests
import streamlit as st
import streamlit.components.v1 as components


# ---------------------------------------------------------------------
# Configuration
# ---------------------------------------------------------------------
DEFAULT_WORKBOOK = "2027 March Madness Player Pool Tracking Skeleton 2026 rnd5.xlsx"
ONEDRIVE_FILE_PATH = "2027MMPlayerPoolStreamlit/2027 March Madness Player Pool Tracking.xlsx"
# Public read-only OneDrive sharing link. Store this in Streamlit Secrets as ONEDRIVE_SHARE_URL.
GRAPH_BASE_URL = "https://graph.microsoft.com/v1.0"
OAUTH_AUTHORITY = "https://login.microsoftonline.com/common/oauth2/v2.0"
GRAPH_SCOPES = "Files.Read User.Read offline_access"

DRAFT_ORDER_SHEET = "DraftOrder"
DRAFT_SHEET = "Draft"
TEAM_DATA_SHEET = "TeamData"
TEAMS_OUT_SHEET = "TeamsOut"
GAME_RESULTS_SHEET = "GameResults"
CONFIG_SHEET = "Config"
TEAM_LIST_SHEET = "TeamList"

DRAFTER_START_ROW = 2

SEED_AVERAGE_GAMES = {
    1: 4.34, 2: 3.35, 3: 2.84, 4: 2.57,
    5: 2.14, 6: 2.05, 7: 1.88, 8: 1.69,
    9: 1.63, 10: 1.58, 11: 1.53, 12: 1.50,
    13: 1.24, 14: 1.15, 15: 1.10, 16: 1.01,
}

# The original Tkinter app uses these same table colors.
BLUE = "#4472C4"
LIGHT_BLUE = "#DDEBF7"
LIGHT_GREEN = "#E2F0D9"
LIGHT_YELLOW = "#FFF2CC"
LIGHT_RED = "#F4CCCC"
LIGHT_ORANGE = "#FCE4D6"
LIGHT_PURPLE = "#D9D2E9"


# ---------------------------------------------------------------------
# Data / calculation layer
# This intentionally contains no Tkinter code, so it can run locally
# and later on Streamlit Community Cloud/Render.
# ---------------------------------------------------------------------
def sum_numeric_values(values):
    numeric_values = []
    for value in values:
        try:
            numeric_values.append(float(value))
        except (TypeError, ValueError):
            continue
    if not numeric_values:
        return None
    total = sum(numeric_values)
    return int(total) if total.is_integer() else total


def sum_existing_values(values):
    return sum_numeric_values(values)


def weighted_fantasy_points(points, rebounds, assists, weights):
    """Calculate fantasy points using the scoring weights from Config!A2:C2."""
    values = [points, rebounds, assists]
    numeric = []
    for value in values:
        try:
            numeric.append(float(value))
        except (TypeError, ValueError):
            numeric.append(None)

    weighted = []
    for value, weight in zip(numeric, weights):
        if value is not None:
            weighted.append(value * weight)

    if not weighted:
        return None

    total = sum(weighted)
    return int(total) if total.is_integer() else total


def format_stat(value):
    if value is None:
        return ""
    if isinstance(value, float):
        return f"{value:.1f}"
    return str(value)


def format_signed_stat(value):
    if value is None:
        return ""
    try:
        return f"{float(value):+.1f}"
    except (TypeError, ValueError):
        return ""


def draft_position_sort_key(value):
    try:
        round_number, pick_number = str(value).split(".", 1)
        return int(round_number), int(pick_number)
    except (TypeError, ValueError):
        return 999999, str(value)


def count_blank_values(values):
    return sum(
        value is None or (isinstance(value, str) and not value.strip())
        for value in values
    )


def count_values_not_out(values):
    count = 0
    for value in values:
        if value is None:
            continue
        if isinstance(value, str):
            normalized = value.strip().casefold()
            if not normalized or normalized == "out":
                continue
        count += 1
    return count


def count_game_statuses(values):
    games_played = 0
    games_not_input = 0

    for start in range(0, len(values), 3):
        game_values = values[start:start + 3]
        has_out_marker = any(
            isinstance(value, str) and value.strip().casefold() == "out"
            for value in game_values
        )
        if has_out_marker:
            continue

        def is_blank(value):
            return value is None or (
                isinstance(value, str) and not value.strip()
            )

        if all(is_blank(value) for value in game_values):
            games_not_input += 1
        else:
            games_played += 1

    return games_played, games_not_input


def average_games_for_seed(seed):
    if seed is None:
        return None
    try:
        seed_number = int(float(seed))
    except (TypeError, ValueError):
        match = re.search(r"\d+", str(seed))
        if match is None:
            return None
        seed_number = int(match.group(0))
    return SEED_AVERAGE_GAMES.get(seed_number)


def color_scale(value, minimum, maximum, reverse=False):
    if value is None or minimum is None or maximum is None:
        return "#FFFFFF"
    if maximum == minimum:
        return "#FFEB84"

    ratio = (float(value) - float(minimum)) / (float(maximum) - float(minimum))
    ratio = max(0.0, min(1.0, ratio))
    if reverse:
        ratio = 1.0 - ratio

    red = (248, 105, 107)
    yellow = (255, 235, 132)
    green = (99, 190, 123)

    if ratio <= 0.5:
        blend = ratio * 2
        start, end = red, yellow
    else:
        blend = (ratio - 0.5) * 2
        start, end = yellow, green

    rgb = tuple(
        round(start[i] + (end[i] - start[i]) * blend)
        for i in range(3)
    )
    return "#%02X%02X%02X" % rgb


def red_white_green_scale(value, minimum, maximum, reverse=False):
    if value is None or minimum is None or maximum is None:
        return "#FFFFFF"
    if maximum == minimum:
        return "#FFFFFF"

    ratio = (float(value) - float(minimum)) / (float(maximum) - float(minimum))
    ratio = max(0.0, min(1.0, ratio))
    if reverse:
        ratio = 1.0 - ratio

    red = (248, 105, 107)
    white = (255, 255, 255)
    green = (99, 190, 123)

    if ratio <= 0.5:
        blend = ratio * 2
        start, end = red, white
    else:
        blend = (ratio - 0.5) * 2
        start, end = white, green

    rgb = tuple(
        round(start[i] + (end[i] - start[i]) * blend)
        for i in range(3)
    )
    return "#%02X%02X%02X" % rgb


# Tournament bracket functions
TEAM_SEEDS = {}

ROUND_LOST_INDEX = [
    ("first four", -1),
    ("first round", 0),
    ("round of 64", 0),
    ("second round", 1),
    ("round of 32", 1),
    ("sweet", 2),
    ("elite", 3),
    ("final four", 4),
    ("semi", 4),
    ("title", 5),
    ("champ", 5),
]


def round_lost_index(round_text):
    """Map a TeamsOut round label to a bracket round index (-1 = First Four)."""
    text = str(round_text or "").strip().casefold()
    for key, idx in ROUND_LOST_INDEX:
        if key in text:
            return idx
    return None


def lost_in_round(team, round_idx, teams_out):
    """True only if the team was eliminated in this specific round."""
    if team is None or "TBD" in str(team):
        return False
    key = str(team).casefold()
    if key not in teams_out:
        return False
    lost_idx = round_lost_index(teams_out[key])
    if lost_idx is None:
        return True
    return lost_idx == round_idx

ESPN_TEAMS_URL = (
    "https://site.api.espn.com/apis/site/v2/sports/basketball/"
    "mens-college-basketball/teams?limit=500"
)
ESPN_LOGO_URL = "https://a.espncdn.com/i/teamlogos/ncaa/500/{}.png"


def _norm_team_name(name):
    name = str(name).casefold().replace("saint", "st")
    return re.sub(r"[^a-z0-9]", "", name)


@st.cache_data(ttl=86400, show_spinner=False)
def get_espn_team_ids():
    """Return {normalized team name: ESPN team id} from ESPN's teams API."""
    ids = {}
    try:
        resp = requests.get(ESPN_TEAMS_URL, timeout=10)
        resp.raise_for_status()
        teams = resp.json()["sports"][0]["leagues"][0]["teams"]
    except Exception:
        return ids

    for entry in teams:
        team = entry.get("team", {})
        team_id = team.get("id")
        if not team_id:
            continue
        for key in ("location", "shortDisplayName", "displayName", "name"):
            value = team.get(key)
            if value:
                ids.setdefault(_norm_team_name(value), team_id)
    return ids


def get_logo_url(team):
    if not team or "TBD" in str(team):
        return None
    team_id = get_espn_team_ids().get(_norm_team_name(team))
    return ESPN_LOGO_URL.format(team_id) if team_id else None


def _logo_img(team):
    url = get_logo_url(team)
    if not url:
        return '<span class="team-logo placeholder"></span>'
    return (
        f'<img class="team-logo" src="{url}" alt="" loading="lazy" '
        f'onerror="this.style.visibility=\'hidden\'">'
    )


def lookup_seed(team, default):
    return TEAM_SEEDS.get(str(team).casefold(), default)


def build_tournament_bracket(bracket_teams, teams_out=None):
    """
    Build tournament bracket structure from teams dict.
    bracket_teams: {seed: {region: [teams]}}
    teams_out: {team_name_lower: round_lost} - teams that were eliminated

    Supports First Four games when multiple teams share the same seed in a region.
    Uses TeamsOut data to determine winners. If a team is in TeamsOut, it lost in that round.
    Returns a list of rounds, where each round contains regions,
    and each region contains matchups.
    """
    if teams_out is None:
        teams_out = {}

    regions = ["EAST", "WEST", "MIDWEST", "SOUTH"]
    rounds = []

    # Create a reverse mapping: team name -> (seed, region) for quick lookup
    # This helps us find which seed a team belongs to when it's in teams_out
    team_to_seed_region = {}
    for seed in bracket_teams:
        for region in bracket_teams[seed]:
            for team in bracket_teams[seed][region]:
                team_to_seed_region[team.casefold()] = (seed, region)
                TEAM_SEEDS[team.casefold()] = seed

    # Track winners at each round to use in future rounds
    # round_winners[round_num][region] = {matchup_idx: winner_team_name}
    round_winners = {}
    round_seeds = {}

    # Track which seeds have First Four games (multiple teams at same seed)
    first_four_seeds = set()

    # First detect First Four games (same seed, same region, multiple teams)
    first_four = {}
    for region in regions:
        first_four[region] = []

    for seed in bracket_teams:
        for region in bracket_teams[seed]:
            teams = bracket_teams[seed][region]
            if len(teams) > 1:
                # Multiple teams at same seed in same region = First Four game
                first_four_seeds.add(seed)
                # Pair consecutive teams for First Four
                for i in range(0, len(teams), 2):
                    if i + 1 < len(teams):
                        team_a = teams[i]
                        team_b = teams[i + 1]

                        # Determine winner based on TeamsOut
                        team_a_lost = lost_in_round(team_a, -1, teams_out)
                        team_b_lost = lost_in_round(team_b, -1, teams_out)

                        first_four[region].append({
                            "team_a": team_a,
                            "seed_a": seed,
                            "team_b": team_b,
                            "seed_b": seed,
                            "round": "First Four",
                            "winner": team_b if team_a_lost else (team_a if team_b_lost else None),
                            "team_a_eliminated": team_a_lost,
                            "team_b_eliminated": team_b_lost,
                        })
                    elif i < len(teams):
                        # Odd team out gets a bye shown as TBD opponent
                        team_a = teams[i]
                        first_four[region].append({
                            "team_a": team_a,
                            "seed_a": seed,
                            "team_b": f"({seed}) TBD (FF)",
                            "seed_b": seed,
                            "round": "First Four",
                            "winner": team_a,
                            "team_a_eliminated": False,
                            "team_b_eliminated": False,
                        })

    # Add First Four round if it has any games
    if any(first_four.values()):
        rounds.append({"name": "First Four", "regions": first_four})

    # Helper function to get team for a seed in a region, considering First Four winners
    def get_team_for_seed(seed, region, round_num):
        # Round of 64 - show actual teams from bracket
        if round_num == 0:
            if seed in bracket_teams and region in bracket_teams[seed]:
                teams = bracket_teams[seed][region]
                if not teams:
                    return f"({seed}) TBD"

                if len(teams) == 1:
                    # Single team at this seed - show it
                    return teams[0]
                else:
                    # Multiple teams (First Four) - check which team is NOT eliminated (the First Four winner)
                    active_teams = [t for t in teams if not lost_in_round(t, -1, teams_out)]

                    if len(active_teams) == 1:
                        # One team won (other is eliminated)
                        return active_teams[0]
                    elif len(active_teams) == 0:
                        # Both eliminated (shouldn't happen) - return first
                        return teams[0]
                    else:
                        # Both still active (First Four game not yet played)
                        return teams[0]

            return f"({seed}) TBD"

        # Future rounds - look up winners from previous round
        if round_num - 1 in round_winners:
            prev_round_winners = round_winners[round_num - 1]
            if region in prev_round_winners:
                # Calculate which matchup index this seed came from
                # This is complex, so we'll use a different approach below
                pass

        return f"({seed}) TBD"

    # Helper to determine winner of a matchup
    def get_winner(team_a, team_b, round_num):
        # If either team is TBD or None, we can't determine a winner
        if team_a is None or "TBD" in str(team_a) or team_b is None or "TBD" in str(team_b):
            return None

        team_a_loss = lost_in_round(team_a, round_num, teams_out)
        team_b_loss = lost_in_round(team_b, round_num, teams_out)

        if team_a_loss and not team_b_loss:
            return team_b
        if team_b_loss and not team_a_loss:
            return team_a
        # Neither (or both) lost in this round - can't determine a winner yet
        return None

    # Helper to check if team lost in this round
    def is_eliminated(team_name, round_num):
        return lost_in_round(team_name, round_num, teams_out)

    # Main bracket rounds
    round_names = ["Round of 64", "Round of 32", "Sweet 16", "Elite Eight", "Final Four", "Championship"]
    matchup_structure = [
        [(1, 16), (8, 9), (5, 12), (4, 13), (6, 11), (3, 14), (7, 10), (2, 15)],  # Round 1
        [(1, 8), (5, 4), (6, 3), (7, 2)],  # Round 2
        [(1, 5), (6, 7)],  # Round 3
        [(1, 6)],  # Round 4
        [(1, 1), (1, 1)],  # Round 5 (semifinal structure, placeholder)
        [(1, 1)],  # Round 6 (championship, placeholder)
    ]

    for round_num, matchups_seeds in enumerate(matchup_structure):
        round_obj = {"name": round_names[round_num], "regions": {}}
        round_winners[round_num] = {}
        round_seeds[round_num] = {}

        for region in regions:
            round_obj["regions"][region] = []
            round_winners[round_num][region] = {}
            round_seeds[round_num][region] = {}

            for matchup_idx, (seed_a, seed_b) in enumerate(matchups_seeds):
                team_a = None
                team_b = None

                # Round of 64 - get from bracket_teams
                if round_num == 0:
                    team_a = get_team_for_seed(seed_a, region, round_num)
                    team_b = get_team_for_seed(seed_b, region, round_num)
                else:
                    # Future rounds - get from previous round's winners
                    # The winner from matchup (matchup_idx * 2) plays winner from (matchup_idx * 2 + 1)
                    prev_matchup_a = matchup_idx * 2
                    prev_matchup_b = matchup_idx * 2 + 1

                    if region in round_winners[round_num - 1]:
                        prev_winners = round_winners[round_num - 1][region]
                        team_a = prev_winners.get(prev_matchup_a)
                        team_b = prev_winners.get(prev_matchup_b)

                    # If no winner found, show TBD with the best (lowest) seed that could advance
                    prev_seeds = round_seeds[round_num - 1].get(region, {})
                    if team_a is None:
                        seed_a = prev_seeds.get(prev_matchup_a, seed_a)
                        team_a = f"({seed_a}) TBD"
                    else:
                        seed_a = lookup_seed(team_a, seed_a)
                    if team_b is None:
                        seed_b = prev_seeds.get(prev_matchup_b, seed_b)
                        team_b = f"({seed_b}) TBD"
                    else:
                        seed_b = lookup_seed(team_b, seed_b)

                # Determine winner of this matchup
                winner = get_winner(team_a, team_b, round_num)
                round_winners[round_num][region][matchup_idx] = winner

                # Seed that advances from this game (best possible seed if undecided)
                if winner:
                    round_seeds[round_num][region][matchup_idx] = seed_a if winner == team_a else seed_b
                else:
                    try:
                        round_seeds[round_num][region][matchup_idx] = min(int(seed_a), int(seed_b))
                    except (TypeError, ValueError):
                        round_seeds[round_num][region][matchup_idx] = seed_a

                # Store matchup with elimination flags
                matchup_obj = {
                    "team_a": team_a,
                    "seed_a": seed_a,
                    "team_b": team_b,
                    "seed_b": seed_b,
                    "winner": winner,
                    "team_a_eliminated": is_eliminated(team_a, round_num),
                    "team_b_eliminated": is_eliminated(team_b, round_num),
                }

                round_obj["regions"][region].append(matchup_obj)

        rounds.append(round_obj)

    return rounds


def build_final_four_brackets(regional_winners, final_four_config, teams_out=None):
    """
    Build Final Four bracket with two semifinal games and one championship game.

    regional_winners: dict mapping region name to winning team (or TBD/None)
    final_four_config: dict with semifinal_1_region_1, semifinal_1_region_2, etc.
    teams_out: dict mapping team name to round lost
    """
    teams_out = teams_out or {}

    # Extract the configured region pairings
    sf1_r1 = final_four_config.get("semifinal_1_region_1", "EAST").upper()
    sf1_r2 = final_four_config.get("semifinal_1_region_2", "WEST").upper()
    sf2_r1 = final_four_config.get("semifinal_2_region_1", "MIDWEST").upper()
    sf2_r2 = final_four_config.get("semifinal_2_region_2", "SOUTH").upper()

    # Get regional winners
    team_sf1_r1 = regional_winners.get(sf1_r1, "TBD")
    team_sf1_r2 = regional_winners.get(sf1_r2, "TBD")
    team_sf2_r1 = regional_winners.get(sf2_r1, "TBD")
    team_sf2_r2 = regional_winners.get(sf2_r2, "TBD")

    # Ensure TBD is shown if None or empty
    if not team_sf1_r1 or "TBD" in str(team_sf1_r1):
        team_sf1_r1 = "TBD"
    if not team_sf1_r2 or "TBD" in str(team_sf1_r2):
        team_sf1_r2 = "TBD"
    if not team_sf2_r1 or "TBD" in str(team_sf2_r1):
        team_sf2_r1 = "TBD"
    if not team_sf2_r2 or "TBD" in str(team_sf2_r2):
        team_sf2_r2 = "TBD"

    rounds = []

    def game_winner(a, b, round_idx):
        if "TBD" in str(a) or "TBD" in str(b):
            return None
        a_lost = lost_in_round(a, round_idx, teams_out)
        b_lost = lost_in_round(b, round_idx, teams_out)
        if a_lost and not b_lost:
            return b
        if b_lost and not a_lost:
            return a
        return None

    # Semifinal Round
    semifinal_round = {
        "name": "Final Four - Semifinals",
        "regions": {
            "FINAL_FOUR": [
                {
                    "team_a": team_sf1_r1,
                    "seed_a": lookup_seed(team_sf1_r1, f"{sf1_r1[0]}"),
                    "team_b": team_sf1_r2,
                    "seed_b": lookup_seed(team_sf1_r2, f"{sf1_r2[0]}"),
                    "winner": game_winner(team_sf1_r1, team_sf1_r2, 4),
                    "team_a_eliminated": lost_in_round(team_sf1_r1, 4, teams_out),
                    "team_b_eliminated": lost_in_round(team_sf1_r2, 4, teams_out),
                },
                {
                    "team_a": team_sf2_r1,
                    "seed_a": lookup_seed(team_sf2_r1, f"{sf2_r1[0]}"),
                    "team_b": team_sf2_r2,
                    "seed_b": lookup_seed(team_sf2_r2, f"{sf2_r2[0]}"),
                    "winner": game_winner(team_sf2_r1, team_sf2_r2, 4),
                    "team_a_eliminated": lost_in_round(team_sf2_r1, 4, teams_out),
                    "team_b_eliminated": lost_in_round(team_sf2_r2, 4, teams_out),
                }
            ]
        }
    }
    rounds.append(semifinal_round)

    # Championship Round
    champ_a = game_winner(team_sf1_r1, team_sf1_r2, 4) or "TBD"
    champ_b = game_winner(team_sf2_r1, team_sf2_r2, 4) or "TBD"

    championship_round = {
        "name": "Championship",
        "regions": {
            "FINAL_FOUR": [
                {
                    "team_a": champ_a,
                    "seed_a": lookup_seed(champ_a, "SF1"),
                    "team_b": champ_b,
                    "seed_b": lookup_seed(champ_b, "SF2"),
                    "winner": game_winner(champ_a, champ_b, 5),
                    "team_a_eliminated": lost_in_round(champ_a, 5, teams_out),
                    "team_b_eliminated": lost_in_round(champ_b, 5, teams_out),
                }
            ]
        }
    }
    rounds.append(championship_round)

    return rounds


def render_regional_winner_card_html(region, winner):
    """Render a regional winner card styled like the bracket game cards."""
    known = bool(winner) and "TBD" not in str(winner)
    seed = lookup_seed(winner, "") if known else ""
    name = str(winner) if known else "TBD"
    logo = _logo_img(winner) if known else '<span class="team-logo placeholder"></span>'
    tbd_class = "" if known else "tbd"
    winner_class = "winner" if known else ""

    return f"""
    <!DOCTYPE html>
    <html>
    <head>
        <meta charset="UTF-8">
        <style>
            body {{
                font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, "Helvetica Neue", Arial, sans-serif;
                margin: 0;
                padding: 0 4px;
                height: 100vh;
                display: flex;
                align-items: center;
                justify-content: center;
                box-sizing: border-box;
                background: white;
            }}
            .matchup-container {{
                width: 100%;
                border: 1px solid #d0d0d0;
                border-radius: 4px;
                overflow: hidden;
                box-shadow: 0 1px 3px rgba(0,0,0,0.08);
                background-color: #f8f9fa;
            }}
            .game-label {{
                font-size: 10px;
                color: #888;
                background: #eef0f3;
                padding: 1px 6px;
                letter-spacing: 0.5px;
            }}
            .team-box {{
                background: #f8f9fa;
                padding: 10px 14px;
                font-size: 18px;
                display: flex;
                align-items: center;
                height: 52px;
                min-height: 52px;
            }}
            .team-info {{ display: flex; align-items: center; gap: 10px; flex: 1; }}
            .team-seed {{ font-weight: 700; color: #4472C4; min-width: 24px; font-size: 18px; }}
            .team-logo {{ width: 38px; height: 38px; object-fit: contain; flex-shrink: 0; }}
            .team-logo.placeholder {{ display: inline-block; }}
            .team-name {{ flex: 1; font-size: 18px; color: #333; }}
            .team-box.tbd {{ color: #999; font-style: italic; }}
            .team-box.tbd .team-seed, .team-box.tbd .team-name {{ color: #999; }}
            .team-box.winner .team-name {{ font-weight: 700; }}
        </style>
    </head>
    <body>
        <div class="matchup-container">
            <div class="game-label">{region} REGIONAL WINNER</div>
            <div class="team-box {tbd_class} {winner_class}">
                <div class="team-info">
                    <span class="team-seed">{seed}</span>
                    {logo}
                    <span class="team-name">{name}</span>
                </div>
            </div>
        </div>
    </body>
    </html>
    """


def render_bracket_html(rounds, region=None):
    """
    Render tournament bracket as an NCAA-style tree.
    If region is specified, only render that region.
    If region is None, render all regions.
    """
    regions_list = ["EAST", "WEST", "MIDWEST", "SOUTH", "FINAL_FOUR"]

    html = """
    <!DOCTYPE html>
    <html>
    <head>
        <meta charset="UTF-8">
        <style>
            body {
                font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, "Helvetica Neue", Arial, sans-serif;
                margin: 0;
                padding: 10px;
                background: white;
            }

            .ncaa-bracket-wrapper {
                position: relative;
                overflow-x: auto;
                background: white;
                border-radius: 4px;
                box-shadow: 0 2px 8px rgba(0,0,0,0.1);
            }

            .ncaa-bracket {
                display: flex;
                gap: 1px;
                padding: 10px;
                min-width: max-content;
            }

            .bracket-column {
                display: flex;
                flex-direction: column;
                align-items: center;
                min-width: 180px;
                padding: 0 15px;
                position: relative;
                flex: 1;
                min-height: 500px;
            }

            .column-matchups {
                flex: 1;
                width: 100%;
                display: flex;
                flex-direction: column;
                justify-content: space-around;
                align-items: center;
            }

            .column-matchups .matchup-container {
                margin: 6px 0;
            }

            .game-label {
                font-size: 10px;
                color: #888;
                background: #eef0f3;
                padding: 1px 6px;
                letter-spacing: 0.5px;
            }

            .bracket-column:last-child {
                border-right: none;
            }

            .bracket-column-title {
                text-align: center;
                font-size: 12px;
                font-weight: 600;
                color: #555;
                margin-bottom: 20px;
                text-transform: uppercase;
                letter-spacing: 0.5px;
                line-height: 1.4;
            }

            .matchup-container {
                position: relative;
                border: 1px solid #d0d0d0;
                border-radius: 4px;
                overflow: hidden;
                box-shadow: 0 1px 3px rgba(0,0,0,0.08);
                flex-shrink: 0;
                width: 100%;
                min-width: 160px;
                background-color: #f8f9fa;
            }

            .team-box {
                background: #f8f9fa;
                padding: 10px 12px;
                font-size: 13px;
                display: flex;
                justify-content: space-between;
                align-items: center;
                border-bottom: 1px solid #e8e8e8;
                transition: all 0.2s ease;
                height: 32px;
                min-height: 32px;
            }

            .team-box:last-child {
                border-bottom: none;
            }

            .team-box:hover {
                background: #eef2f7;
            }

            .matchup-container:hover {
                border-color: #d0d0d0;
                box-shadow: 0 2px 6px rgba(68, 114, 196, 0.15);
            }

            .team-info {
                display: flex;
                align-items: center;
                gap: 8px;
                flex: 1;
            }

            .team-seed {
                font-weight: 700;
                color: #4472C4;
                min-width: 20px;
                font-size: 12px;
            }

            .team-logo {
                width: 22px;
                height: 22px;
                object-fit: contain;
                flex-shrink: 0;
            }

            .team-logo.placeholder {
                display: inline-block;
            }

            .team-box.eliminated .team-logo {
                filter: grayscale(100%);
                opacity: 0.5;
            }

            .team-name {
                flex: 1;
                overflow: hidden;
                text-overflow: ellipsis;
                white-space: nowrap;
                font-size: 13px;
                color: #333;
            }

            .team-score {
                font-weight: 600;
                color: #666;
                margin-left: 12px;
                min-width: 30px;
                text-align: right;
            }

            .team-box.tbd {
                background: #f9f9f9;
                color: #999;
                font-style: italic;
            }

            .team-box.tbd .team-seed {
                color: #ccc;
            }

            .team-box.tbd .team-name {
                color: #999;
            }

            .team-box.tbd .team-score {
                color: #ccc;
            }

            .team-box.eliminated .team-name {
                color: #999;
                font-weight: 400;
            }

            .team-box.eliminated .team-seed {
                color: #999;
                font-weight: 700;
            }

            .team-box.winner .team-name {
                font-weight: 700;
            }

            .team-box.winner .team-seed {
                font-weight: 700;
            }

            .connector-line {
                position: absolute;
                height: 2px;
                background: #d0d0d0;
                top: 50%;
                transform: translateY(-50%);
                pointer-events: none;
            }

            .vertical-connector {
                position: absolute;
                width: 2px;
                background: #d0d0d0;
                left: 100%;
                pointer-events: none;
            }
        </style>
    </head>
    <body>
        <div class="ncaa-bracket-wrapper">
            <svg class="bracket-connectors" style="position: absolute; top: 0; left: 0; pointer-events: none;"></svg>
            <div class="ncaa-bracket">
    """

    # If region specified, only render that region
    regions_to_render = [region] if region else regions_list

    game_counter = 1

    # Flatten all matchups by round across specified regions
    for round_idx, round_data in enumerate(rounds):
        round_name = round_data['name']

        # Get all matchups for this round across specified regions
        all_matchups = []
        for r in regions_to_render:
            if r in round_data['regions']:
                all_matchups.extend(round_data['regions'][r])

        # Hide the First Four column when the region has no First Four games
        if round_name == "First Four" and not all_matchups:
            continue

        display_name = "Opening Round" if round_name == "First Four" else round_name

        html += f"""
                <div class="bracket-column">
                    <div class="bracket-column-title">{display_name}</div>
                    <div class="column-matchups">
        """

        # Display matchups
        for matchup_idx, matchup in enumerate(all_matchups):
            team_a = matchup['team_a']
            team_b = matchup['team_b']
            seed_a = matchup['seed_a']
            seed_b = matchup['seed_b']
            winner = matchup.get('winner')

            is_tbd_a = "TBD" in str(team_a)
            is_tbd_b = "TBD" in str(team_b)
            is_eliminated_a = matchup.get('team_a_eliminated', False)
            is_eliminated_b = matchup.get('team_b_eliminated', False)
            is_winner_a = winner == team_a
            is_winner_b = winner == team_b

            tbd_class_a = "tbd" if is_tbd_a else ""
            eliminated_class_a = "eliminated" if is_eliminated_a else ""
            winner_class_a = "winner" if is_winner_a else ""
            classes_a = f"{tbd_class_a} {eliminated_class_a} {winner_class_a}".strip()

            tbd_class_b = "tbd" if is_tbd_b else ""
            eliminated_class_b = "eliminated" if is_eliminated_b else ""
            winner_class_b = "winner" if is_winner_b else ""
            classes_b = f"{tbd_class_b} {eliminated_class_b} {winner_class_b}".strip()

            game_num = game_counter
            game_counter += 1
            feeds_to = matchup_idx // 2 if len(all_matchups) > 1 and round_name != "First Four" else -1

            html += f"""
                    <div class="matchup-container" data-round="{round_idx}" data-index="{matchup_idx}" data-feeds-to="{feeds_to}">
                        <div class="game-label">G{game_num}</div>
                        <div class="team-box {classes_a}">
                            <div class="team-info">
                                <span class="team-seed">{seed_a}</span>
                                {_logo_img(team_a)}
                                <span class="team-name">{str(team_a)[:28]}</span>
                            </div>
                            <span class="team-score"></span>
                        </div>
                        <div class="team-box {classes_b}">
                            <div class="team-info">
                                <span class="team-seed">{seed_b}</span>
                                {_logo_img(team_b)}
                                <span class="team-name">{str(team_b)[:28]}</span>
                            </div>
                            <span class="team-score"></span>
                        </div>
                    </div>
            """


        html += """
                    </div>
                </div>
        """

    html += """
            </div>
        </div>
        <script>
            function drawConnectors() {
                const svg = document.querySelector('.bracket-connectors');
                const wrapper = document.querySelector('.ncaa-bracket-wrapper');
                const columns = document.querySelectorAll('.bracket-column');

                if (columns.length < 2 || !wrapper) return;

                // Clear existing lines
                svg.innerHTML = '';

                // Set SVG size to match wrapper
                svg.style.width = '0px';
                svg.style.height = '0px';
                svg.style.width = wrapper.scrollWidth + 'px';
                svg.style.height = wrapper.scrollHeight + 'px';
                svg.setAttribute('width', wrapper.scrollWidth);
                svg.setAttribute('height', wrapper.scrollHeight);

                // Function to draw connector lines between two columns
                function connectColumns(colA, colB) {
                    const matchupsA = colA.querySelectorAll('.matchup-container');
                    const matchupsB = colB.querySelectorAll('.matchup-container');

                    // Connect each pair of feeder games to the game they feed (by data-feeds-to)
                    for (let i = 0; i < matchupsB.length; i++) {
                        const feeders = Array.from(matchupsA).filter(m => m.dataset.feedsTo === String(i));
                        const mAa = feeders[0];
                        const mAb = feeders[1];
                        const mB = matchupsB[i];

                        if (mAa && mAb && mB) {
                            // Get positions relative to wrapper
                            const rectAa = mAa.getBoundingClientRect();
                            const rectAb = mAb.getBoundingClientRect();
                            const rectB = mB.getBoundingClientRect();
                            const wrapperRect = wrapper.getBoundingClientRect();

                            const offX = wrapper.scrollLeft - wrapperRect.left;
                            const offY = wrapper.scrollTop - wrapperRect.top;

                            const yAa = rectAa.top + offY + rectAa.height / 2;
                            const yAb = rectAb.top + offY + rectAb.height / 2;
                            const yB = rectB.top + offY + rectB.height / 2;
                            const xA = rectAa.right + offX;
                            const xB = rectB.left + offX;

                            // Midpoint x coordinate
                            const xMid = xA + (xB - xA) * 0.5;

                            // Line from first game right edge to midpoint
                            const line1 = document.createElementNS('http://www.w3.org/2000/svg', 'line');
                            line1.setAttribute('x1', xA);
                            line1.setAttribute('y1', yAa);
                            line1.setAttribute('x2', xMid);
                            line1.setAttribute('y2', yAa);
                            line1.setAttribute('stroke', '#d0d0d0');
                            line1.setAttribute('stroke-width', '2');
                            svg.appendChild(line1);

                            // Line from second game right edge to midpoint
                            const line2 = document.createElementNS('http://www.w3.org/2000/svg', 'line');
                            line2.setAttribute('x1', xA);
                            line2.setAttribute('y1', yAb);
                            line2.setAttribute('x2', xMid);
                            line2.setAttribute('y2', yAb);
                            line2.setAttribute('stroke', '#d0d0d0');
                            line2.setAttribute('stroke-width', '2');
                            svg.appendChild(line2);

                            // Vertical line from first game to winner position
                            const line3 = document.createElementNS('http://www.w3.org/2000/svg', 'line');
                            line3.setAttribute('x1', xMid);
                            line3.setAttribute('y1', yAa);
                            line3.setAttribute('x2', xMid);
                            line3.setAttribute('y2', yAb);
                            line3.setAttribute('stroke', '#d0d0d0');
                            line3.setAttribute('stroke-width', '2');
                            svg.appendChild(line3);

                            // Line from midpoint to winner box
                            const line4 = document.createElementNS('http://www.w3.org/2000/svg', 'line');
                            line4.setAttribute('x1', xMid);
                            line4.setAttribute('y1', yB);
                            line4.setAttribute('x2', xB);
                            line4.setAttribute('y2', yB);
                            line4.setAttribute('stroke', '#d0d0d0');
                            line4.setAttribute('stroke-width', '2');
                            svg.appendChild(line4);
                        }
                    }
                }

                // Connect every pair of adjacent rounds (handles First Four and later rounds)
                for (let c = 0; c < columns.length - 1; c++) {
                    connectColumns(columns[c], columns[c + 1]);
                }
            }

            window.addEventListener('load', () => {
                setTimeout(drawConnectors, 100);
            });
            window.addEventListener('resize', drawConnectors);
            if (window.ResizeObserver) {
                new ResizeObserver(() => drawConnectors()).observe(document.body);
            }
        </script>
    </body>
    </html>
    """

    return html


def load_data(file_path, updated_time=None):
    workbook = openpyxl.load_workbook(
        file_path,
        read_only=True,
        data_only=True,
        keep_vba=Path(file_path).suffix.lower() == ".xlsm",
        keep_links=False,
    )

    required_sheets = [DRAFT_ORDER_SHEET, DRAFT_SHEET, TEAM_DATA_SHEET, CONFIG_SHEET]
    missing = [s for s in required_sheets if s not in workbook.sheetnames]
    if missing:
        raise ValueError("Missing required sheet(s): " + ", ".join(missing))

    # Fantasy scoring weights. Config!A2:C2 correspond to PTS, REB, AST.
    config_ws = workbook[CONFIG_SHEET]
    fantasy_weights = []
    for cell in ("A2", "B2", "C2"):
        value = config_ws[cell].value
        try:
            fantasy_weights.append(float(value))
        except (TypeError, ValueError):
            raise ValueError(
                f"Config!{cell} must contain a numeric fantasy-point value."
            )

    # Payout settings. Config!D2:G2 correspond to base payout,
    # last-place multiplier, third-place multiplier, and second-place multiplier.
    # Historical average tournament games by seed.
    # Config!I2:I17 correspond to seeds 1 through 16.
    seed_average_games = {}
    for seed_number, row_number in enumerate(range(2, 18), start=1):
        value = config_ws[f"I{row_number}"].value
        try:
            seed_average_games[seed_number] = float(value)
        except (TypeError, ValueError):
            raise ValueError(
                f"Config!I{row_number} must contain a numeric historical average games value."
            )

    # Config!J2 is the tournament FPts weighting.
    try:
        tournament_fpts_weight = float(config_ws["J2"].value)
    except (TypeError, ValueError):
        raise ValueError("Config!J2 must contain a numeric tournament FPts weight.")

    payout_config = {}
    for cell in ("D2", "E2", "F2", "G2"):
        value = config_ws[cell].value
        try:
            payout_config[cell] = float(value)
        except (TypeError, ValueError):
            raise ValueError(f"Config!{cell} must contain a numeric payout value.")

    # Final Four bracket configuration
    final_four_config = {}
    try:
        # Semifinal 1: M2 vs M3
        final_four_config["semifinal_1_region_1"] = str(config_ws["M2"].value).strip().upper() if config_ws["M2"].value else "EAST"
        final_four_config["semifinal_1_region_2"] = str(config_ws["M3"].value).strip().upper() if config_ws["M3"].value else "WEST"
        # Semifinal 2: N2 vs N3
        final_four_config["semifinal_2_region_1"] = str(config_ws["N2"].value).strip().upper() if config_ws["N2"].value else "MIDWEST"
        final_four_config["semifinal_2_region_2"] = str(config_ws["N3"].value).strip().upper() if config_ws["N3"].value else "SOUTH"
    except (TypeError, AttributeError):
        # Default configuration if not found
        final_four_config = {
            "semifinal_1_region_1": "EAST",
            "semifinal_1_region_2": "WEST",
            "semifinal_2_region_1": "MIDWEST",
            "semifinal_2_region_2": "SOUTH",
        }
    drafters = []
    seen = set()
    ws = workbook[DRAFT_ORDER_SHEET]
    for values in ws.iter_rows(
        min_row=DRAFTER_START_ROW, min_col=1, max_col=1, values_only=True
    ):
        value = values[0]
        if value is None:
            continue
        draftee = str(value).strip()
        key = draftee.casefold()
        if draftee and key not in seen:
            drafters.append(draftee)
            seen.add(key)

    # Draft rows
    ws = workbook[DRAFT_SHEET]
    draft_rows = []
    for row_number, values in enumerate(
        ws.iter_rows(min_row=2, min_col=1, max_col=21, values_only=True),
        start=2,
    ):
        draft_position, draftee, player_name = values[:3]
        if draftee is None or player_name is None:
            continue

        tournament_values = values[3:21]
        _, games_not_input = count_game_statuses(tournament_values)
        games_played = count_values_not_out(values[3:9])

        draft_rows.append({
            "draft_row": row_number,
            "draft_position": draft_position,
            "draftee": str(draftee).strip(),
            "player_name": str(player_name).strip(),
            "games_not_input": games_not_input,
            "possible_games_remaining": count_blank_values(values[3:9]),
            "games_played": games_played,
            "tourn_pts": sum_numeric_values(tournament_values[0:6]),
            "tourn_reb": sum_numeric_values(tournament_values[6:12]),
            "tourn_ast": sum_numeric_values(tournament_values[12:18]),
            "tourn_pts_rounds": tournament_values[0:6],
            "tourn_reb_rounds": tournament_values[6:12],
            "tourn_ast_rounds": tournament_values[12:18],
        })

    # Player/team data
    ws = workbook[TEAM_DATA_SHEET]
    header_values = next(
        ws.iter_rows(min_row=1, max_row=1, values_only=True), ()
    )
    headers = {
        str(value).strip(): index
        for index, value in enumerate(header_values)
        if value is not None
    }

    required_headers = ["Player", "Team", "Seed", "PPG", "RPG", "APG"]
    missing = [h for h in required_headers if h not in headers]
    if missing:
        raise ValueError("Missing TeamData header(s): " + ", ".join(missing))

    max_column = max(headers.values()) + 1
    player_stats = {}
    player_names = []

    for values in ws.iter_rows(min_row=2, max_col=max_column, values_only=True):
        def get_value(header):
            index = headers[header]
            return values[index] if index < len(values) else None

        player_name = get_value("Player")
        if player_name is None:
            continue

        display_name = str(player_name).strip()
        if not display_name:
            continue

        player_key = display_name.casefold()
        if player_key not in player_stats:
            player_names.append(display_name)

        ppg = get_value("PPG")
        rpg = get_value("RPG")
        apg = get_value("APG")
        # Season FPts are calculated from the season PPG/RPG/APG values
        # using the configurable PTS/REB/AST weights in Config!A2:C2.
        total_fantasy_points = weighted_fantasy_points(
            ppg, rpg, apg, fantasy_weights
        )

        player_stats[player_key] = {
            "Team": get_value("Team"),
            "Seed": get_value("Seed"),
            "Region": get_value("Region") if "Region" in headers else None,
            "Tot. FPts": total_fantasy_points,
            "PPG": ppg,
            "RPG": rpg,
            "APG": apg,
        }

    player_names.sort(key=str.casefold)

    # Eliminated teams
    teams_out = {}  # {team_name_lower: round_lost}
    if TEAMS_OUT_SHEET in workbook.sheetnames:
        ws = workbook[TEAMS_OUT_SHEET]
        for values in ws.iter_rows(min_row=2, min_col=1, max_col=2, values_only=True):
            team_name = values[0]
            round_lost = values[1] if len(values) > 1 else None
            if team_name is not None and str(team_name).strip():
                team_lower = str(team_name).strip().casefold()
                round_text = str(round_lost).strip() if round_lost else None
                teams_out[team_lower] = round_text

    # Workbook last-edited timestamp is maintained by the local stats updater
    # in GameResults!D2. This is the authoritative timestamp displayed by
    # the dashboard; it does not depend on OneDrive's file modified time.
    workbook_updated_time = None
    if GAME_RESULTS_SHEET not in workbook.sheetnames:
        raise ValueError(f"Missing required sheet: {GAME_RESULTS_SHEET}")

    ws = workbook[GAME_RESULTS_SHEET]
    raw_updated_time = ws["D2"].value
    # A brand-new/pre-draft workbook may not have a stats-updater timestamp yet.
    # Treat a blank D2 as "not updated yet" instead of rejecting the workbook.
    if raw_updated_time is None or str(raw_updated_time).strip() == "":
        workbook_updated_time = None
    elif isinstance(raw_updated_time, datetime):
        workbook_updated_time = raw_updated_time
    else:
        raw_updated_text = str(raw_updated_time).strip()
        try:
            workbook_updated_time = datetime.fromisoformat(
                raw_updated_text.replace("Z", "+00:00")
            )
        except ValueError:
            for date_format in (
                "%m/%d/%Y %I:%M %p",
                "%m/%d/%Y %H:%M",
                "%Y-%m-%d %H:%M:%S",
                "%Y-%m-%d %H:%M",
            ):
                try:
                    workbook_updated_time = datetime.strptime(raw_updated_text, date_format)
                    break
                except ValueError:
                    continue

        if workbook_updated_time is None:
            raise ValueError(
                "GameResults!D2 does not contain a recognizable date/time. "
                "Expected an Excel date/time value or a standard date/time string."
            )

    # A blank D2 is valid for a pre-draft workbook, so only normalize
    # timezone information when a timestamp was actually supplied.
    if workbook_updated_time is not None and workbook_updated_time.tzinfo is not None:
        workbook_updated_time = workbook_updated_time.astimezone().replace(tzinfo=None)

    # Most recent game result
    latest_game_result = {"Round": "", "Winner": "", "Loser": ""}
    if GAME_RESULTS_SHEET in workbook.sheetnames:
        ws = workbook[GAME_RESULTS_SHEET]
        headers2 = {}
        for column in range(1, ws.max_column + 1):
            value = ws.cell(row=1, column=column).value
            if value is not None:
                headers2[str(value).strip()] = column

        required = ["Round", "Winner", "Loser"]
        if all(h in headers2 for h in required):
            for row in range(ws.max_row, 1, -1):
                vals = {
                    h: ws.cell(row=row, column=headers2[h]).value
                    for h in required
                }
                if any(v not in (None, "") for v in vals.values()):
                    latest_game_result = {
                        h: "" if vals[h] is None else str(vals[h]).strip()
                        for h in required
                    }
                    break

    # Tournament bracket teams (from TeamList sheet if available)
    bracket_teams = {}  # {seed: {region: [teams]}}
    if TEAM_LIST_SHEET in workbook.sheetnames:
        ws = workbook[TEAM_LIST_SHEET]
        header_values = next(
            ws.iter_rows(min_row=1, max_row=1, values_only=True), ()
        )
        headers3 = {
            str(value).strip(): index
            for index, value in enumerate(header_values)
            if value is not None
        }

        if "Seed" in headers3 and "Region" in headers3 and "Team" in headers3:
            for values in ws.iter_rows(min_row=2, max_col=max(headers3.values()) + 1, values_only=True):
                try:
                    seed_value = values[headers3["Seed"]]
                    if seed_value is None:
                        continue

                    # Handle seed values like "11 - FF" or "16 - FF" by extracting the number
                    seed_str = str(seed_value).strip()
                    # Extract just the numeric part before any " - " or other suffix
                    if " - " in seed_str:
                        seed_str = seed_str.split(" - ")[0].strip()

                    seed = int(float(seed_str))
                    region = str(values[headers3["Region"]]).strip().upper() if values[headers3["Region"]] else None
                    team = str(values[headers3["Team"]]).strip() if values[headers3["Team"]] else None

                    if seed and region and team and 1 <= seed <= 16:
                        if seed not in bracket_teams:
                            bracket_teams[seed] = {}
                        if region not in bracket_teams[seed]:
                            bracket_teams[seed][region] = []
                        bracket_teams[seed][region].append(team)
                except (TypeError, ValueError, IndexError):
                    continue

    return {
        "fantasy_weights": fantasy_weights,
        "seed_average_games": seed_average_games,
        "tournament_fpts_weight": tournament_fpts_weight,

        "payout_config": payout_config,
        "final_four_config": final_four_config,
        "drafters": drafters,
        "draft_rows": draft_rows,
        "player_stats": player_stats,
        "player_names": player_names,
        "teams_out": teams_out,
        "latest_game_result": latest_game_result,
        "updated_time": workbook_updated_time,
        "bracket_teams": bracket_teams,
    }


def player_total(row, fantasy_weights):
    return weighted_fantasy_points(
        row.get("tourn_pts"),
        row.get("tourn_reb"),
        row.get("tourn_ast"),
        fantasy_weights,
    )


def make_team_rows(data, selected_draftee):
    rows = [
        row for row in data["draft_rows"]
        if row["draftee"].casefold() == selected_draftee.casefold()
    ]
    rows.sort(key=lambda row: draft_position_sort_key(row["draft_position"]))

    result = []
    for row in rows:
        stats = data["player_stats"].get(row["player_name"].casefold(), {})
        result.append({
            "Draft Position": row["draft_position"],
            "Player Name": row["player_name"],
            "Team": stats.get("Team"),
            "Seed": stats.get("Seed"),
            "FPts": stats.get("Tot. FPts"),
            "PPG": stats.get("PPG"),
            "RPG": stats.get("RPG"),
            "APG": stats.get("APG"),
            "Tourn. FPts": player_total(row, data["fantasy_weights"]),
            "Tourn. PTS": row.get("tourn_pts"),
            "Tourn. REB": row.get("tourn_reb"),
            "Tourn. AST": row.get("tourn_ast"),
            "_team_out": (
                stats.get("Team") is not None
                and str(stats.get("Team")).strip().casefold() in data["teams_out"]
            ),
        })
    return result


def make_standings_round_history(data):
    """Return each draftee's ranking after each tournament round."""
    round_names = [
        "First Round",
        "Second Round",
        "Sweet Sixteen",
        "Elite Eight",
        "Final Four",
        "National Title",
    ]

    # Draft columns are grouped as 6 rounds of PTS, followed by 6 REB,
    # followed by 6 AST. For each round, add PTS + REB + AST for every
    # drafted player, then accumulate the result through the tournament.
    round_totals = {
        draftee: [0.0] * len(round_names)
        for draftee in data["drafters"]
    }

    for row in data["draft_rows"]:
        owner = row["draftee"]
        pts = row.get("tourn_pts_rounds", [None] * 6)
        reb = row.get("tourn_reb_rounds", [None] * 6)
        ast = row.get("tourn_ast_rounds", [None] * 6)

        for i in range(6):
            round_total = weighted_fantasy_points(
                pts[i], reb[i], ast[i], data["fantasy_weights"]
            )
            if round_total is not None:
                round_totals.setdefault(owner, [0.0] * len(round_names))[i] += round_total

    # Only include rounds that have at least one numeric stat recorded for
    # at least one drafted player. This prevents future/unplayed rounds from
    # appearing in the progression chart before any games have been played.
    played_round_indexes = []
    for i in range(len(round_names)):
        has_numeric_value = False
        for row in data["draft_rows"]:
            for key in ("tourn_pts_rounds", "tourn_reb_rounds", "tourn_ast_rounds"):
                values = row.get(key, [None] * len(round_names))
                if i < len(values):
                    value = values[i]
                    if isinstance(value, (int, float)) and not isinstance(value, bool):
                        has_numeric_value = True
                        break
            if has_numeric_value:
                break
        if has_numeric_value:
            played_round_indexes.append(i)

    # Convert round-by-round scoring into cumulative totals, then rank
    # everyone after each played round. Rank 1 is the best (highest) total.
    result = []
    running_total = {owner: 0.0 for owner in round_totals}
    for i in played_round_indexes:
        round_name = round_names[i]
        for owner, values in round_totals.items():
            running_total[owner] += values[i]

        ordered = sorted(
            data["drafters"],
            key=lambda owner: (-running_total.get(owner, 0.0), owner.casefold()),
        )
        rankings = {owner: rank for rank, owner in enumerate(ordered, start=1)}

        row = {"Round": round_name}
        for owner in data["drafters"]:
            row[owner] = rankings.get(owner)
        result.append(row)

    return pd.DataFrame(result)


def make_standings(data):
    positions = {
        draftee.casefold(): position
        for position, draftee in enumerate(data["drafters"], start=1)
    }
    totals = {}

    for row in data["draft_rows"]:
        key = row["draftee"].casefold()
        if key not in totals:
            totals[key] = {
                "Team Owner": row["draftee"],
                "Draft Position": positions.get(key, ""),
                "Total FPts": None,
                "Active Players": 0,
                "Poss. Games Remaining": 0,
                "MVP": "",
                "MVP FPts": None,
            }

        t = totals[key]
        t["Total FPts"] = sum_existing_values([
            t["Total FPts"], player_total(row, data["fantasy_weights"])
        ])

        stats = data["player_stats"].get(row["player_name"].casefold(), {})
        team = stats.get("Team")
        if team is not None and str(team).strip().casefold() not in data["teams_out"]:
            t["Active Players"] += 1

        t["Poss. Games Remaining"] += row.get("possible_games_remaining") or 0

        pf = player_total(row, data["fantasy_weights"])
        if pf is not None and (t["MVP FPts"] is None or pf > t["MVP FPts"]):
            t["MVP"] = row["player_name"]
            t["MVP FPts"] = pf

    standings = list(totals.values())
    standings.sort(key=lambda x: (
        -(x["Total FPts"] if x["Total FPts"] is not None else 0),
        x["Draft Position"] if isinstance(x["Draft Position"], int) else 999999,
        x["Team Owner"].casefold(),
    ))
    for i, row in enumerate(standings, start=1):
        row["Rank"] = i

    # A pre-draft workbook can have no rows at all. Return an empty
    # DataFrame with the expected schema so downstream page rendering
    # does not raise a KeyError when selecting these columns.
    standings_columns = [
        "Rank", "Draft Position", "Team Owner", "Total FPts",
        "Active Players", "Poss. Games Remaining", "MVP", "MVP FPts"
    ]
    if not standings:
        return pd.DataFrame(columns=standings_columns)

    return pd.DataFrame(standings)[standings_columns]


def make_power_rankings(data):
    positions = {
        draftee.casefold(): position
        for position, draftee in enumerate(data["drafters"], start=1)
    }
    totals = {}

    for row in data["draft_rows"]:
        key = row["draftee"].casefold()
        totals.setdefault(key, {
            "Team Owner": row["draftee"],
            "Draft Position": positions.get(key, ""),
            "Total FPts": None,
            "Active Players": 0,
            "Games Played": 0,
        })

        t = totals[key]
        current_fpts = player_total(row, data["fantasy_weights"])
        t["Total FPts"] = sum_existing_values([t["Total FPts"], current_fpts])

        stats = data["player_stats"].get(row["player_name"].casefold(), {})
        team = stats.get("Team")
        if team is not None and str(team).strip().casefold() not in data["teams_out"]:
            t["Active Players"] += 1

        t["Games Played"] += row.get("games_played") or 0

    result = []
    tournament_weight = data["tournament_fpts_weight"]

    for key, t in totals.items():
        team_projected_fpts = 0.0
        expected_games_total = 0.0

        for row in data["draft_rows"]:
            if row["draftee"].casefold() != key:
                continue

            stats = data["player_stats"].get(row["player_name"].casefold(), {})
            team = stats.get("Team")
            team_is_active = (
                team is not None
                and str(team).strip().casefold() not in data["teams_out"]
            )

            current_player_fpts = float(player_total(row, data["fantasy_weights"]) or 0.0)
            games_played = row.get("games_played") or 0

            tourn_fpts_per_game = (
                current_player_fpts / games_played
                if games_played > 0 else 0.0
            )

            # "Tot. FPts" in player_stats is the season FPts per game
            # calculated from the season PPG/RPG/APG fantasy components.
            season_fpts_per_game = float(stats.get("Tot. FPts") or 0.0)

            # If the player has not recorded any tournament games yet,
            # use the season FPts/game by itself. There is no tournament
            # sample to blend with the season rate.
            if games_played == 0:
                weighted_fpts_per_game = season_fpts_per_game
            else:
                weighted_fpts_per_game = (
                    tournament_weight * tourn_fpts_per_game
                    + (1.0 - tournament_weight) * season_fpts_per_game
                )

            seed_value = stats.get("Seed")
            try:
                seed_number = int(float(seed_value))
            except (TypeError, ValueError):
                match = re.search(r"\d+", str(seed_value))
                seed_number = int(match.group(0)) if match else None

            historical_avg_games = (
                data["seed_average_games"].get(seed_number)
                if seed_number is not None else None
            )

            if not team_is_active:
                expected_remaining_games = 0.0
            elif historical_avg_games is None:
                expected_remaining_games = 1.0
            else:
                expected_remaining_games = max(
                    historical_avg_games - games_played,
                    1.0,
                )

            expected_games_total += expected_remaining_games

            team_projected_fpts += (
                current_player_fpts
                + weighted_fpts_per_game * expected_remaining_games
            )

        total_fpts = t["Total FPts"]
        fpg = (
            total_fpts / t["Games Played"]
            if total_fpts is not None and t["Games Played"] > 0
            else None
        )

        result.append({
            **t,
            "Projected Score": round(team_projected_fpts, 1),
            "FPts per game": round(fpg, 1) if fpg is not None else None,
            "Expected Games Remaining": round(expected_games_total, 1),
        })

    result.sort(key=lambda x: (
        -(x["Projected Score"] if x["Projected Score"] is not None else 0),
        x["Draft Position"] if isinstance(x["Draft Position"], int) else 999999,
        x["Team Owner"].casefold(),
    ))

    for i, row in enumerate(result, start=1):
        row["Rank"] = i

    power_columns = [
        "Rank", "Draft Position", "Team Owner", "Projected Score",
        "Total FPts", "Active Players", "FPts per game",
        "Games Played", "Expected Games Remaining"
    ]
    if not result:
        return pd.DataFrame(columns=power_columns)

    return pd.DataFrame(result)[power_columns]


def make_player_rankings(data):
    result = []
    for row in data["draft_rows"]:
        stats = data["player_stats"].get(row["player_name"].casefold(), {})
        team = stats.get("Team")
        result.append({
            "Player Name": row["player_name"],
            "Team": team,
            "Total FPts": player_total(row, data["fantasy_weights"]),
            "Team Owner": row["draftee"],
            "Games Played": row.get("games_played", 0),
            "Poss. Games Remaining": row.get("possible_games_remaining", 0),
            "_team_out": (
                team is not None
                and str(team).strip().casefold() in data["teams_out"]
            ),
            "_draft_position": row["draft_position"],
        })

    result.sort(key=lambda p: (
        -(p["Total FPts"] if p["Total FPts"] is not None else 0),
        draft_position_sort_key(p["_draft_position"]),
    ))
    result = result[:10]
    for i, row in enumerate(result, start=1):
        row["Rank"] = i

    player_columns = [
        "Rank", "Player Name", "Team", "Total FPts",
        "Team Owner", "Games Played", "Poss. Games Remaining",
        "_team_out",
    ]
    if not result:
        return pd.DataFrame(columns=player_columns), result

    return pd.DataFrame(result)[player_columns], result


def make_performers(data, positive=True):
    result = []
    for row in data["draft_rows"]:
        total = player_total(row, data["fantasy_weights"])
        games = row.get("games_played", 0)
        if games <= 0 or total is None:
            continue

        stats = data["player_stats"].get(row["player_name"].casefold(), {})
        season_avg = sum_existing_values([stats.get("Tot. FPts")])
        if season_avg is None:
            continue

        tournament_avg = total / games
        delta = tournament_avg - season_avg

        if positive and delta <= 0:
            continue
        if not positive and delta >= 0:
            continue

        team = stats.get("Team")
        result.append({
            "Player Name": row["player_name"],
            "Team": team,
            "Avg. FPts/Tourn. Game": round(tournament_avg, 1),
            "Delta From Average": round(delta, 1),
            "Team Owner": row["draftee"],
            "Games Played": games,
            "Poss. Games Remaining": row.get("possible_games_remaining", 0),
            "_team_out": (
                team is not None
                and str(team).strip().casefold() in data["teams_out"]
            ),
            "_draft_position": row["draft_position"],
        })

    result.sort(key=lambda p: (
        -p["Delta From Average"] if positive else p["Delta From Average"],
        draft_position_sort_key(p["_draft_position"]),
    ))
    result = result[:5]
    for i, row in enumerate(result, start=1):
        row["Rank"] = i

    performer_columns = [
        "Rank", "Player Name", "Team", "Avg. FPts/Tourn. Game",
        "Delta From Average", "Team Owner", "Games Played",
        "Poss. Games Remaining", "_team_out",
    ]
    if not result:
        return pd.DataFrame(columns=performer_columns), result

    return pd.DataFrame(result)[performer_columns], result


def make_draft_ratings(data):
    positions = {
        draftee.casefold(): position
        for position, draftee in enumerate(data["drafters"], start=1)
    }
    totals = {
        d.casefold(): {
            "Team Owner": d,
            "Draft Position": positions.get(d.casefold(), ""),
            "Regions": set(),
            "Projected FPts": None,
        }
        for d in data["drafters"]
    }

    for row in data["draft_rows"]:
        key = row["draftee"].casefold()
        totals.setdefault(key, {
            "Team Owner": row["draftee"],
            "Draft Position": positions.get(key, ""),
            "Regions": set(),
            "Projected FPts": None,
        })

        stats = data["player_stats"].get(row["player_name"].casefold(), {})
        avg_games = average_games_for_seed(stats.get("Seed"))
        season_avg = sum_existing_values([stats.get("Tot. FPts")])

        region = stats.get("Region")
        if region:
            totals[key]["Regions"].add(str(region).strip().casefold())

        if avg_games is None or season_avg is None:
            continue

        projected = avg_games * season_avg
        totals[key]["Projected FPts"] = sum_existing_values([
            totals[key]["Projected FPts"], projected
        ])

    result = [{
        "Team Owner": t["Team Owner"],
        "Draft Position": t["Draft Position"],
        "Projected FPts": t["Projected FPts"],
        "Regions": len(t["Regions"]),
    } for t in totals.values()]

    result.sort(key=lambda x: (
        -(x["Projected FPts"] if x["Projected FPts"] is not None else 0),
        x["Draft Position"] if isinstance(x["Draft Position"], int) else 999999,
        x["Team Owner"].casefold(),
    ))
    for i, row in enumerate(result, start=1):
        row["Rank"] = i

    ratings_columns = [
        "Rank", "Draft Position", "Team Owner", "Projected FPts", "Regions"
    ]
    if not result:
        return pd.DataFrame(columns=ratings_columns)

    return pd.DataFrame(result)[ratings_columns]


# ---------------------------------------------------------------------
# Streamlit presentation helpers
# ---------------------------------------------------------------------
def style_table(df, scale_columns=None, reverse_columns=None, medal_rank=True, pink_last=False, rank_rwg=False, performance_gradient=None, column_colors=None):
    scale_columns = scale_columns or {}
    reverse_columns = reverse_columns or {}

    # `_team_out` is an internal helper used only to identify eliminated
    # players. Remove it before creating the Styler so Streamlit cannot
    # render it as a visible column, while retaining the flags for styling.
    team_out_flags = None
    if "_team_out" in df.columns:
        team_out_flags = df["_team_out"].copy()
        df = df.drop(columns=["_team_out"])

    def apply_styles(row):
        styles = [""] * len(row)
        for col, (minimum, maximum) in scale_columns.items():
            if col in row.index:
                value = row[col]
                bg = color_scale(
                    value, minimum, maximum,
                    reverse=reverse_columns.get(col, False)
                )
                styles[row.index.get_loc(col)] = f"background-color: {bg}"
        # Optional performance gradient applied ONLY to the Rank column.
        if performance_gradient in ("green", "red") and "Rank" in row.index:
            rank_values = pd.to_numeric(df["Rank"], errors="coerce")
            rank = pd.to_numeric(row["Rank"], errors="coerce")
            if pd.notna(rank) and rank_values.notna().any():
                lo = rank_values.min()
                hi = rank_values.max()
                if hi > lo:
                    # For both tables, Rank #1 gets the strongest color.
                    ratio = (hi - rank) / (hi - lo)
                    if performance_gradient == "green":
                        start = (235, 247, 238)
                        end = (112, 173, 71)
                    else:
                        start = (252, 235, 235)
                        end = (192, 80, 80)

                    rgb = tuple(
                        round(start[i] + (end[i] - start[i]) * ratio)
                        for i in range(3)
                    )
                    rank_color = "#%02X%02X%02X" % rgb
                    styles[row.index.get_loc("Rank")] = (
                        f"background-color: {rank_color}; color: #000000;"
                    )

        if rank_rwg and "Rank" in row.index:
            rank_values = pd.to_numeric(df["Rank"], errors="coerce")
            rank = pd.to_numeric(row["Rank"], errors="coerce")
            if pd.notna(rank) and rank_values.notna().any():
                lo = rank_values.min()
                hi = rank_values.max()
                if hi > lo:
                    midpoint = (lo + hi) / 2
                    if rank <= midpoint:
                        ratio = (rank - lo) / max(midpoint - lo, 1)
                        r = int(round(0 + 255 * ratio))
                        g = int(round(176 + 79 * ratio))
                        b = int(round(80 + 175 * ratio))
                    else:
                        ratio = (rank - midpoint) / max(hi - midpoint, 1)
                        r = 255
                        g = int(round(255 - 120 * ratio))
                        b = int(round(255 - 120 * ratio))
                    styles[row.index.get_loc("Rank")] = (
                        f"background-color: rgb({r}, {g}, {b})"
                    )

        if medal_rank and "Rank" in row.index and not rank_rwg:
            rank = row["Rank"]
            if pink_last and pd.notna(rank) and rank == df["Rank"].max():
                bg = "#FF87F7"
            else:
                bg = {
                    1: "#FFD700",
                    2: "#C0C0C0",
                    3: "#CD7F32",
                }.get(rank, "#FFFFFF")
            styles[row.index.get_loc("Rank")] = f"background-color: {bg}"

        # Performer gradient must be applied last so it overrides the
        # standard gold/silver/bronze rank colors. It affects ONLY Rank.
        if performance_gradient in ("green", "red") and "Rank" in row.index:
            rank_values = pd.to_numeric(df["Rank"], errors="coerce")
            rank = pd.to_numeric(row["Rank"], errors="coerce")
            if pd.notna(rank) and rank_values.notna().any():
                lo = rank_values.min()
                hi = rank_values.max()
                if hi > lo:
                    ratio = (hi - rank) / (hi - lo)
                    if performance_gradient == "green":
                        start = (235, 247, 238)
                        end = (112, 173, 71)
                    else:
                        start = (252, 235, 235)
                        end = (192, 80, 80)

                    rgb = tuple(
                        round(start[i] + (end[i] - start[i]) * ratio)
                        for i in range(3)
                    )
                    rank_color = "#%02X%02X%02X" % rgb
                    styles[row.index.get_loc("Rank")] = (
                        f"background-color: {rank_color}; color: #000000;"
                    )

        # Eliminated players: bold red text across the entire player row.
        # `_team_out` is kept outside the displayed dataframe.
        if team_out_flags is not None:
            try:
                is_out = bool(team_out_flags.loc[row.name])
            except (KeyError, TypeError):
                is_out = False
            if is_out:
                for idx in range(len(row)):
                    styles[idx] += "; color: #C00000; font-weight: 700"

        return styles

    styled = df.style.apply(apply_styles, axis=1)

    # Display numeric values cleanly. Calculations remain unchanged.
    # One-decimal fields show exactly one decimal; integer fields do not
    # expose pandas' trailing floating-point zeros.
    one_decimal_columns = {
        "FPts per game",
        "Avg. FPts/Tourn. Game",
        "Delta From Average",
        "Delta From Season Avg.",
        "Projected FPts",
        "Projected Score",
        "Total FPts",
        "MVP FPts",
    }
    format_map = {}
    for col in df.columns:
        if col in one_decimal_columns:
            format_map[col] = lambda value: "" if pd.isna(value) else f"{float(value):.1f}"
        elif pd.api.types.is_float_dtype(df[col]):
            format_map[col] = lambda value: "" if pd.isna(value) else f"{float(value):g}"

    if format_map:
        styled = styled.format(format_map)

    # Display positive performer deltas with an explicit "+" sign.
    if "Delta From Season Avg." in df.columns:
        styled = styled.format({
            "Delta From Season Avg.": lambda x: (
                f"+{x:.1f}" if pd.notna(x) and x > 0
                else (f"{x:.1f}" if pd.notna(x) else "")
            )
        })

    # Center all table headers and cell contents while preserving the
    # existing conditional formatting/background colors.
    styled = styled.set_properties(**{"text-align": "center"})
    styled = styled.set_table_styles([
        {"selector": "th", "props": [("text-align", "center")]}
    ])

    # Optional column-specific background colors. These are applied after the
    # row/scale styling so the requested Team View stat groups are actually
    # visible, while leaving the existing scale styling intact elsewhere.
    if column_colors:
        for col, bg in column_colors.items():
            if col in df.columns:
                styled = styled.set_properties(
                    subset=[col],
                    **{"background-color": bg}
                )

    return styled


def display_df(df, height=420, scale_columns=None, reverse_columns=None, hide_index=True, rank_rwg=False, pink_last=False, performance_gradient=None, column_colors=None):
    if df.empty:
        st.info("No data available yet.")
        return

    styled = style_table(
        df,
        scale_columns=scale_columns,
        reverse_columns=reverse_columns,
        pink_last=pink_last,
        rank_rwg=rank_rwg,
        performance_gradient=performance_gradient,
        column_colors=column_colors,
    )

    # Size the table to its actual row count so short tables do not
    # show a large block of blank cells underneath the data.
    row_height = 35
    header_height = 38
    compact_height = header_height + (len(df) * row_height) + 4

    # Give every supplied row enough room to render without a vertical
    # scrollbar. The height argument is retained only as a safety cap for
    # unusually large tables.
    actual_height = compact_height if compact_height <= height else height

    st.dataframe(
        styled,
        use_container_width=True,
        hide_index=hide_index,
        height=actual_height,
    )


# ---------------------------------------------------------------------
# Public OneDrive helpers
# ---------------------------------------------------------------------
def get_onedrive_share_url():
    """Read the public OneDrive sharing URL from Streamlit Secrets."""
    try:
        if "onedrive" in st.secrets:
            url = st.secrets["onedrive"].get("share_url")
        else:
            url = st.secrets.get("ONEDRIVE_SHARE_URL")
    except (KeyError, AttributeError, FileNotFoundError):
        # No secrets file found - return None to use local mode
        return None

    if not url:
        return None

    url = str(url).strip()
    return url or None


def download_onedrive_workbook_from_share(share_url):
    """Download an anonymously shared OneDrive workbook.

    The OneDrive link is an ``Anyone with the link can view`` link, so no
    Microsoft login or Microsoft Graph token is required.  OneDrive can
    return an HTML viewer page for a normal share URL; adding ``download=1``
    asks OneDrive for the underlying file instead.
    """
    from urllib.parse import parse_qsl, urlencode, urlsplit, urlunsplit

    parts = urlsplit(share_url)
    query = dict(parse_qsl(parts.query, keep_blank_values=True))
    query["download"] = "1"
    download_url = urlunsplit(
        (parts.scheme, parts.netloc, parts.path, urlencode(query), parts.fragment)
    )

    response = requests.get(
        download_url,
        timeout=90,
        allow_redirects=True,
        headers={
            "User-Agent": "Mozilla/5.0 (Streamlit OneDrive workbook reader)",
        },
    )

    if not response.ok:
        raise RuntimeError(
            f"OneDrive returned HTTP {response.status_code} while downloading the shared workbook."
        )

    content = response.content
    content_type = response.headers.get("Content-Type", "unknown")

    # .xlsx files are ZIP containers and therefore begin with PK.
    if not content.startswith(b"PK"):
        # A public link should be anonymously downloadable. If OneDrive
        # instead returned an HTML page, surface a useful message rather
        # than feeding the HTML to openpyxl.
        preview = content[:200].decode("utf-8", errors="ignore").replace("\n", " ").strip()
        raise RuntimeError(
            "The OneDrive share link did not return the Excel workbook. "
            "Make sure the file is shared as 'Anyone with the link can view' "
            "and that downloads are allowed. "
            f"Content-Type: {content_type}."
            + (f" Response began with: {preview[:120]}" if preview else "")
        )

    temp_path = Path(".streamlit_onedrive_workbook.xlsx")
    temp_path.write_bytes(content)

    # The workbook itself contains the authoritative last-edited timestamp
    # in GameResults!D2. load_data() reads that cell after the download.
    return temp_path, None


# ---------------------------------------------------------------------
# App
# ---------------------------------------------------------------------
st.set_page_config(
    page_title="MM Player Pool Dashboard",
    page_icon="🏀",
    layout="wide",
    initial_sidebar_state="collapsed",
)

st.markdown(
    """
    <style>
    /* Hide Streamlit's top app toolbar/header and footer for a cleaner
       dashboard, especially on mobile. */
    header[data-testid="stHeader"],
    div[data-testid="stToolbar"],
    #MainMenu,
    footer {
        display: none !important;
    }

    .block-container {padding-top: 1.2rem; padding-bottom: 2rem;}

    /* Enhanced metric card styling with colored left border and background */
    div[data-testid="stMetric"] {
        background-color: #f8f9fa;
        border: 1px solid #e0e0e0;
        border-left: 4px solid #4472C4;
        border-radius: 8px;
        padding: 10px;
        box-shadow: 0 2px 4px rgba(68, 114, 196, 0.08);
        transition: box-shadow 0.3s ease;
    }

    div[data-testid="stMetric"]:hover {
        box-shadow: 0 4px 8px rgba(68, 114, 196, 0.15);
    }

    /* Enhanced title styling */
    h1 {
        color: #2E5C8A;
        border-bottom: 3px solid #4472C4;
        padding-bottom: 0.5rem;
    }

    /* Enhanced header styling with accent colors */
    h2 {
        color: #2E5C8A;
        border-bottom: 2px solid #1B9E9A;
        padding-bottom: 0.3rem;
    }

    h3 {
        color: #2E5C8A;
    }
    </style>
    """,
    unsafe_allow_html=True,
)

st.title("🏀 2027 March Madness Player Pool Dashboard")

# Production mode: read the workbook directly from the public OneDrive
# sharing link. Visitors do NOT need to sign in with Microsoft.
#
# If the share URL is not configured, retain the local uploader flow so the
# app remains convenient to run during development.
workbook_path = Path(DEFAULT_WORKBOOK)
cloud_updated_time = None
onedrive_share_url = get_onedrive_share_url()

if onedrive_share_url:
    try:
        workbook_path, cloud_updated_time = download_onedrive_workbook_from_share(
            onedrive_share_url
        )
    except Exception as exc:
        st.error(f"Could not load the OneDrive workbook: {exc}")
        st.stop()
else:
    with st.sidebar:
        st.header("Data Source")
        uploaded = st.file_uploader(
            "Upload an Excel workbook",
            type=["xlsx", "xlsm"],
            help="For local testing, upload either the empty or populated example workbook.",
        )
        st.caption("Local development mode")

    if uploaded is not None:
        temp_path = Path(".streamlit_uploaded_workbook.xlsx")
        temp_path.write_bytes(uploaded.getvalue())
        workbook_path = temp_path

    if not workbook_path.exists():
        st.warning(
            f"Put `{DEFAULT_WORKBOOK}` in the same folder as `dashboard.py`, "
            "or upload an Excel workbook using the sidebar."
        )
        st.stop()

try:
    data = load_data(workbook_path, updated_time=cloud_updated_time)
except Exception as exc:
    st.error(f"Could not load the workbook: {exc}")
    st.stop()

latest = data["latest_game_result"]

if data["updated_time"] is None:
    st.caption("Standings Updated: Not yet updated")
else:
    st.caption(
        "Standings Updated: "
        + data["updated_time"].strftime("%m/%d/%Y %I:%M %p")
    )

standings = make_standings(data)
power = make_power_rankings(data)
players, player_rows = make_player_rankings(data)
over, over_rows = make_performers(data, True)
under, under_rows = make_performers(data, False)
ratings = make_draft_ratings(data)

# Top-line summary
c1, c2, c3, c4 = st.columns(4)
c1.metric("Pool Members", len(data["drafters"]))
c2.metric("Drafted Players", len(data["draft_rows"]))
c3.metric("Players Active", int(standings["Active Players"].sum()) if not standings.empty else 0)
c4.metric("Teams Eliminated", len(data["teams_out"]))

tab0, tab1, tab2, tab3, tab4, tab5 = st.tabs([
    "🏠 Pool Overview",
    "🏆 Standings",
    "📊 Player Statistics",
    "👥 Team View",
    "🎯 Draft Ratings",
    "⛹️‍♂️ Tournament Bracket",
])

with tab0:
    st.subheader("Pool Overview")
    st.caption("A quick look at the current state of the pool.")

    # Current leaders and headline stats. These are all derived from the
    # same standings/power-ranking data used elsewhere in the dashboard.
    # After the draft, standings rows exist even before any game stats have
    # been entered. In that state there is no meaningful leader/MVP/projected
    # leader yet, so keep these cards blank rather than attempting to format
    # None as a number.
    leader = None
    if not standings.empty and standings["Total FPts"].notna().any():
        leader_candidates = standings[standings["Total FPts"].notna()]
        if not leader_candidates.empty:
            leader = leader_candidates.iloc[0]

    mvp_row = None
    if not standings.empty:
        mvp_candidates = standings.dropna(subset=["MVP FPts"])
        if not mvp_candidates.empty:
            mvp_row = mvp_candidates.loc[mvp_candidates["MVP FPts"].idxmax()]

    projected_leader = None
    if not power.empty and power["Projected Score"].notna().any():
        projected_candidates = power[power["Projected Score"].notna()]
        if not projected_candidates.empty:
            projected_leader = projected_candidates.iloc[0]
    active_total = int(standings["Active Players"].sum()) if not standings.empty else 0

    overview_history = make_standings_round_history(data)
    biggest_riser_name = "—"
    biggest_riser_text = "No completed rounds yet"
    biggest_faller_name = "—"
    biggest_faller_text = "No completed rounds yet"
    if not overview_history.empty and len(overview_history) >= 2:
        # Once a tournament round has been completed, compare the latest
        # round's standings rank with the previous completed round.
        previous = overview_history.iloc[-2]
        current = overview_history.iloc[-1]
        changes = []
        for owner in data["drafters"]:
            old_rank = previous.get(owner)
            new_rank = current.get(owner)
            if pd.notna(old_rank) and pd.notna(new_rank):
                changes.append((float(old_rank) - float(new_rank), owner))
        if changes:
            risers = [item for item in changes if item[0] > 0]
            fallers = [item for item in changes if item[0] < 0]
            if risers:
                movement = max(item[0] for item in risers)
                tied_risers = [
                    owner for change, owner in risers
                    if change == movement
                ]
                tied_risers.sort(key=str.casefold)
                biggest_riser_name = ", ".join(tied_risers)
                biggest_riser_text = f"+{int(movement)} spot{'s' if movement != 1 else ''}"
            else:
                biggest_riser_name = "No riser"
                biggest_riser_text = "No one moved up"
            if fallers:
                movement = min(item[0] for item in fallers)
                movement_down = abs(movement)
                tied_fallers = [
                    owner for change, owner in fallers
                    if change == movement
                ]
                tied_fallers.sort(key=str.casefold)
                biggest_faller_name = ", ".join(tied_fallers)
                biggest_faller_text = f"-{int(movement_down)} spot{'s' if movement_down != 1 else ''}"
            else:
                biggest_faller_name = "No faller"
                biggest_faller_text = "No one moved down"
    else:
        # Before the first round is complete, use the pre-tournament Draft
        # Ratings rank as the baseline and compare it with the current Power
        # Rankings rank. Positive movement means a team is ranked higher now.
        draft_rank_by_owner = {}
        if not ratings.empty and {"Team Owner", "Rank"}.issubset(ratings.columns):
            draft_rank_by_owner = {
                str(row["Team Owner"]).casefold(): row["Rank"]
                for _, row in ratings.iterrows()
            }

        changes = []
        if not power.empty and {"Team Owner", "Rank"}.issubset(power.columns):
            for _, row in power.iterrows():
                owner = str(row["Team Owner"])
                draft_rank = draft_rank_by_owner.get(owner.casefold())
                current_rank = row["Rank"]
                if pd.notna(draft_rank) and pd.notna(current_rank):
                    changes.append((float(draft_rank) - float(current_rank), owner))

        if changes:
            risers = [item for item in changes if item[0] > 0]
            fallers = [item for item in changes if item[0] < 0]
            if risers:
                movement = max(item[0] for item in risers)
                tied_risers = sorted(
                    [owner for change, owner in risers if change == movement],
                    key=str.casefold,
                )
                biggest_riser_name = ", ".join(tied_risers)
                biggest_riser_text = f"+{int(movement)} spot{'s' if movement != 1 else ''} vs. draft rating"
            else:
                biggest_riser_name = "No riser"
                biggest_riser_text = "No one ranked higher than draft rating"

            if fallers:
                movement = min(item[0] for item in fallers)
                movement_down = abs(movement)
                tied_fallers = sorted(
                    [owner for change, owner in fallers if change == movement],
                    key=str.casefold,
                )
                biggest_faller_name = ", ".join(tied_fallers)
                biggest_faller_text = f"- {int(movement_down)} spot{'s' if movement_down != 1 else ''} vs. draft rating"
            else:
                biggest_faller_name = "No faller"
                biggest_faller_text = "No one ranked lower than draft rating"

    c1, c2, c3, c4 = st.columns(4)
    if leader is not None:
        c1.metric("🏆 Current Leader", str(leader["Team Owner"]), f"{float(leader['Total FPts']):.1f} FPts")
    else:
        c1.metric("🏆 Current Leader", "—")
    if projected_leader is not None:
        c2.metric("🔮 Projected Leader", str(projected_leader["Team Owner"]), f"{float(projected_leader['Projected Score']):.1f} Proj. FPts")
    else:
        c2.metric("🔮 Projected Leader", "—")
    if mvp_row is not None:
        c3.metric("⭐ Pool MVP", str(mvp_row["MVP"]), f"{float(mvp_row['MVP FPts']):.1f} FPts")
    else:
        c3.metric("⭐ Pool MVP", "—")
    if latest["Winner"] and latest["Loser"]:
        # Use Streamlit's native metric card so this card has the exact same
        # height, padding, and responsive behavior as the other overview cards.
        # A leading minus sign makes the delta render with the red down arrow;
        # Streamlit strips the sign from the displayed text.
        c4.metric(
            "🏀 Most Recent Game Input",
            latest["Winner"],
            f"-- def. {latest['Loser']} [ {latest['Round']} ] --",
        )
    else:
        c4.metric("🏀 Most Recent Game", "No game recorded")

    st.markdown("### 💸 Current Payouts 💰")
    if not standings.empty:
        payout_rows = []

        payout_base = data["payout_config"]["D2"]
        last_payout = payout_base * data["payout_config"]["E2"]
        third_payout = payout_base * data["payout_config"]["F2"]
        second_payout = payout_base * data["payout_config"]["G2"]
        first_payout = (
            payout_base * len(data["drafters"])
            - last_payout
            - third_payout
            - second_payout
        )
        payout_by_rank = {
            1: first_payout,
            2: second_payout,
            3: third_payout,
            len(standings): last_payout,
        }

        # Show the top three teams plus the team currently in last place.
        payout_indices = list(range(min(3, len(standings))))
        if len(standings) > 3:
            payout_indices.append(len(standings) - 1)

        rank_emojis = {1: "🥇", 2: "🥈", 3: "🥉", (len(standings)):"🦨"}
        for index in payout_indices:
            row = standings.iloc[index]
            rank = int(row["Rank"])
            owner = row["Team Owner"]
            fpts = row["Total FPts"]

            change_text = "—"
            if not overview_history.empty and len(overview_history) >= 2:
                previous = overview_history.iloc[-2]
                current = overview_history.iloc[-1]
                old_rank = previous.get(owner)
                new_rank = current.get(owner)
                if pd.notna(old_rank) and pd.notna(new_rank):
                    movement = float(old_rank) - float(new_rank)
                    if movement > 0:
                        change_text = f"↑ {int(movement)}"
                    elif movement < 0:
                        change_text = f"↓ {int(abs(movement))}"

            payout_rows.append({
                "Rank": f"{rank_emojis.get(rank, rank)} {rank}",
                "Payout": f"${payout_by_rank.get(rank, 0):,.0f}",
                "Team Owner": owner,
                "FPts": f"{float(fpts):.1f}" if pd.notna(fpts) else "",
                "Change": change_text,
            })

        payout_display = pd.DataFrame(payout_rows)
        display_df(payout_display, height=210)
    # Biggest riser / faller / highest team round / highest individual game
    c5, c6, c7, c8 = st.columns(4)
    if biggest_riser_name != "—":
        c5.metric("📈 Biggest Riser This Round", biggest_riser_name, biggest_riser_text)
    else:
        c5.metric("📈 Biggest Riser This Round", "—", biggest_riser_text)

    if biggest_faller_name != "—":
        c6.metric("📉 Biggest Faller This Round", biggest_faller_name, biggest_faller_text)
    else:
        c6.metric("📉 Biggest Faller This Round", "—", biggest_faller_text)

    # Highest single-round team score: sum the FPts from all five players
    # on each draftee's team for every tournament round.
    highest_round_name = "—"
    highest_round_text = "No round scores recorded yet"
    round_names = [
        "First Round",
        "Second Round",
        "Sweet Sixteen",
        "Elite Eight",
        "Final Four",
        "National Title",
    ]

    round_totals = {
        owner: [0.0] * len(round_names)
        for owner in data["drafters"]
    }

    for row in data["draft_rows"]:
        owner = row["draftee"]
        pts_rounds = row.get("tourn_pts_rounds", [None] * 6)
        reb_rounds = row.get("tourn_reb_rounds", [None] * 6)
        ast_rounds = row.get("tourn_ast_rounds", [None] * 6)

        for i in range(len(round_names)):
            round_fpts = weighted_fantasy_points(
                pts_rounds[i],
                reb_rounds[i],
                ast_rounds[i],
                data["fantasy_weights"],
            )
            if round_fpts is not None:
                round_totals.setdefault(
                    owner, [0.0] * len(round_names)
                )[i] += float(round_fpts)

    round_scores = [
        (score, owner, round_names[i])
        for owner, scores in round_totals.items()
        for i, score in enumerate(scores)
        if score > 0
    ]

    if round_scores:
        highest_score, highest_owner, highest_round = max(
            round_scores, key=lambda x: (x[0], x[1].casefold())
        )
        highest_round_name = highest_owner
        highest_round_text = f"{highest_score:.1f} FPts [ {highest_round} ]"

    c7.metric("🧨 Highest Single-Round Score", highest_round_name, highest_round_text)

    # Highest single-game individual performance. Each player's tournament
    # round stat line represents that player's performance in that round/game.
    highest_game_player = "—"
    highest_game_text = "No game scores recorded yet"
    individual_game_scores = []

    for row in data["draft_rows"]:
        pts_rounds = row.get("tourn_pts_rounds", [None] * 6)
        reb_rounds = row.get("tourn_reb_rounds", [None] * 6)
        ast_rounds = row.get("tourn_ast_rounds", [None] * 6)

        for i, round_name in enumerate(round_names):
            game_fpts = weighted_fantasy_points(
                pts_rounds[i],
                reb_rounds[i],
                ast_rounds[i],
                data["fantasy_weights"],
            )
            if game_fpts is not None:
                individual_game_scores.append(
                    (float(game_fpts), row["player_name"], round_name)
                )

    if individual_game_scores:
        highest_game_score, highest_game_player_name, highest_game_round = max(
            individual_game_scores,
            key=lambda x: (x[0], x[1].casefold())
        )
        highest_game_player = highest_game_player_name
        highest_game_text = f"{highest_game_score:.1f} FPts [ {highest_game_round} ]"

    c8.metric(
        "🔥 Highest Single-Game Performance",
        highest_game_player,
        highest_game_text,
    )


    st.subheader("🚀 Standings Progression")
    st.caption("League ranking after each tournament round.")
    standings_history = make_standings_round_history(data)
    if not standings_history.empty:
        # Plot ranking rather than points so the chart is much easier to read.
        # Reverse the Y axis so Rank 1 appears at the top of the chart.
        chart_data = standings_history.melt(
            id_vars="Round",
            var_name="Team Owner",
            value_name="Rank",
        )
        chart_data["Round"] = pd.Categorical(
            chart_data["Round"],
            categories=[
                "First Round",
                "Second Round",
                "Sweet Sixteen",
                "Elite Eight",
                "Final Four",
                "National Title",
            ],
            ordered=True,
        )
        try:
            import altair as alt
            round_order = [
                "First Round", "Second Round", "Sweet Sixteen",
                "Elite Eight", "Final Four", "National Title"
            ]
            owner_order = list(data["drafters"])
            rank_max = max(len(owner_order), int(chart_data["Rank"].max()))

            # Generate exactly one deterministic, unique color per pool member.
            # The palette is sized from the actual number of pool members, so
            # it automatically expands/shrinks with the pool.
            def make_owner_colors(owners):
                count = len(owners)
                if count == 0:
                    return []

                colors = []
                used = set()

                for i in range(count):
                    # Evenly distribute hues around the color wheel.
                    hue = i / count
                    red, green, blue = colorsys.hsv_to_rgb(hue, 0.72, 0.82)
                    color = "#{:02X}{:02X}{:02X}".format(
                        round(red * 255),
                        round(green * 255),
                        round(blue * 255),
                    )

                    # Guard against an accidental RGB rounding collision.
                    # This keeps the guarantee of one unique hex color per owner.
                    if color in used:
                        step = 1 / max(count * 1000, 1)
                        attempt = 1
                        while color in used:
                            adjusted_hue = (hue + step * attempt) % 1.0
                            red, green, blue = colorsys.hsv_to_rgb(
                                adjusted_hue, 0.72, 0.82
                            )
                            color = "#{:02X}{:02X}{:02X}".format(
                                round(red * 255),
                                round(green * 255),
                                round(blue * 255),
                            )
                            attempt += 1

                    used.add(color)
                    colors.append(color)

                return colors

            owner_colors = make_owner_colors(owner_order)

            base = alt.Chart(chart_data).encode(
                x=alt.X(
                    "Round:N",
                    sort=round_order,
                    title=None,
                ),
                y=alt.Y(
                    "Rank:Q",
                    scale=alt.Scale(domain=[1, rank_max], reverse=True),
                    title=None,
                    axis=alt.Axis(values=list(range(1, rank_max + 1))),
                ),
                color=alt.Color(
                    "Team Owner:N",
                    title=None,
                    sort=owner_order,
                    legend=None,
                    scale=alt.Scale(domain=owner_order, range=owner_colors),
                ),
                tooltip=["Round:N", "Team Owner:N", "Rank:Q"],
            )

            lines = base.mark_line()
            points = base.mark_point(filled=True, size=55)

            # Put a small rounded label behind every point so the owner's name
            # is visible without having to hover. The label background uses
            # the same series color as the corresponding line, while the
            # owner name is always white and bold for maximum contrast.
            label_bg = base.mark_rect(
                width=54,
                height=18,
                cornerRadius=5,
                opacity=0.92,
                strokeWidth=0,
            )
            labels = base.mark_text(
                fontSize=9,
                fontWeight="bold",
                baseline="middle",
            ).encode(
                text="Team Owner:N",
                color=alt.value("#FFFFFF"),
            )

            chart = alt.layer(lines, points, label_bg, labels).properties(height=500)
            st.altair_chart(chart, use_container_width=True)
        except ImportError:
            # Fallback if Altair is unavailable in the deployment environment.
            st.line_chart(
                standings_history.set_index("Round"),
                height=450,
                use_container_width=True,
            )

with tab1:
    st.subheader("Tournament Standings")
    display_df(
        standings,
        height=520,
        pink_last=True,
        scale_columns={
            "Draft Position": (
                standings["Draft Position"].min(),
                standings["Draft Position"].max(),
            ),
            "Total FPts": (
                standings["Total FPts"].min(),
                standings["Total FPts"].max(),
            ),
            "Active Players": (
                standings["Active Players"].min(),
                standings["Active Players"].max(),
            ),
            "Poss. Games Remaining": (
                standings["Poss. Games Remaining"].min(),
                standings["Poss. Games Remaining"].max(),
            ),
            "MVP FPts": (
                standings["MVP FPts"].min(),
                standings["MVP FPts"].max(),
            ),
        },
        reverse_columns={"Draft Position": True},
    )


    st.subheader("Power Rankings")
    display_df(
        power,
        height=520,
        rank_rwg=True,
        scale_columns={
            "Draft Position": (
                power["Draft Position"].min(),
                power["Draft Position"].max(),
            ),
            "Active Players": (
                power["Active Players"].min(),
                power["Active Players"].max(),
            ),
        },
        reverse_columns={"Draft Position": True},
    )


with tab2:
    st.subheader("Player Rankings — Top 10")
    display_df(players, height=400
    )

    st.markdown("### <span style='color:#2E8B57; font-weight:800;'>↑</span> Top Overperformers 🦄", unsafe_allow_html=True)
    over_display = over.rename(columns={"Delta From Average": "Delta From Season Avg."}).copy()
    under_display = under.rename(columns={"Delta From Average": "Delta From Season Avg."}).copy()

    # Format these two dashboard columns explicitly as display strings.
    # This avoids Streamlit/pandas rendering the underlying floats with
    # six decimal places on some pandas/Streamlit versions.
    for display_df_ in (over_display, under_display):
        if "Avg. FPts/Tourn. Game" in display_df_.columns:
            display_df_["Avg. FPts/Tourn. Game"] = display_df_["Avg. FPts/Tourn. Game"].apply(
                lambda x: f"{float(x):.1f}" if pd.notna(x) else ""
            )

    display_df(
        over_display,
        height=240,
        performance_gradient="green",
    )

    st.markdown("### <span style='color:#C62828; font-weight:800;'>↓</span> Top Underperformers 🪠", unsafe_allow_html=True)
    display_df(
        under_display,
        height=240,
        performance_gradient="red",
    )

    st.divider()
    st.subheader("Player Finder")
    selected_player = st.selectbox(
        "Search for a player",
        [""] + data["player_names"],
        index=0,
    )

    if selected_player:
        matching = [
            row for row in data["draft_rows"]
            if row["player_name"].casefold() == selected_player.casefold()
        ]
        if matching:
            row = matching[0]
            stats = data["player_stats"].get(selected_player.casefold(), {})
            result = pd.DataFrame([{
                "Draft Position": row["draft_position"],
                "Draftee": row["draftee"],
                "Team": stats.get("Team"),
                "Seed": stats.get("Seed"),
                "FPts": stats.get("Tot. FPts"),
                "PPG": stats.get("PPG"),
                "RPG": stats.get("RPG"),
                "APG": stats.get("APG"),
                "Tourn. FPts": player_total(row, data["fantasy_weights"]),
                "Tourn. PTS": row.get("tourn_pts"),
                "Tourn. REB": row.get("tourn_reb"),
                "Tourn. AST": row.get("tourn_ast"),
                "_team_out": (
                    stats.get("Team") is not None
                    and str(stats.get("Team")).strip().casefold() in data["teams_out"]
                ),
            }])
            for column in ["FPts", "PPG", "RPG", "APG", "Tourn. FPts"]:
                if column in result.columns:
                    result[column] = result[column].apply(
                        lambda value: "" if pd.isna(value) else f"{float(value):.1f}"
                    )
            st.caption(f"Showing Player Finder statistics for {selected_player}.")
            display_df(
                result,
                height=90,
                column_colors={
                    "FPts": LIGHT_BLUE,
                    "PPG": LIGHT_BLUE,
                    "RPG": LIGHT_BLUE,
                    "APG": LIGHT_BLUE,
                    "Tourn. FPts": LIGHT_GREEN,
                    "Tourn. PTS": LIGHT_GREEN,
                    "Tourn. REB": LIGHT_GREEN,
                    "Tourn. AST": LIGHT_GREEN,
                },
            )

with tab3:
    st.subheader("Team View")

    selected_draftee = st.selectbox(
        "Select a team owner",
        [""] + data["drafters"],
        index=0,
    )

    if selected_draftee:
        team_rows = make_team_rows(data, selected_draftee)
        team_df = pd.DataFrame(team_rows)

        # Team View only: display FPts, PPG, RPG, APG, and Tournament FPts
        # to exactly one decimal place.
        for column in ["FPts", "PPG", "RPG", "APG", "Tourn. FPts"]:
            if column in team_df.columns:
                team_df[column] = team_df[column].apply(
                    lambda value: "" if pd.isna(value) else f"{float(value):.1f}"
                )

        if team_df.empty:
            st.info("No drafted players found for this team owner.")
        else:
            st.caption(
                f"Showing {len(team_df)} player(s) drafted by {selected_draftee}."
            )
            display_df(
                team_df,
                height=500,
                column_colors={
                    "FPts": LIGHT_BLUE,
                    "PPG": LIGHT_BLUE,
                    "RPG": LIGHT_BLUE,
                    "APG": LIGHT_BLUE,
                    "Tourn. FPts": LIGHT_GREEN,
                    "Tourn. PTS": LIGHT_GREEN,
                    "Tourn. REB": LIGHT_GREEN,
                    "Tourn. AST": LIGHT_GREEN,
                },
            )

with tab4:
    st.subheader("Draft Ratings")
    st.caption(
        "Projected FPts = seed average games × season-average FPts, summed by draftee."
    )
    display_df(
        ratings,
        height=520,
        rank_rwg=True,
        scale_columns={
            "Draft Position": (
                ratings["Draft Position"].min(),
                ratings["Draft Position"].max(),
            ),
            "Projected FPts": (
                ratings["Projected FPts"].min(),
                ratings["Projected FPts"].max(),
            ),
        },
        reverse_columns={"Draft Position": True},
    )

with tab5:
    st.subheader("Tournament Bracket")
    st.caption("March Madness tournament bracket based on teams entered in the TeamList sheet.")

    if data["bracket_teams"]:
        try:
            # Build the bracket with teams_out data to determine winners
            rounds = build_tournament_bracket(data["bracket_teams"], data["teams_out"])

            # Create subtabs for regions and final four
            region_tabs = st.tabs(["🔵 EAST", "🟤 WEST", "🟢 MIDWEST", "🟡 SOUTH", "🏆 FINAL FOUR"])

            # Show each region's bracket (exclude Final Four and Championship)
            regions = ["EAST", "WEST", "MIDWEST", "SOUTH"]
            regional_rounds = rounds[:-2]  # All rounds except Final Four and Championship

            for tab_idx, (region_tab, region) in enumerate(zip(region_tabs[:4], regions)):
                with region_tab:
                    # Show Elite Eight winner box at the top
                    elite_eight_matchups = regional_rounds[-1]["regions"][region]
                    if elite_eight_matchups:
                        matchup = elite_eight_matchups[0]
                        team_a = matchup["team_a"]
                        team_b = matchup["team_b"]

                        # Determine winner
                        team_a_lost = lost_in_round(team_a, 3, data["teams_out"])
                        team_b_lost = lost_in_round(team_b, 3, data["teams_out"])

                        winner = None
                        if team_a_lost and not team_b_lost:
                            winner = team_b
                        elif team_b_lost and not team_a_lost:
                            winner = team_a
                        elif not team_a_lost and not team_b_lost:
                            winner = None  # Game not played yet

                        # Display winner box at the top
                        col1, col2, col3 = st.columns([1.3, 1, 1.3])
                        with col2:
                            components.html(
                                render_regional_winner_card_html(region, winner),
                                height=90,
                            )

                    st.markdown("---")

                    # Show bracket
                    bracket_html = render_bracket_html(regional_rounds, region=region)
                    components.html(bracket_html, height=1200, scrolling=True)

            # Final Four tab - show only Final Four and Championship
            with region_tabs[4]:
                # Compute regional winners
                regional_winners = {}
                for region in regions:
                    elite_eight_matchups = regional_rounds[-1]["regions"][region]
                    if elite_eight_matchups:
                        matchup = elite_eight_matchups[0]
                        team_a = matchup["team_a"]
                        team_b = matchup["team_b"]

                        # Determine winner
                        team_a_lost = lost_in_round(team_a, 3, data["teams_out"])
                        team_b_lost = lost_in_round(team_b, 3, data["teams_out"])

                        winner = None
                        if team_a_lost and not team_b_lost:
                            winner = team_b
                        elif team_b_lost and not team_a_lost:
                            winner = team_a
                        # else: winner is None (TBD)

                        regional_winners[region] = winner if winner else "TBD"

                # Build Final Four bracket with configured region pairings
                final_four_rounds = build_final_four_brackets(
                    regional_winners, 
                    data["final_four_config"], 
                    data["teams_out"]
                )
                bracket_html = render_bracket_html(final_four_rounds)
                components.html(bracket_html, height=550, scrolling=False)

        except Exception as e:
            st.error(f"Error rendering bracket: {str(e)}")
            st.write(f"Details: {e}")
    else:
        st.warning("No teams found in the TeamList sheet.")
        st.info("Make sure your Excel file has a 'TeamList' sheet with columns:\n- **Region** (EAST, WEST, MIDWEST, SOUTH)\n- **Seed** (1-16)\n- **Team** (team name)")

st.divider()
st.caption(
    "The dashboard reads the Excel workbook from OneDrive. "
    "No workbook changes are made by the app."
)



